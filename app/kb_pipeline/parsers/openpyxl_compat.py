from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation


def _patch_data_validation_id() -> None:
    if getattr(DataValidation, "_kb_accepts_wps_id_attr", False):
        return

    original_init = DataValidation.__init__

    def patched_init(self: DataValidation, *args: Any, **kwargs: Any) -> None:
        # WPS/modern Excel files may add a dataValidation id attribute. It does
        # not affect cell values, but openpyxl 3.1.x cannot deserialize it.
        kwargs.pop("id", None)
        original_init(self, *args, **kwargs)

    DataValidation.__init__ = patched_init  # type: ignore[method-assign]
    DataValidation._kb_accepts_wps_id_attr = True  # type: ignore[attr-defined]


def load_workbook_compat(path: str | Path, **kwargs: Any):
    _patch_data_validation_id()
    return load_workbook(path, **kwargs)
