"""Parsing layer: parse results and merge rules for MinerU / native tables / HTML / images with VLM
captions / code files."""
from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from kb_pipeline.models import KBSource, ParsedBlock, SourceFile
from kb_pipeline.parsers.common import page_idx, parser_profile_for_path
from kb_pipeline.parsers.html_dom import parse_html_dom
from kb_pipeline.pipeline.detect_changes import detect_change
from kb_pipeline.pipeline.parse_job import _is_intentionally_empty
from kb_pipeline.vision import vlm

from _support import _CodexAudit20260906TestsSupport, _CodexFinalTestsSupport, _block, _mb, _repo_file, _text_of_tokens


class ParserRegressionTests(unittest.TestCase):
    def test_mineru_page_indexes_are_one_based(self) -> None:
        self.assertEqual(page_idx({"page_idx": 0}), 1)
        self.assertEqual(page_idx({"page_idx": 1}), 2)
        self.assertEqual(page_idx({"page_id": 3}), 4)
        self.assertEqual(page_idx({"page_no": 1}), 1)

    def test_extension_changing_rename_requires_parse(self) -> None:
        old = {
            "status": "indexed",
            "content_version": "same-checksum",
            "size": 12,
            "indexed_version": "same-checksum",
            "indexed_parser_profile": parser_profile_for_path(Path("before.txt")),
            "filename": "before.txt",
            "metadata_fingerprint": "old-metadata",
        }
        new = SourceFile(
            kb_id="project_materials",
            collection="kb_project",
            source_root="项目资料",
            source_type="local_mirror",
            file_key=1,
            source_path="项目资料/after.html",
            rel_path="after.html",
            filename="after.html",
            dir="",
            physical_path="/tmp/after.html",
            mime_type="text/html",
            size=12,
            mtime=1,
            checksum="same-checksum",
        )
        change = detect_change(old, new)
        self.assertEqual(change.change_type, "parser_changed")
        self.assertEqual(change.job_type, "parse")

    def test_empty_text_is_valid_but_empty_pdf_is_not(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            empty_text = root / "empty.py"
            empty_pdf = root / "empty.pdf"
            empty_text.write_text("\n\t", encoding="utf-8")
            empty_pdf.write_bytes(b"")
            self.assertTrue(_is_intentionally_empty(empty_text))
            self.assertFalse(_is_intentionally_empty(empty_pdf))

    def test_empty_files_of_every_text_suffix_are_valid(self) -> None:
        """.mjs / .cjs / .lua / .cxx / .hh are supported by scan and parse, yet they were missing from the table of
        suffixes whose files may legitimately be empty: an empty placeholder file was recorded as a parse failure,
        while an equally empty .js finished normally."""
        from kb_pipeline.localfs.scanner import SUPPORTED_EXTS
        from kb_pipeline.parsers.router import parse_native
        from kb_pipeline.vision.images import IMAGE_SUFFIXES

        binary = {".pdf", ".pptx", ".docx", ".xlsx", ".xls", *IMAGE_SUFFIXES}
        with tempfile.TemporaryDirectory() as tmp:
            for suffix in sorted(SUPPORTED_EXTS - binary - {".html"}):          # an empty .html parses into a block holding only the title
                path = Path(tmp) / f"empty{suffix}"
                path.write_text("\n", encoding="utf-8")
                self.assertTrue(_is_intentionally_empty(path), suffix)
            for suffix in (".mjs", ".lua", ".hh"):
                self.assertEqual(parse_native(Path(tmp) / f"empty{suffix}"), [], suffix)

    def test_nested_special_html_block_is_not_duplicated(self) -> None:
        html_text = """
        <html><head><title>Test</title></head><body>
          <section class="scene">
            <h2 class="scene-title">Outer scene</h2>
            <div class="faq-item"><span class="q-text">Question</span><div class="faq-a">Answer</div></div>
            <p>This paragraph provides enough substantial body text for deterministic DOM extraction behavior.</p>
          </section>
        </body></html>
        """
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "nested.html"
            path.write_text(html_text, encoding="utf-8")
            blocks = parse_html_dom(path)
        self.assertEqual([block.block_type for block in blocks], ["scene"])


class ParseCacheTests(unittest.TestCase):
    def test_vlm_cache_depends_on_model_and_prompt(self) -> None:
        calls: list[dict] = []

        class FakeCompletions:
            def create(self, **kwargs):
                calls.append(kwargs)
                message = SimpleNamespace(
                    content=json.dumps(
                        {"summary": "ok", "entities": [], "facts": [], "keywords": [], "confidence": "high"}
                    )
                )
                return SimpleNamespace(choices=[SimpleNamespace(message=message)])

        class FakeOpenAI:
            def __init__(self, **kwargs):
                self.chat = SimpleNamespace(completions=FakeCompletions())

        with tempfile.TemporaryDirectory() as tmp, patch.object(vlm, "OpenAI", FakeOpenAI):
            root = Path(tmp)
            image = root / "image.png"
            cache = root / "cache.json"
            image.write_bytes(b"not-a-real-image")
            first = vlm.caption_image(
                image_path=image,
                base_url="http://example.invalid",
                api_key="test",
                model_id="model-a",
                prompt="prompt-a",
                cache_json=cache,
            )
            second = vlm.caption_image(
                image_path=image,
                base_url="http://example.invalid",
                api_key="test",
                model_id="model-a",
                prompt="prompt-a",
                cache_json=cache,
            )
            third = vlm.caption_image(
                image_path=image,
                base_url="http://example.invalid",
                api_key="test",
                model_id="model-a",
                prompt="prompt-b",
                cache_json=cache,
            )
        self.assertFalse(first["vlm_cache_hit"])
        self.assertTrue(second["vlm_cache_hit"])
        self.assertFalse(third["vlm_cache_hit"])
        self.assertEqual(len(calls), 2)

    def test_mineru_request_defaults_to_the_engine_this_image_ships(self) -> None:
        """Without MINERU_BACKEND the request carries the backend the parser container published (a backend the
        parser cannot serve is rejected with 409); an explicit setting takes effect as usual."""
        from kb_pipeline.parsers import service_clients

        sent: list[dict] = []

        class FakeSession:
            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

            def post(self, url, *, files, data, timeout):
                sent.append(dict(data))
                return SimpleNamespace(raise_for_status=lambda: None, json=lambda: {"md_content": "# 标题"})

        with tempfile.TemporaryDirectory() as tmp, patch.object(service_clients, "_local_service_session", FakeSession), \
                patch.dict(os.environ, {"MINERU_REUSE_CACHE": "0", "MINERU_INFO_FILE": str(Path(tmp) / "carrel-mineru.json")}):
            (Path(tmp) / "carrel-mineru.json").write_text(json.dumps({"backend": "vlm-engine"}), encoding="utf-8")
            source = Path(tmp) / "a.pdf"
            source.write_bytes(b"%PDF-1.4")
            os.environ.pop("MINERU_BACKEND", None)
            service_clients.call_mineru_sync("http://127.0.0.1:8765", source, Path(tmp) / "o.json", Path(tmp) / "o.md")
            os.environ["MINERU_BACKEND"] = "pipeline"
            service_clients.call_mineru_sync("http://127.0.0.1:8765", source, Path(tmp) / "o.json", Path(tmp) / "o.md")
        self.assertEqual([d["backend"] for d in sent], ["vlm-engine", "pipeline"])


class MarkdownRoutingTests(unittest.TestCase):
    def test_markdown_splits_on_headings_and_carries_section_path(self) -> None:
        from kb_pipeline.parsers.common import markdown_to_blocks

        md = "# 卷二\n\n前言\n\n## 核心战略\n\n战略正文\n\n```bash\n# 这是注释不是标题\npip install x\n```\n\n## 市场地位\n\n地位正文"
        blocks = markdown_to_blocks(md, parser="native", parser_profile="md-sections-v1", doc_type="md")
        self.assertEqual([b.title for b in blocks], ["卷二", "核心战略", "市场地位"])
        self.assertEqual(blocks[1].metadata["section_path"], ["卷二", "核心战略"])
        # heading carried by title/section_path, not duplicated into the body
        self.assertFalse(any(b.text.startswith("#") for b in blocks))
        # a "# comment" inside a fence is code, not a heading
        self.assertNotIn("这是注释", " ".join(b.title or "" for b in blocks))
        self.assertIn("# 这是注释不是标题", blocks[1].text)

    def test_router_sends_markdown_through_section_splitter(self) -> None:
        from kb_pipeline.parsers.router import parse_native

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "doc.md"
            p.write_text("# A\n\nx\n\n## B\n\ny\n", encoding="utf-8")
            blocks = parse_native(p)
        self.assertEqual(len(blocks), 2)
        self.assertEqual(blocks[0].parser_profile, "md-sections-v3")
        self.assertEqual(blocks[1].metadata["section_path"], ["A", "B"])


class TableHeaderTests(unittest.TestCase):
    def _rows(self, *rows):
        return [(i + 1, list(r)) for i, r in enumerate(rows)]

    def test_tiered_header_is_merged_on_verbatim_repeat(self) -> None:
        from kb_pipeline.parsers.native_table import detect_header_rows, merge_header_rows

        rows = self._rows(
            ["镜像名称", "v6.0应用版本", "v6.0应用版本", "备注"],
            ["镜像名称", "x86_64", "arm64", "备注"],
            ["Kylin-Server-10", "SP2B09-x86-20210524", "SP2B09-aarch64-20210524", "单机"],
        )
        hdr = detect_header_rows(rows)
        self.assertEqual([rn for rn, _ in hdr], [1, 2])
        self.assertEqual(merge_header_rows([r for _, r in hdr]), ["镜像名称", "v6.0应用版本/x86_64", "v6.0应用版本/arm64", "备注"])

    def test_vertical_merge_from_the_header_row_admits_a_version_shaped_second_tier(self) -> None:
        """Codex 2026-09-14 R02: the second row of a feature comparison table holds product column names (with
        version strings), which the shape check treated as data; the vertical merge of C1:C2 in the workbook
        shows it is a header."""
        from kb_pipeline.parsers.native_table import chunk_rows_to_blocks, detect_header_rows, merge_header_rows, vertical_spans

        rows = self._rows(
            ["一级模块", "二级模块", "功能", "子功能", "支持情况", "", ""],
            ["", "", "", "", "示例协作v4.3->v4.11", "友商A", "友商B"],
            ["公共能力", "登录", "普通登录方式", "微信扫码登录", "1", "0", "1"],
        )
        self.assertEqual([rn for rn, _ in detect_header_rows(rows)], [1])                       # shape only: the second row is taken as data
        spans = vertical_spans([(3, 1, 3, 2), (5, 1, 5, 2), (2, 137, 2, 138), (1, 1, 1, 60)])     # a 60-row category label does not count; the 2-row merge in the data area is recorded but unused by the header check
        self.assertEqual(spans, {1: 2, 137: 138})
        hdr = detect_header_rows(rows, spans=spans)
        self.assertEqual([rn for rn, _ in hdr], [1, 2])
        merged = merge_header_rows([r for _, r in hdr])
        self.assertTrue(any("友商A" in h for h in merged) and any("示例协作v4.3->v4.11" in h for h in merged), merged)
        blocks = chunk_rows_to_blocks("xlsx", rows, "对比", max_tokens=200, overlap_tokens=0, header_spans=spans)
        body = [b for b in blocks if not b.metadata.get("summary")]
        self.assertTrue(all("友商A" in b.text.splitlines()[2] for b in body), [b.text[:120] for b in body])   # the HEADER row carries the product column names
        self.assertNotIn("示例协作v4.3->v4.11 | 友商A", "\n".join(l for b in body for l in b.text.splitlines()[3:]))   # no longer emitted as a data row
        # Without merge footprints the behaviour is unchanged: a data-shaped second row stays out of the header
        self.assertEqual([rn for rn, _ in detect_header_rows(rows, spans={})], [1])

    def test_plain_data_row_never_joins_the_header(self) -> None:
        from kb_pipeline.parsers.native_table import detect_header_rows

        rows = self._rows(
            ["PC 平台", "支持浏览器", "版本", "备注"],
            ["Windows", "Chrome", "Chrome ≥ 80", ""],   # looks label-like but shares no cell verbatim
        )
        self.assertEqual(len(detect_header_rows(rows)), 1)
        rows = self._rows(
            ["功能", "V7 20260122", "上线版本"],
            ["登录", "✅", "v7.0.2412a"],                 # date / version shaped cells
        )
        self.assertEqual(len(detect_header_rows(rows)), 1)

    def test_prefix_is_budgeted_and_long_rows_are_split(self) -> None:
        from kb_pipeline.parsers.native_table import chunk_rows_to_blocks
        from kb_pipeline.utils import count_tokens

        header = [f"列{i}" for i in range(20)]
        long_cell = "这是一个非常长的功能描述单元格。" * 40
        rows = self._rows(header, *[[f"值{i}" for i in range(20)] for _ in range(30)], ["长", long_cell])
        blocks = chunk_rows_to_blocks("xlsx", rows, "sheet", max_tokens=200, overlap_tokens=0)
        body = [b for b in blocks if not b.metadata.get("summary")]
        pieces = [b for b in body if "这是一个非常长" in b.text]
        self.assertTrue(all("HEADER:" in b.text for b in body if b not in pieces))
        # the block text (prefix included) respects max_tokens; a row that alone
        # exceeds it is split at cell boundaries, an over-long cell by sentence
        self.assertTrue(all(count_tokens(b.text) <= 200 for b in body), [count_tokens(b.text) for b in body])
        # every split piece already carries the column names and the row identity, not the whole header
        self.assertGreater(len(pieces), 1)
        self.assertTrue(all(b.text.splitlines()[2].startswith("列0: 长 | 列1: 这是一个非常长的功能描述单元格。") for b in pieces))
        self.assertTrue(all("HEADER:" not in b.text for b in pieces))
        self.assertEqual("".join(b.text.splitlines()[2].split("列1: ", 1)[1] for b in pieces), long_cell)   # nothing lost, nothing repeated


class ImageNormalizationTests(unittest.TestCase):
    def _png(self, path: Path, size, mode="RGB") -> None:
        from PIL import Image

        Image.new(mode, size, (200, 30, 30) if mode == "RGB" else (200, 30, 30, 128)).save(path, format="PNG")

    def test_small_png_passes_through_untouched(self) -> None:
        from kb_pipeline.vision.images import load_image_for_model

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "a.png"; self._png(p, (64, 48))
            data, mime = load_image_for_model(p, max_pixels=3_686_400)
            self.assertEqual((data, mime), (p.read_bytes(), "image/png"))

    def test_oversized_image_is_downscaled_within_budget(self) -> None:
        from PIL import Image

        from kb_pipeline.vision.images import load_image_for_model

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "big.png"; self._png(p, (1200, 900))
            data, mime = load_image_for_model(p, max_pixels=100_000)
            self.assertEqual(mime, "image/png")
            with Image.open(__import__("io").BytesIO(data)) as im:
                self.assertLessEqual(im.width * im.height, 100_000)
                self.assertAlmostEqual(im.width / im.height, 1200 / 900, places=1)

    def test_bmp_and_gif_are_reencoded_as_png(self) -> None:
        from PIL import Image

        from kb_pipeline.vision.images import image_data_url, load_image_for_model

        with tempfile.TemporaryDirectory() as tmp:
            bmp = Path(tmp) / "x.bmp"; Image.new("RGB", (20, 20), (1, 2, 3)).save(bmp, format="BMP")
            gif = Path(tmp) / "x.gif"; Image.new("P", (20, 20)).save(gif, format="GIF")
            for p in (bmp, gif):
                data, mime = load_image_for_model(p)
                self.assertEqual(mime, "image/png"); self.assertTrue(data.startswith(b"\x89PNG"))
            self.assertTrue(image_data_url(bmp).startswith("data:image/png;base64,"))


class VisualEmbeddingClientTests(unittest.TestCase):
    def _client_with_fake_server(self, dim, calls):
        from kb_pipeline.embedding import visual

        class FakeResponse:
            status_code = 200
            text = ""

            def __init__(self, body):
                self._body = body

            def json(self):
                return self._body

        class FakeSession:
            def post(self, url, json=None, headers=None, timeout=None, allow_redirects=True):
                calls.append({"url": url, "json": json, "headers": headers})
                return FakeResponse({"data": [{"embedding": [0.5] * dim}]})

        return visual, patch.object(visual, "_session_for", return_value=FakeSession())

    def test_request_shape_uses_chat_messages_with_system_instruction(self) -> None:
        from PIL import Image

        calls: list = []
        visual, patched = self._client_with_fake_server(8, calls)
        with tempfile.TemporaryDirectory() as tmp, patched:
            img = Path(tmp) / "f.png"; Image.new("RGB", (10, 10)).save(img)
            client = visual.VisualEmbeddingClient(base_url="http://h/v1", api_key="k", model_id="m", dim=8, retry=1)
            vec = client.embed_image(img)
        self.assertEqual(len(vec), 8)
        body = calls[0]["json"]
        self.assertEqual(calls[0]["url"], "http://h/v1/embeddings")
        self.assertEqual(body["model"], "m"); self.assertEqual(body["dimensions"], 8)
        self.assertNotIn("input", body)  # the plain input path lands in a different space
        self.assertEqual(body["messages"][0], {"role": "system", "content": visual.DEFAULT_INSTRUCTION})
        content = body["messages"][1]["content"]
        self.assertEqual(content[0]["type"], "image_url")
        self.assertTrue(content[0]["image_url"]["url"].startswith("data:image/png;base64,"))
        self.assertEqual(calls[0]["headers"]["Authorization"], "Bearer k")

    def test_cache_reuse_and_dim_mismatch(self) -> None:
        from PIL import Image

        calls: list = []
        visual, patched = self._client_with_fake_server(8, calls)
        with tempfile.TemporaryDirectory() as tmp, patched:
            img = Path(tmp) / "f.png"; Image.new("RGB", (10, 10)).save(img)
            cache = Path(tmp) / "c" / "f.embed.json"
            client = visual.VisualEmbeddingClient(base_url="http://h/v1", api_key="", model_id="m", dim=8, retry=1)
            first = client.embed_image(img, cache_json=cache)
            second = client.embed_image(img, cache_json=cache)
            self.assertEqual(first, second); self.assertEqual(len(calls), 1)  # served from cache
            self.assertTrue(cache.exists())
            # a different model id invalidates the cache
            other = visual.VisualEmbeddingClient(base_url="http://h/v1", api_key="", model_id="m2", dim=8, retry=1)
            other.embed_image(img, cache_json=cache); self.assertEqual(len(calls), 2)
            # server returning the wrong width is an error, not silently stored
            wrong = visual.VisualEmbeddingClient(base_url="http://h/v1", api_key="", model_id="m3", dim=16, retry=1)
            with self.assertRaises(RuntimeError):
                wrong.embed_image(img)
            # embed_images keeps order
            vecs = client.embed_images([(img, None), (img, None)])
            self.assertEqual(len(vecs), 2)


class ImageFileParserTests(unittest.TestCase):
    def test_png_is_supported_and_routed_to_image_profile(self) -> None:
        from kb_pipeline.localfs.scanner import SUPPORTED_EXTS

        for ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"):
            self.assertIn(ext, SUPPORTED_EXTS, ext)
            self.assertEqual(parser_profile_for_path(Path(f"x{ext}")), "image-vlm-v1", ext)
        self.assertNotIn(".svg", SUPPORTED_EXTS); self.assertNotIn(".heic", SUPPORTED_EXTS)

    def test_standalone_image_becomes_one_visual_block_under_cache_root(self) -> None:
        from PIL import Image

        from kb_pipeline.parsers import image_file

        def fake_enrich(blocks, **kwargs):
            self.assertFalse(kwargs["filter_decorative"])
            self.assertIn("14-finance-workflows.png", kwargs["prompt"])
            for b in blocks:
                b.text = "VISUAL SUMMARY: 流程图"; b.visual_summary = "流程图"
                b.metadata.update({"vlm_status": "success", "visual_sha256": "ab" * 32})
            return blocks

        with tempfile.TemporaryDirectory() as tmp, patch.object(image_file, "enrich_blocks_with_vlm", fake_enrich):
            src = Path(tmp) / "mirror" / "14-finance-workflows.png"; src.parent.mkdir()
            Image.new("RGB", (32, 16)).save(src)
            cache = Path(tmp) / "cache"
            blocks = image_file.parse_image_file(
                path=src, cache_dir=cache, vlm_base_url="u", vlm_api_key="k", vlm_model_id="m", vlm_concurrency=1,
            )
        self.assertEqual(len(blocks), 1)
        b = blocks[0]
        self.assertEqual((b.block_type, b.parser_profile, b.doc_type), ("image", "image-vlm-v1", "png"))
        self.assertEqual(b.title, "14-finance-workflows")
        self.assertEqual(Path(b.visual_ref), cache / "image" / "14-finance-workflows.png")
        self.assertEqual(b.metadata["image_width"], 32)

    def test_failed_caption_fails_the_job(self) -> None:
        from PIL import Image

        from kb_pipeline.parsers import image_file

        def failing_enrich(blocks, **kwargs):
            for b in blocks:
                b.metadata.update({"vlm_status": "failed", "vlm_error": "boom"})
            return blocks

        with tempfile.TemporaryDirectory() as tmp, patch.object(image_file, "enrich_blocks_with_vlm", failing_enrich):
            src = Path(tmp) / "p.png"; Image.new("RGB", (8, 8)).save(src)
            with self.assertRaises(RuntimeError):
                image_file.parse_image_file(path=src, cache_dir=Path(tmp) / "c", vlm_base_url="u",
                                            vlm_api_key="k", vlm_model_id="m", vlm_concurrency=1)
            # garbage with an image suffix is rejected up front
            bad = Path(tmp) / "bad.jpg"; bad.write_bytes(b"not an image")
            with self.assertRaises(RuntimeError):
                image_file.parse_image_file(path=bad, cache_dir=Path(tmp) / "c", vlm_base_url="u",
                                            vlm_api_key="k", vlm_model_id="m", vlm_concurrency=1)

    def test_failures_that_a_retry_cannot_fix_do_not_retry(self) -> None:
        """A broken image, a strip the model cannot take: these used to raise a plain exception, were retried 5
        times with backoff, and once the cooldown was over the scan queued them again. A failed model call or a
        missing API key remain retryable."""
        from PIL import Image

        from kb_pipeline.parsers import image_file
        from kb_pipeline.pipeline.worker import _should_retry

        settings = SimpleNamespace(job_max_retries=5)
        job = {"retry_count": 0}

        def failure(path, api_key="k"):
            with self.assertRaises(RuntimeError) as caught:
                image_file.parse_image_file(path=path, cache_dir=path.parent / "cache", vlm_base_url="u",
                                            vlm_api_key=api_key, vlm_model_id="m", vlm_concurrency=1)
            return _should_retry(job, settings, caught.exception), str(caught.exception)

        def failing_enrich(blocks, **kwargs):
            for b in blocks:
                b.metadata.update({"vlm_status": "failed", "vlm_error": "boom"})
            return blocks

        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "bad.jpg"; bad.write_bytes(b"not an image")
            strip = Path(tmp) / "divider.png"; Image.new("RGB", (3000, 20)).save(strip)
            good = Path(tmp) / "fig.png"; Image.new("RGB", (64, 48)).save(good)
            self.assertEqual(failure(bad)[0], False)
            retry, message = failure(strip)
            self.assertEqual((retry, message), (False, "VLM caption skipped for image file divider.png: degenerate_aspect"))
            self.assertEqual(failure(good, api_key=""), (True, "VLM caption skipped for image file fig.png: missing_api_key"))
            with patch.object(image_file, "enrich_blocks_with_vlm", failing_enrich):
                self.assertEqual(failure(good)[0], True)
            # A file that cannot be opened is an environment problem (permissions, disk), not a broken image: retried as usual
            with patch.object(Path, "open", side_effect=PermissionError("permission denied")), self.assertRaises(PermissionError) as caught:
                image_file.parse_image_file(path=good, cache_dir=Path(tmp) / "cache", vlm_base_url="u",
                                            vlm_api_key="k", vlm_model_id="m", vlm_concurrency=1)
            self.assertTrue(_should_retry(job, settings, caught.exception))


class VisualChunkSelectionTests(unittest.TestCase):
    def test_first_chunk_of_each_vlm_seen_picture_is_selected(self) -> None:
        from kb_pipeline.models import ParsedBlock, UnifiedChunk
        from kb_pipeline.pipeline.parse_job import select_visual_chunks

        with tempfile.TemporaryDirectory() as tmp:
            img = Path(tmp) / "i.png"; img.write_bytes(b"x")
            seen = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="image", text="t",
                               block_id="img1", visual_ref=str(img), metadata={"vlm_status": "success"})
            skipped = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="image", text="t",
                                  block_id="img2", visual_ref=str(img), metadata={"vlm_status": "skipped"})
            failed = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="image", text="t",
                                 block_id="img3", visual_ref=str(img), metadata={"vlm_status": "failed"})
            table = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="table", text="t",
                                block_id="tab", visual_ref=str(img), metadata={"vlm_status": "success"})
            chunks = [
                UnifiedChunk("k:1:v:p:img1:0", 0, "a", seen), UnifiedChunk("k:1:v:p:img1:1", 1, "b", seen),
                UnifiedChunk("k:1:v:p:img2:0", 2, "c", skipped), UnifiedChunk("k:1:v:p:img3:0", 3, "d", failed),
                UnifiedChunk("k:1:v:p:tab:0", 5, "f", table),
            ]
            picked = [c.chunk_uid for c in select_visual_chunks(chunks)]
            # one vector per picture; skipped/decorative and tables stay out;
            # a picture whose caption failed still went through the VLM and is embedded
            self.assertEqual(picked, ["k:1:v:p:img1:0", "k:1:v:p:img3:0"])
            # a VLM-seen block whose crop vanished mid-parse is a loud failure,
            # not a silent hole in the visual index
            gone = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="image", text="t",
                               block_id="img4", visual_ref=str(Path(tmp) / "gone.png"), metadata={"vlm_status": "success"})
            with self.assertRaises(RuntimeError):
                select_visual_chunks([UnifiedChunk("k:1:v:p:img4:0", 0, "e", gone)])

    def test_payload_marks_points_that_carry_the_visual_vector(self) -> None:
        from kb_pipeline.models import ParsedBlock, UnifiedChunk
        from kb_pipeline.pipeline.parse_job import _payload_for_chunk

        settings = SimpleNamespace(cache_dir=Path("/cache"), embedding_model_id="m")
        row = {"kb_id": "k", "file_key": 1, "source_path": "r/a.png", "rel_path": "a.png", "filename": "a.png",
               "dir": "", "content_version": "v", "size": 1, "mtime": 1, "mime_type": "image/png"}
        block = ParsedBlock(parser="image-file", parser_profile="image-vlm-v1", doc_type="png", block_type="image",
                            text="VISUAL SUMMARY: s", block_id="image-0001", title="a",
                            visual_ref="/cache/parse/k/1/v/image/a.png", visual_summary="s",
                            metadata={"visual_sha256": "cd" * 32, "vlm_status": "success"})
        chunk = UnifiedChunk("u", 0, "t", block, 1)
        with_vec = _payload_for_chunk(settings, row, chunk, 1, visual_model="qwen3-vl-embedding-2b")
        without = _payload_for_chunk(settings, row, chunk, 1)
        self.assertEqual(with_vec["visual_embedding_model"], "qwen3-vl-embedding-2b")
        self.assertEqual(with_vec["visual_sha256"], "cd" * 32)
        self.assertEqual(with_vec["visual_ref"], "parse/k/1/v/image/a.png")
        self.assertNotIn("visual_embedding_model", without)

    def test_degenerate_crops_are_skipped_instead_of_failing_the_file(self) -> None:
        """2026-09-11 kb_005: MinerU cropped a 797×2 horizontal rule out of a docx, the Qwen3-VL processor
        answered a plain 400, and "VLM caption failed 1/1" sent the whole file into endless retries.
        Images with an aspect ratio above 100 are not sent to the model: skipped, not counted as failures, and
        folded away as decorative on the body route; a standalone image file only skips the caption."""
        from kb_pipeline.models import ParsedBlock
        from kb_pipeline.parsers import visual_blocks
        from kb_pipeline.vision.filter import degenerate_image

        self.assertEqual(degenerate_image(797, 2), "degenerate_aspect")
        self.assertEqual(degenerate_image(40, 6000), "degenerate_aspect")
        self.assertIsNone(degenerate_image(797, 40))
        self.assertIsNone(degenerate_image(16, 16))                                 # the processor upscales small images; not degenerate
        self.assertIsNone(degenerate_image(None, None))
        calls: list[int] = []

        def fake_caption(jobs, **kwargs):
            calls.append(len(jobs))
            return {job_id: {"kind": "diagram", "summary": "s", "text_verbatim": "", "entities": [], "facts": [], "keywords": [], "confidence": "high"}
                    for job_id, *_ in jobs}

        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp, patch.object(visual_blocks, "caption_images_parallel", fake_caption):
            strip = Path(tmp) / "line.jpg"; Image.new("RGB", (797, 2)).save(strip)
            real = Path(tmp) / "fig.png"; Image.new("RGB", (400, 300)).save(real)
            mk = lambda bid, ref: ParsedBlock(parser="p", parser_profile="p", doc_type="docx", block_type="image", text="", block_id=bid, visual_ref=str(ref))
            a, b = mk("line", strip), mk("fig", real)
            visual_blocks.enrich_blocks_with_vlm([a, b], base_url="u", api_key="k", model_id="m", cache_dir=Path(tmp), filter_decorative=True)
            self.assertEqual(calls, [1])                                            # only the normal one went to the model
            self.assertEqual((a.metadata["vlm_status"], a.metadata["vlm_skip_reason"], a.metadata["decorative"]), ("skipped", "degenerate_aspect", True))
            self.assertFalse(visual_blocks.went_through_vlm(a))
            self.assertEqual(visual_blocks.vlm_failure_summary([a, b]), {"visual_blocks": 1, "vlm_failed": 0})
            c = mk("alone", strip)
            visual_blocks.enrich_blocks_with_vlm([c], base_url="u", api_key="k", model_id="m", cache_dir=Path(tmp), filter_decorative=False)
            self.assertEqual(c.metadata["vlm_status"], "skipped")
            self.assertNotIn("decorative", c.metadata)                              # standalone image file: caption skipped only, not folded away

    def test_decorative_filter_before_the_vlm(self) -> None:
        """Before the model is called only size and file name decide; logos / watermarks have no reliable feature
        up front (cropped images are named by content hash), so the model judges them."""
        from kb_pipeline.vision.filter import conservative_decorative_filter as decide

        self.assertEqual(decide(width=79, height=79, object_name="ab12.jpg").decorative_reason, "tiny_icon")
        self.assertFalse(decide(width=79, height=80, object_name="ab12.jpg").decorative_skip)
        self.assertEqual(decide(width=800, height=600, object_name="page-background.png").decorative_reason, "decorative_name")
        self.assertFalse(decide(width=800, height=600, object_name="company-logo.png").decorative_skip)
        self.assertFalse(decide(width=None, height=None).decorative_skip)

    def test_vlm_title_appears_once_as_caption(self) -> None:
        from kb_pipeline.chunking.chunker import text_for_block
        from kb_pipeline.models import ParsedBlock
        from kb_pipeline.parsers import visual_blocks

        def fake_caption(jobs, **kwargs):
            return {job_id: {"kind": "diagram", "title": "六个财务 Agent", "summary": "摘要", "text_verbatim": "",
                             "entities": [], "facts": ["f1"], "keywords": ["k1"], "confidence": "high"}
                    for job_id, *_ in jobs}

        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp, patch.object(visual_blocks, "caption_images_parallel", fake_caption):
            img = Path(tmp) / "i.png"; Image.new("RGB", (16, 16)).save(img)
            no_caption = ParsedBlock(parser="p", parser_profile="p", doc_type="png", block_type="image", text="",
                                     block_id="a", title="14-finance-workflows", visual_ref=str(img))
            with_caption = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="image", text="",
                                       block_id="b", caption="图 3 架构", visual_ref=str(img))
            visual_blocks.enrich_blocks_with_vlm([no_caption, with_caption], base_url="u", api_key="k", model_id="m",
                                                 cache_dir=Path(tmp), filter_decorative=False)
            expected_sha = visual_blocks.file_hash(img)
        t1 = text_for_block(no_caption)
        self.assertEqual(t1.count("六个财务 Agent"), 1)          # once, as CAPTION
        self.assertIn("CAPTION: 六个财务 Agent", t1)
        self.assertEqual(t1.count("摘要"), 1)                     # summary once
        self.assertTrue(t1.startswith("TITLE: 14-finance-workflows"))
        t2 = text_for_block(with_caption)
        self.assertIn("CAPTION: 图 3 架构", t2); self.assertIn("TITLE: 六个财务 Agent", t2)  # both kept: different info
        self.assertEqual(no_caption.metadata["visual_sha256"], expected_sha)  # picture hash recorded for the payload
        self.assertTrue(visual_blocks.went_through_vlm(no_caption))

    def test_visual_summary_not_repeated_in_chunk_text(self) -> None:
        from kb_pipeline.chunking.chunker import text_for_block
        from kb_pipeline.models import ParsedBlock

        block = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="image",
                            text="TITLE: x\nVISUAL SUMMARY: 一张架构图", block_id="b", visual_summary="一张架构图")
        self.assertEqual(text_for_block(block).count("一张架构图"), 1)
        bare = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="image",
                           text="", block_id="b", caption="c", visual_summary="一张架构图")
        self.assertIn("VISUAL SUMMARY: 一张架构图", text_for_block(bare))


class EquationMergeTests(unittest.TestCase):
    """Equations must be mergeable into the body-text buffer.

    MinerU breaks every connecting sentence between equations into its own text block. As long as equations
    cannot be merged, the buffer is flushed at every equation -- "where / $$equation$$ / therefore" becomes
    three chunks, two of which hold a single token. Measured on the library KB: 1,424 of 44,043 chunks were
    ≤5 tokens.

    But the real loss is at the other end: 14,338 equation chunks were separated from the prose explaining
    them, and nobody searches in LaTeX -- what makes an equation findable is precisely the two sentences
    around it.
    """

    def test_an_equation_no_longer_splits_the_paragraph(self) -> None:
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks

        merged = merge_mineru_text_blocks(
            [_mb("t1", "其中"),
             _mb("e1", "x^2", block_type="equation", latex="x^2"),
             _mb("t2", "所以")],
            target_tokens=800,
        )
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0].metadata["merged_block_count"], 3)

    def test_the_equation_marker_survives_the_merge(self) -> None:
        """The buffer joins the rendered text, not block.text.

        An equation's content lives in latex; block.text is either empty or the same latex -- joining
        block.text directly would drop the "EQUATION: " marker fed to the extraction model, and the extraction
        prompt relies on exactly that marker to recognise an equation.
        """
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks

        merged = merge_mineru_text_blocks(
            [_mb("t1", "其中"), _mb("e1", "", block_type="equation", latex="E=mc^2")],
            target_tokens=800,
        )
        self.assertEqual(len(merged), 1)
        self.assertIn("EQUATION: E=mc^2", merged[0].text)

    def test_tables_and_images_still_break_the_run(self) -> None:
        """They are large on their own and carry their own 2048-dim visual vector, so retrieving them
        separately makes sense.

        The body text is made long enough (above the adoption floor), otherwise the following block would claim
        it as a title -- that is a different rule, see LeadInAdoptionTests. What is tested here is the break
        itself.
        """
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks

        long_a, long_b = _text_of_tokens(300), _text_of_tokens(300)
        for kind in ("table", "image", "chart"):
            with self.subTest(kind=kind):
                merged = merge_mineru_text_blocks(
                    [_mb("t1", long_a), _mb("x1", "阻断", block_type=kind), _mb("t2", long_b)],
                    target_tokens=800,
                )
                self.assertEqual([b.block_id for b in merged],
                                 ["mineru-merged-text-00001", "x1", "mineru-merged-text-00002"])

    def test_plain_text_input_is_untouched_by_the_render_switch(self) -> None:
        """Switching the join from block.text to text_for_block must be verbatim-equivalent for plain prose.

        Otherwise this change would also touch every document without equations -- once the ready text
        changes, GraphRAG's text_unit changes with it and the whole LLM cache is invalidated. kb_004 has not a
        single equation; its ready documents must not change by one byte.
        """
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks

        blocks = [_mb("t1", "第一段"), _mb("t2", "第二段"), _mb("t3", "第三段")]
        merged = merge_mineru_text_blocks(blocks, target_tokens=800)
        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0].text, "第一段\n\n第二段\n\n第三段")


class ShortBufferFoldTests(unittest.TestCase):
    """The remainder left by an overflow does not become its own block; it folds back into the previous
    merged block.

    750+100 crossing 800 flushes 750, and the remaining 100 would become a standalone chunk. Folding it back
    is enough, but **only when adjacent**: the rendered ready document is joined in block order, and folding
    back across a table would scramble the document order -- and that text is what the graph build reads.
    """

    def _run(self, blocks):
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks
        return merge_mineru_text_blocks(blocks, target_tokens=800)

    def test_the_leftover_folds_back(self) -> None:
        merged = self._run([
            _mb("a", _text_of_tokens(750)),
            _mb("b", _text_of_tokens(60)),          # 750+60 overflows -> flush a, b is stranded
            _mb("tbl", "表", block_type="table"),
        ])
        self.assertEqual(len(merged), 2)
        self.assertEqual(merged[0].metadata["merged_block_count"], 2)
        self.assertEqual(merged[0].metadata["source_block_ids"], ["a", "b"])

    def test_it_does_not_fold_across_a_table(self) -> None:
        """Short prose wedged between tables never folds back -- crossing a table scrambles the document
        order, and the rendered ready document is joined in exactly that block order.

        Where it goes instead is **forward**, merged into the next block (see LeadInAdoptionTests): position
        unchanged, order intact. This test pins just one thing: it does not appear before or inside tbl1.
        """
        merged = self._run([
            _mb("tbl1", "表一", block_type="table"),
            _mb("b", _text_of_tokens(60)),
            _mb("tbl2", "表二", block_type="table"),
        ])
        self.assertEqual([b.block_id for b in merged], ["tbl1", "tbl2"])
        self.assertIsNone(merged[0].title)          # did not fold back
        self.assertIn("书", merged[1].title or "")   # merged forward into tbl2

    def test_folding_never_invents_a_block_id(self) -> None:
        """chunk_uid is accounted by block_id. Inventing a new id when folding back would break graph
        provenance on the spot (chunk_provenance_links would count it as unmapped)."""
        merged = self._run([
            _mb("a", _text_of_tokens(750)),
            _mb("b", _text_of_tokens(60)),
            _mb("tbl", "表", block_type="table"),
        ])
        self.assertEqual(merged[0].block_id, "mineru-merged-text-00001")
        self.assertEqual(len({b.block_id for b in merged}), len(merged))

    def test_a_long_enough_buffer_still_stands_alone(self) -> None:
        merged = self._run([
            _mb("a", _text_of_tokens(750)),
            _mb("b", _text_of_tokens(300)),         # 300 >= floor(100)
            _mb("tbl", "表", block_type="table"),
        ])
        self.assertEqual(len(merged), 3)
        self.assertEqual([b.block_id for b in merged[:2]],
                         ["mineru-merged-text-00001", "mineru-merged-text-00002"])

    def test_a_title_at_the_break_goes_with_the_text_below_it(self) -> None:
        """When the buffer is cut because it is full, the titles at its tail are left to the next buffer. They used
        to go out with the previous block: the title landed at the end of the previous section's block and then
        vanished from every chunk during chunking (in production 42 of 116 PDFs, 326 lines)."""
        from kb_pipeline.chunking.chunker import blocks_to_chunks

        def section(name):
            return {"section_path": [name]}

        merged = self._run([
            _mb("a", _text_of_tokens(700), metadata=section("第1章 概述")),
            _mb("h", "第2章 部署说明", block_type="title", metadata=section("第2章 部署说明")),
            _mb("b", _text_of_tokens(300), metadata=section("第2章 部署说明")),     # 700 + title + 300 overflows: cut before the title
            _mb("h2", "2.1 环境", block_type="title", metadata=section("2.1 环境")),
        ])
        self.assertEqual([b.metadata["source_block_ids"] for b in merged], [["a"], ["h", "b", "h2"]])
        self.assertEqual(merged[0].metadata["heading_lines"], [])
        self.assertTrue(merged[1].text.startswith("第2章 部署说明\n\n书"))
        self.assertEqual(merged[1].metadata["section_path"], ["第2章 部署说明"])
        chunks = blocks_to_chunks(kb_id="k", file_key=1, content_version="v", parser_profile="p", blocks=merged,
                                  max_tokens=400, overlap_tokens=0)
        texts = [c.text for c in chunks]
        self.assertTrue(any(t.startswith("第2章 部署说明\n书") for t in texts), [t[:20] for t in texts])
        self.assertEqual(sum(t.count("2.1 环境") for t in texts), 1)               # the title at the end has no next block to go to, and is not lost either
        # The title exactly fills the buffer: it is left to the next block as well
        from kb_pipeline.utils import count_tokens

        merged = self._run([_mb("a", _text_of_tokens(800 - count_tokens("第2章 部署说明"))),
                            _mb("h", "第2章 部署说明", block_type="title"), _mb("b", _text_of_tokens(300))])
        self.assertEqual([b.metadata["source_block_ids"] for b in merged], [["a"], ["h", "b"]])
        # A table follows the title: the title still becomes the table's TITLE and is not carried forward
        merged = self._run([_mb("a", _text_of_tokens(798)), _mb("h", "第2章 部署说明", block_type="title"),
                            _mb("tbl", "| a | b |", block_type="table")])
        self.assertEqual([(b.block_id, b.title) for b in merged], [("mineru-merged-text-00001", None), ("tbl", "第2章 部署说明")])
        # A long run of titles (an outline, an index page marked as titles throughout): at most 3 are left to the
        # next block, and the buffer is still emitted by size
        outline = [_mb(f"h{i}", f"{i} 条目{i}", block_type="title") for i in range(1000)] + [_mb("p", _text_of_tokens(300))]
        merged = self._run(outline)
        self.assertGreater(len(merged), 5)
        self.assertLess(max(count_tokens(b.text) for b in merged), 1200)        # the blank lines between blocks cost tokens too, so a little over 800
        self.assertEqual([i for b in merged for i in b.metadata["source_block_ids"]], [b.block_id for b in outline])
        self.assertEqual(merged[-1].metadata["source_block_ids"][-2:], ["h999", "p"])


class DocxMergeTests(unittest.TestCase):
    """docx used to have no merge step at all.

    Every MinerU paragraph and every one-line subheading was its own block, so a 12-token subheading like the
    bold "3.2.2.1 ..." line in _fake_blocks became a standalone chunk and a standalone vector. Measured on
    mineru-3.4.4-docx-v1: a median chunk of 20 tokens, 70% under 50 tokens -- the most fragmented of all the
    parse paths (the PDF path has a merger, median 115).

    The block structure is identical to PDF, so the same merger is reused rather than writing another one.
    """

    def _fake_blocks(self):
        return [
            _mb("mineru-docx-title-00001", "**3.2.2.1 接入QQ**", block_type="title"),
            _mb("mineru-docx-text-00002", "先在管理后台开通渠道。"),
            _mb("mineru-docx-text-00003", "然后扫码绑定。"),
        ]

    def test_the_docx_path_actually_merges(self) -> None:
        import tempfile
        from kb_pipeline.parsers import docx_enhanced

        blocks = self._fake_blocks()
        original = docx_enhanced.mineru_docx_blocks
        docx_enhanced.mineru_docx_blocks = lambda **kw: list(blocks)
        try:
            with tempfile.TemporaryDirectory() as tmp:
                # Not a real zip -> media / chart extraction takes the BadZipFile branch and returns empty,
                # never touching the network
                fake = Path(tmp) / "x.docx"
                fake.write_bytes(b"not a zip")
                out = docx_enhanced.parse_docx_enhanced(
                    mineru_url="", vlm_base_url="", vlm_api_key="", vlm_model_id="",
                    vlm_concurrency=1, path=fake, cache_dir=Path(tmp),
                    merge_target_tokens=800,
                )
        finally:
            docx_enhanced.mineru_docx_blocks = original

        self.assertEqual(len(out), 1, "三个块应该并成一个")
        self.assertEqual(out[0].metadata["merged_block_count"], 3)
        self.assertIn("接入QQ", out[0].text)
        self.assertIn("扫码绑定", out[0].text)

    def test_merged_docx_ids_do_not_collide_with_the_pdf_ones(self) -> None:
        """Both paths share the merger, but the block_id must still reveal which path it came from.

        block_id is the accounting unit of chunk_uid; if the same prefix collided across the two document
        kinds, provenance would point at the wrong file.
        """
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks

        pdf = merge_mineru_text_blocks([_mb("t1", "正文")], target_tokens=800)
        docx = merge_mineru_text_blocks(
            [_mb("t1", "正文")], target_tokens=800, id_prefix="mineru-docx-merged-text")
        self.assertEqual(pdf[0].block_id, "mineru-merged-text-00001")
        self.assertEqual(docx[0].block_id, "mineru-docx-merged-text-00001")

    def test_the_merge_target_is_shared_with_pdf_not_reinvented(self) -> None:
        """Both paths use the same block_merge_tokens (2x max_tokens) instead of each configuring its own."""
        from kb_pipeline.models import KBSource

        source = KBSource(kb_id="kb_x", collection="kb_x", source_root="r",
                          source_type="local", max_tokens=400, overlap_tokens=80)
        self.assertEqual(source.block_merge_tokens, 800)
        call = _repo_file("app/kb_pipeline/pipeline/parse_job.py")
        docx_call = call.split("return parse_docx_enhanced(", 1)[1].split(")", 1)[0]
        self.assertIn("merge_target_tokens=source.block_merge_tokens", docx_call)


class OfficeRepeatedContentTests(unittest.TestCase):
    """docx / pptx used to de-duplicate by "type + first 160 characters". In most docx files the page number is
    the same value throughout, so the de-duplication covered the whole document, and paragraphs, tables and
    pictures in later sections identical to earlier ones were silently dropped (three FAQ questions each had a
    "Solution" line, and after indexing only the first question kept it); pptx kept only one of the short
    labels repeated on a slide. What the original repeats is content, and is not de-duplicated."""

    def test_docx_keeps_paragraphs_tables_and_pictures_that_repeat(self) -> None:
        import base64
        from unittest import mock

        from kb_pipeline.parsers import mineru_docx

        table = "<table><tr><td>参数</td><td>取值</td></tr><tr><td>仓库</td><td>GitHub</td></tr></table>"
        png = base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"0" * 32).decode()
        content = []
        for n in (1, 2, 3):
            content += [
                {"type": "text", "text": f"问题 {n}", "text_level": 1, "page_idx": 0},
                {"type": "text", "text": f"问题描述:第 {n} 种现象。", "page_idx": 0},
                {"type": "text", "text": "解决方案:重启服务后重新登录。", "page_idx": 0},
                {"type": "table", "table_body": table, "page_idx": 0},
                {"type": "image", "img_path": "images/flow.png", "page_idx": 0},
            ]
        payload = {"content_list": content, "images": {"images/flow.png": png}}
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(mineru_docx, "call_mineru_sync", return_value=payload):
            blocks = mineru_docx.mineru_docx_blocks(mineru_url="http://x", path=Path(tmp) / "faq.docx", cache_dir=Path(tmp))
        self.assertEqual(len(blocks), 15)
        self.assertEqual(len({b.block_id for b in blocks}), 15)
        for kind, section in (("text", "问题 2"), ("table", "问题 3"), ("image", "问题 3")):
            self.assertTrue([b for b in blocks if b.block_type == kind and b.metadata["section_path"] == [section]
                             and (kind != "text" or b.text == "解决方案:重启服务后重新登录。")], (kind, section))
        self.assertEqual(sum(b.text == "解决方案:重启服务后重新登录。" for b in blocks), 3)
        self.assertEqual(len({b.visual_ref for b in blocks if b.block_type == "image"}), 1)      # the same picture, referenced in three places

    def test_pptx_keeps_short_labels_that_repeat_on_a_slide(self) -> None:
        from unittest import mock

        from kb_pipeline.parsers import mineru_pptx

        content = [{"type": "text", "text": "功能对比", "page_idx": 0}]
        for name in ("单点登录", "审计日志", "离线包", "国密算法"):
            content += [{"type": "text", "text": name, "page_idx": 0}, {"type": "text", "text": "支持", "page_idx": 0}]
        content += [{"type": "text", "text": "功能对比", "page_idx": 1}, {"type": "text", "text": "支持", "page_idx": 1}]
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(mineru_pptx, "call_mineru_sync", return_value={"content_list": content}):
            blocks = mineru_pptx.mineru_pptx_blocks(mineru_url="http://x", path=Path(tmp) / "deck.pptx", cache_dir=Path(tmp))
            slides = mineru_pptx.aggregate_mineru_slide_blocks(blocks, Path(tmp) / "deck.pptx")
        self.assertEqual([b.text for b in blocks if b.slide_idx == 1],
                         ["功能对比", "单点登录", "支持", "审计日志", "支持", "离线包", "支持", "国密算法", "支持"])
        self.assertEqual(slides[0].text.count("支持"), 4)                                        # the label under every card is kept
        self.assertEqual(len({b.block_id for b in blocks}), len(blocks))
        self.assertEqual(parser_profile_for_path(Path("a.pptx")), "pptx-mineru-slide-vlm-v4")


class DocxChartTests(unittest.TestCase):
    """The text of a native docx chart used to be every value in the XML collected together and de-duplicated by
    value: repeated data points were deleted, every later value shifted, and formula references were written
    into the body as well; chart blocks were always appended at the end of the document, without a section."""

    NS = ('xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')

    @staticmethod
    def _cache(kind: str, values, number_format: str = "General") -> str:
        pts = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in enumerate(values) if v is not None)
        code = f"<c:formatCode>{number_format}</c:formatCode>" if kind == "num" else ""
        return (f"<c:{kind}Ref><c:f>Sheet1!$B$2:$B$5</c:f><c:{kind}Cache>{code}"
                f'<c:ptCount val="{len(values)}"/>{pts}</c:{kind}Cache></c:{kind}Ref>')

    def _series(self, name: str, cats, vals, number_format: str = "General") -> str:
        return (f"<c:ser><c:tx>{self._cache('str', [name])}</c:tx><c:cat>{self._cache('str', cats)}</c:cat>"
                f"<c:val>{self._cache('num', vals, number_format)}</c:val></c:ser>")

    def _chart(self, title: str, *series: str) -> bytes:
        return (f"<c:chartSpace {self.NS}><c:chart><c:title><c:tx><c:rich><a:p><a:r><a:t>{title}</a:t></a:r></a:p></c:rich></c:tx></c:title>"
                f"<c:plotArea><c:barChart>{''.join(series)}</c:barChart>"
                "<c:catAx><c:title><c:tx><c:rich><a:p><a:r><a:t>季度</a:t></a:r></a:p></c:rich></c:tx></c:title></c:catAx>"
                "</c:plotArea></c:chart></c:chartSpace>").encode("utf-8")

    def test_every_point_keeps_its_series_and_category(self) -> None:
        from kb_pipeline.parsers.docx_enhanced import chart_xml_table

        quarters = ["Q1", "Q2", "Q3", "Q4"]
        title, axes, table = chart_xml_table(self._chart(
            "销售额", self._series("2023", quarters, [10, 20, 10, 30]), self._series("2024", quarters, [10, None, 0.1 + 0.2, 30]),
            self._series("完成率", quarters, [0.06, 0.916666666666667, 1, 0.5], "0%")))
        self.assertEqual((title, axes), ("销售额", ["季度"]))
        self.assertEqual(table.splitlines(), [
            "| Series | Category | Value |", "| --- | --- | --- |",
            "| 2023 | Q1 | 10 |", "| 2023 | Q2 | 20 |", "| 2023 | Q3 | 10 |", "| 2023 | Q4 | 30 |",     # the 10 of Q3 is still there
            "| 2024 | Q1 | 10 |", "| 2024 | Q2 |  |", "| 2024 | Q3 | 0.3 |", "| 2024 | Q4 | 30 |",       # values equal to the previous series are kept; an empty point keeps its place
            "| 完成率 | Q1 | 6% |", "| 完成率 | Q2 | 92% |", "| 完成率 | Q3 | 100% |", "| 完成率 | Q4 | 50% |",
        ])
        self.assertNotIn("Sheet1!", table)                                                       # formula references are not content
        # Multi-level categories (the inner level first in the file) are joined from the outer to the inner level,
        # an outer label written only on the first point of its group and carried forward; a date axis caches
        # serial numbers. However large the indexes, only the points that exist count
        far = 3_000_000_000
        levels = ("<c:multiLvlStrRef><c:f>Sheet1!$A$2:$B$5</c:f><c:multiLvlStrCache>"
                  + "<c:lvl>" + "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in ((0, "Q1"), (1, "Q2"), (2, "Q1"), (far, "Q2"))) + "</c:lvl>"
                  + '<c:lvl><c:pt idx="0"><c:v>2023</c:v></c:pt><c:pt idx="2"><c:v>2024</c:v></c:pt></c:lvl>'
                  + "</c:multiLvlStrCache></c:multiLvlStrRef>")
        values = "".join(f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>' for i, v in ((0, 1), (1, 2), (2, 3), (far, 4)))
        nested = (f"<c:ser><c:tx>{self._cache('str', ['销量'])}</c:tx><c:cat>{levels}</c:cat>"
                  f"<c:val><c:numRef><c:numCache><c:formatCode>General</c:formatCode>{values}</c:numCache></c:numRef></c:val></c:ser>")
        dated = (f"<c:ser><c:tx>{self._cache('str', ['日活'])}</c:tx><c:cat>{self._cache('num', [45292, 45323], 'yyyy/m/d')}</c:cat>"
                 f"<c:val>{self._cache('num', [120, 135])}</c:val></c:ser>")
        self.assertEqual(chart_xml_table(self._chart("走势", nested, dated))[2].splitlines()[2:], [
            "| 销量 | 2023/Q1 | 1 |", "| 销量 | 2023/Q2 | 2 |", "| 销量 | 2024/Q1 | 3 |", "| 销量 | 2024/Q2 | 4 |",
            "| 日活 | 2024-01-01 | 120 |", "| 日活 | 2024-02-01 | 135 |",
        ])
        # A chart with formula references only and no cached values: there is no data to write
        bare = f"<c:chartSpace {self.NS}><c:chart><c:plotArea><c:barChart><c:ser><c:val><c:numRef><c:f>S!B2:B5</c:f></c:numRef></c:val></c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>"
        self.assertEqual(chart_xml_table(bare.encode("utf-8")), ("", [], ""))
        # A chart with the other structure (chartEx): no pairing, the cached values are listed in their original
        # order, and repeated values are kept
        other = ('<cx:chartSpace xmlns:cx="http://schemas.microsoft.com/office/drawing/2014/chartex"><cx:chartData><cx:data id="0">'
                 '<cx:strDim type="cat"><cx:f>Sheet1!$A$2:$A$4</cx:f><cx:lvl ptCount="3"><cx:pt idx="0">期初</cx:pt><cx:pt idx="1">增加</cx:pt><cx:pt idx="2">期末</cx:pt></cx:lvl></cx:strDim>'
                 '<cx:numDim type="val"><cx:f>Sheet1!$B$2:$B$4</cx:f><cx:lvl ptCount="3"><cx:pt idx="0">10</cx:pt><cx:pt idx="1">10</cx:pt><cx:pt idx="2">20</cx:pt></cx:lvl></cx:numDim>'
                 "</cx:data></cx:chartData></cx:chartSpace>")
        self.assertEqual(chart_xml_table(other.encode("utf-8"))[2].splitlines(), ["期初", "增加", "期末", "10", "10", "20"])

    def test_charts_go_back_to_where_they_sit_in_the_document(self) -> None:
        import zipfile
        from unittest import mock

        from kb_pipeline.chunking.chunker import blocks_to_chunks
        from kb_pipeline.parsers import docx_enhanced

        w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        r = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        c = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"'

        def para(text: str = "", chart: str = "") -> str:
            run = f"<w:r><w:t>{text}</w:t></w:r>" if text else ""
            ref = f'<w:r><w:drawing><c:chart {c} r:id="{chart}"/></w:drawing></w:r>' if chart else ""
            return f"<w:p>{run}{ref}</w:p>"

        document = (f"<w:document {w} {r}><w:body>" + para("一、华东区") + para("华东区各季度销售额如下:") + para(chart="rId7")
                    + para("二、华南区") + para("华南区各季度销售额如下:") + para(chart="rId5") + para("三、结论") + "</w:body></w:document>")
        rels = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/>'
                '<Relationship Id="rId7" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="/word/charts/chart2.xml"/>'
                "</Relationships>")
        quarters = ["Q1", "Q2", "Q3", "Q4"]
        blocks = [
            _mb("t1", "**一、华东区**", block_type="title", metadata={"section_path": ["一、华东区"]}),
            _mb("p1", "华东区各季度销售额如下:", metadata={"section_path": ["一、华东区"]}),
            _mb("t2", "**二、华南区**", block_type="title", metadata={"section_path": ["二、华南区"]}),
            _mb("p2", "华南区各季度销售额如下:", metadata={"section_path": ["二、华南区"]}),
            _mb("t3", "**三、结论**", block_type="title", metadata={"section_path": ["三、结论"]}),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "报告.docx"
            with zipfile.ZipFile(path, "w") as zf:
                zf.writestr("word/document.xml", document)
                zf.writestr("word/_rels/document.xml.rels", rels)
                zf.writestr("word/charts/chart1.xml", self._chart("华南区销售额", self._series("2024", quarters, [7, 7, 8, 9])))
                zf.writestr("word/charts/chart2.xml", self._chart("华东区销售额", self._series("2024", quarters, [10, 20, 10, 30])))
            with mock.patch.object(docx_enhanced, "mineru_docx_blocks", lambda **kw: list(blocks)), \
                    mock.patch.object(docx_enhanced, "merge_mineru_text_blocks", lambda found, **kw: found):
                out = docx_enhanced.parse_docx_enhanced(mineru_url="", vlm_base_url="", vlm_api_key="", vlm_model_id="",
                                                        vlm_concurrency=1, path=path, cache_dir=Path(tmp))
        self.assertEqual([b.block_id for b in out], ["t1", "p1", "docx-chart-0001", "t2", "p2", "docx-chart-0002", "t3"])
        east, south = out[2], out[5]
        self.assertEqual((east.title, east.metadata["section_path"]), ("华东区销售额", ["一、华东区"]))
        self.assertEqual((south.title, south.metadata["section_path"]), ("华南区销售额", ["二、华南区"]))
        self.assertEqual((east.block_type, east.caption), ("table", "Chart data; axes: 季度"))
        chunks = blocks_to_chunks(kb_id="kb", file_key=1, content_version="v", parser_profile="p", blocks=out,
                                  max_tokens=400, overlap_tokens=0)
        text = next(c.text for c in chunks if c.block.block_id == "docx-chart-0001")
        self.assertIn("TITLE: 华东区销售额\nCAPTION: Chart data; axes: 季度\n| Series | Category | Value |", text)
        self.assertIn("| 2024 | Q3 | 10 |", text)
        self.assertEqual(parser_profile_for_path(Path("a.docx")), "docx-mineru-ooxml-vlm-v10")


class LeadInAdoptionTests(unittest.TestCase):
    """A short heading wedged between tables / images is merged into the block that follows it.

    This is the last kind of fragment in the chain, and the loss cuts both ways: the heading alone becomes
    a 4-token chunk that cannot be retrieved (BM25 favours short documents, so it actually surfaces too
    easily), while the table it describes has nothing in the store saying what it is. Measured on kb_004,
    67% of the figures / tables were in the state "no caption of its own, heading stranded next door".

    The criterion is about **content**, not the knowledge base's identity: if the target has its own
    caption / title, it is not taken over. MinerU extracts captions well on datasheets and books (kb_003 92%,
    kb_005 69% carry their own), so the rule steps aside there; it only applies to product manuals laid out
    with the heading sitting above the figure.
    """

    def _run(self, blocks):
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks
        return merge_mineru_text_blocks(blocks, target_tokens=800)

    def test_a_stranded_heading_is_adopted_by_the_next_block(self) -> None:
        merged = self._run([
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("h", "直流电气特性"),
            _mb("tbl1", "| a | b |", block_type="table"),
        ])
        self.assertEqual([b.block_id for b in merged], ["tbl0", "tbl1"])
        self.assertEqual(merged[1].title, "直流电气特性")
        self.assertTrue(merged[1].metadata.get("adopted_lead_in"))
        self.assertEqual(merged[1].metadata.get("adopted_block_ids"), ["h"])

    def test_a_block_that_already_has_a_caption_declines(self) -> None:
        """Leave alone what is already right -- which is exactly why it needs no per-KB policy."""
        merged = self._run([
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("h", "直流电气特性"),
            _mb("tbl1", "| a | b |", block_type="table", caption="表 3-1 直流电气特性"),
        ])
        self.assertEqual([b.block_id for b in merged],
                         ["tbl0", "mineru-merged-text-00001", "tbl1"])
        self.assertEqual(merged[2].caption, "表 3-1 直流电气特性")

    def test_it_goes_into_title_not_text(self) -> None:
        """A table block's text is its table_markdown. Stuffing the heading into text would make
        text_for_block see text != table_markdown and emit the whole table a second time."""
        from kb_pipeline.chunking.chunker import text_for_block

        merged = self._run([
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("h", "直流电气特性"),
            _mb("tbl1", "| a | b |", block_type="table", table_markdown="| a | b |"),
        ])
        rendered = text_for_block(merged[1])
        self.assertTrue(rendered.startswith("TITLE: 直流电气特性"))
        self.assertEqual(rendered.count("| a | b |"), 1)

    def test_long_text_is_not_treated_as_a_heading(self) -> None:
        merged = self._run([
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("body", _text_of_tokens(300)),
            _mb("tbl1", "| a | b |", block_type="table"),
        ])
        self.assertEqual([b.block_id for b in merged],
                         ["tbl0", "mineru-merged-text-00001", "tbl1"])
        self.assertIsNone(merged[2].title)

    def test_folding_back_wins_over_adoption(self) -> None:
        """A buffer right behind a merged block that was just emitted is the tail of prose cut off by
        target_tokens, not this table's heading. In that case it belongs back with the preceding text."""
        merged = self._run([
            _mb("a", _text_of_tokens(750)),
            _mb("b", _text_of_tokens(60)),      # 750+60 overflows -> flush a, b is stranded
            _mb("tbl", "| a | b |", block_type="table"),
        ])
        self.assertEqual([b.block_id for b in merged], ["mineru-merged-text-00001", "tbl"])
        self.assertEqual(merged[0].metadata["source_block_ids"], ["a", "b"])
        self.assertIsNone(merged[1].title)

    def test_the_document_order_never_changes(self) -> None:
        """Adoption means "carried by a different block", not "moved to a different position". The heading
        came before the table and still comes before it after the merge -- chunks are joined in block order,
        and once the order slips the extraction unit reads a different document."""
        blocks = [
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("h", "直流电气特性"),
            _mb("tbl1", "| a | b |", block_type="table"),
        ]
        text = "\n".join(f"{b.title or ''}\n{b.text}" for b in self._run(blocks))
        self.assertLess(text.index("直流电气特性"), text.index("| a | b |"))


class MergedCellFillTests(unittest.TestCase):
    """Merged cells are filled down vertically, not across.

    Filling across is pure duplication: the same value copied N times within one row renders as
    note|note|note|note, and after chunking that is dozens of verbatim-identical chunks -- one vector stored
    dozens of times, surfacing in clusters at retrieval. Measured on table-native-v1, 9.2% of the chunks came
    from this, 66% for the "filling instructions" sheet.

    Vertical fill must stay: a category label such as "WPS AI" spanning 54 rows is exactly what gives every
    row chunk its own category; without it every row but the first has no idea where it belongs.
    """

    def _sheet(self, merges, cells, dims):
        """A minimal openpyxl stand-in: merged_cells.ranges / iter_rows / cell / max_row / max_column are all
        it takes to drive the merge-fill branch of read_xlsx_rows."""
        class Rng:
            def __init__(self, b): self.bounds = b
        class Cell:
            def __init__(self, r, c, v): self.row, self.column, self.value = r, c, v
        class WS:
            title = "S"
            max_row, max_column = dims
            merged_cells = type("M", (), {"ranges": [Rng(b) for b in merges]})()
            def cell(self, r, c): return Cell(r, c, cells.get((r, c)))
            def iter_rows(self):
                for r in range(1, dims[0] + 1):
                    yield [Cell(r, c, cells.get((r, c))) for c in range(1, dims[1] + 1)]
        return WS()

    def _fill(self, ws):
        """Replays the fill logic of read_xlsx_rows: calling the real function is too heavy (it needs a zip),
        so this steps back -- it lifts that loop out of the source and runs it, guaranteeing the same code is
        under test."""
        from kb_pipeline.parsers import native_table

        src = _repo_file("app/kb_pipeline/parsers/native_table.py")
        body = src.split("for merged in ws.merged_cells.ranges:", 1)[1]
        body = body.split("rows: list[tuple[int, list[str]]] = []", 1)[0]
        code = "for merged in ws.merged_cells.ranges:\n" + body
        code = "\n".join(line[12:] if line.startswith(" " * 12) else line
                         for line in code.splitlines())
        ns = {"ws": ws, "merged_values": {}, "used_rows": ws.max_row,
              "used_cols": ws.max_column, "normalize_cell": native_table.normalize_cell}
        exec(compile(code, "fill.py", "exec"), ns)
        return ns["merged_values"]

    def test_a_block_merge_is_not_copied_across_columns(self) -> None:
        # A 7-row x 4-column note block, like the "filling instructions" sheet
        ws = self._sheet([(1, 2, 4, 8)], {(2, 1): "填写说明：一、二、三"}, (24, 4))
        filled = self._fill(ws)
        self.assertEqual({c for _, c in filled}, {1}, "只该落在锚点列")
        self.assertEqual(sorted({r for r, _ in filled}), list(range(2, 9)), "每一行都该有")

    def test_a_vertical_merge_still_labels_every_row(self) -> None:
        # The "WPS AI" category label spanning 54 rows -- this branch must be kept as is
        ws = self._sheet([(1, 5, 1, 58)], {(5, 1): "WPS AI"}, (110, 8))
        filled = self._fill(ws)
        self.assertEqual(len(filled), 54)
        self.assertTrue(all(v == "WPS AI" for v in filled.values()))
        self.assertEqual({c for _, c in filled}, {1})

    def test_an_empty_merge_fills_nothing(self) -> None:
        ws = self._sheet([(1, 2, 4, 8)], {}, (24, 4))
        self.assertEqual(self._fill(ws), {})


class HtmlTableToMarkdownTests(unittest.TestCase):
    """For tables MinerU gives only the table_body key, containing <table><tr><td>..., with no markdown
    alternative. That markup goes into the chunk text verbatim: measured, tags made up a median 57%~61% of
    the body characters, 312,105 tags in kb_005 alone. The tokenizer compresses the repeated </td><td>
    fairly well, so the real waste in tokens is 33% -- 3,374 real tables, 1,391,102 -> 928,717 tokens.

    The waste is not just budget: one 399-token chunk was nothing but <td></td>, and its vector represents
    "this is a table" rather than what the table contains.
    """

    def _md(self, html):
        from kb_pipeline.parsers.common import html_table_to_markdown
        return html_table_to_markdown(html)

    def _tbl(self, *rows):
        lt, gt = chr(60), chr(62)
        body = "".join("%str%s%s%s/tr%s" % (lt, gt, r, lt, gt) for r in rows)
        return "%stable%s%s%s/table%s" % (lt, gt, body, lt, gt)

    def _cell(self, text, **attrs):
        lt, gt = chr(60), chr(62)
        a = "".join(' %s="%s"' % kv for kv in attrs.items())
        return "%std%s%s%s%s/td%s" % (lt, a, gt, text, lt, gt)

    def test_a_plain_table_becomes_pipe_rows(self) -> None:
        md = self._md(self._tbl(self._cell("a") + self._cell("b"),
                                self._cell("1") + self._cell("2")))
        self.assertEqual(md.splitlines(),
                         ["| a | b |", "| --- | --- |", "| 1 | 2 |"])

    def test_rowspan_repeats_instead_of_leaving_a_hole(self) -> None:
        """markdown has no notion of spans. Leaving a hole shifts every later column, and a misaligned table
        is harder to read than a repetitive one -- and more likely to make the model pair values with the
        wrong row."""
        md = self._md(self._tbl(self._cell("a") + self._cell("b"),
                                self._cell("x", rowspan="2") + self._cell("1"),
                                self._cell("2")))
        self.assertEqual(md.splitlines()[-2:], ["| x | 1 |", "| x | 2 |"])

    def test_colspan_fills_the_row_width(self) -> None:
        md = self._md(self._tbl(self._cell("标题", colspan="3"),
                                self._cell("1") + self._cell("2") + self._cell("3")))
        self.assertEqual(md.splitlines()[0], "| 标题 | 标题 | 标题 |")

    def test_latex_with_a_less_than_sign_survives(self) -> None:
        """For a cell like $S_T<K$ in a maths table, "is the character after < a letter" cannot decide -- K
        is a letter. HTMLParser would take <K$... as the start of a tag name and the whole cell would vanish
        silently. So the check uses a whitelist of known tags, and any < outside it is a plain character.
        """
        md = self._md(self._tbl(self._cell("$S_T<K$") + self._cell("b")))
        self.assertIn("$S_T<K$", md)
        self.assertEqual(len(md.splitlines()[0].split("|")) - 2, 2)

    def test_a_pipe_inside_a_cell_is_escaped(self) -> None:
        md = self._md(self._tbl(self._cell("a|b") + self._cell("c")))
        self.assertIn("a\\|b", md)
        # After escaping there are still only two real column separators -- a plain split("|") would count
        # the \| as well
        import re as _re
        self.assertEqual(len(_re.split(r"(?<!\\)\|", md.splitlines()[0])) - 2, 2)

    def test_truncated_html_still_yields_what_was_read(self) -> None:
        lt, gt = chr(60), chr(62)
        md = self._md("%stable%s%str%s%std%s只有半张" % (lt, gt, lt, gt, lt, gt))
        self.assertIn("只有半张", md)

    def test_non_tables_and_empty_tables_return_none(self) -> None:
        lt, gt = chr(60), chr(62)
        for value in ("没有表格", "", "%stable%s%s/table%s" % (lt, gt, lt, gt)):
            with self.subTest(value=value[:12]):
                self.assertIsNone(self._md(value))

    def test_item_text_converts_table_body(self) -> None:
        """The wiring itself must be pinned too: a converter written but never hooked up is as good as none."""
        from kb_pipeline.parsers.common import item_text

        text = item_text({"table_body": self._tbl(self._cell("a") + self._cell("b"))})
        self.assertIn("| a | b |", text)
        self.assertNotIn("td" + chr(62), text)

    def test_an_unconvertible_body_is_kept_verbatim(self) -> None:
        """Better to keep the markup than to lose the whole content over one unrecognisable table."""
        from kb_pipeline.parsers.common import item_text

        weird = chr(60) + "table" + chr(62) + chr(60) + "/table" + chr(62)
        self.assertEqual(item_text({"table_body": weird}), weird)


class CodeAndFrontmatterParsingTests(unittest.TestCase):
    """Code repository knowledge bases (2026-09-04): markdown frontmatter becomes fields + its own block; Python
    is chunked by symbol without loss."""

    def test_frontmatter_becomes_fields_and_its_own_block(self) -> None:
        from kb_pipeline.parsers.common import markdown_to_blocks, parse_frontmatter

        md = "---\nname: 示例365-wiki\ndescription: >-\n  云端知识库编译与检索。\n  覆盖构建、搜索。\ntags: a,b\n---\n\n# 用法\n\n正文。\n"
        fields, body = parse_frontmatter(md)
        self.assertEqual(fields["name"], "示例365-wiki")
        self.assertEqual(fields["description"], "云端知识库编译与检索。 覆盖构建、搜索。")
        self.assertEqual(fields["tags"], "a,b")
        self.assertTrue(body.lstrip().startswith("# 用法"))
        blocks = markdown_to_blocks(md, parser="native", parser_profile="md-sections-v2", doc_type="md")
        self.assertEqual(blocks[0].block_id, "md-fm")
        self.assertEqual(blocks[0].metadata["frontmatter"]["name"], "示例365-wiki")
        self.assertIn("description: 云端知识库编译与检索", blocks[0].text)
        self.assertEqual([b.title for b in blocks[1:]], ["用法"])
        self.assertEqual(parse_frontmatter("# no frontmatter\n---\n"), ({}, "# no frontmatter\n---\n"))

    def test_python_symbol_blocks_are_lossless_and_carry_metadata(self) -> None:
        from kb_pipeline.parsers.code_python import python_symbol_blocks

        src = (
            "import os\nfrom lib.util import helper\n\nLIMIT = 3\n\n\n"
            "def run(x: int, *, flag=False) -> int:\n    \"\"\"Run it.\"\"\"\n    return helper(x) + os.getpid()\n\n\n"
            "@dataclass\nclass Engine(Base):\n    \"\"\"An engine.\"\"\"\n    name: str = \"e\"\n\n    def start(self):\n        return self.check()\n\n    CONST = 1\n\n    def check(self):\n        return run(1)\n\n\nif __name__ == \"__main__\":\n    run(2)\n"
        )
        blocks = python_symbol_blocks(src)
        kinds = [b.metadata["symbol"]["kind"] for b in blocks]
        self.assertEqual(kinds, ["module", "function", "class", "method", "method"])
        self.assertEqual([b.title for b in blocks], ["(module)", "run", "Engine", "Engine.start", "Engine.check"])
        # Lossless: every line lands in exactly one block (the module block collects all lines outside symbols)
        joined = sorted(l for b in blocks for l in b.text.splitlines() if l.strip())
        self.assertEqual(joined, sorted(l for l in src.splitlines() if l.strip()))
        run = blocks[1].metadata["symbol"]
        self.assertEqual(run["signature"], "def run(x: int, *, flag = False) -> int")
        self.assertEqual(run["docstring"], "Run it.")
        self.assertEqual(run["calls"], ["helper", "os.getpid"])
        cls = blocks[2].metadata["symbol"]
        self.assertEqual((cls["bases"], cls["methods"], cls["fields"]), (["Base"], ["start", "check"], ["name", "CONST"]))
        self.assertEqual(blocks[3].metadata["symbol"]["calls"], ["self.check"])
        self.assertEqual(blocks[3].metadata["section_path"], ["Engine"])
        self.assertIn("CONST = 1", blocks[3].text)          # class-level statements between methods go with the preceding method
        mod = blocks[0].metadata["symbol"]
        self.assertEqual(mod["constants"], ["LIMIT"])
        self.assertEqual(mod["imports"][1], {"module": "lib.util", "name": "helper", "asname": None, "level": 0})
        self.assertEqual(mod["calls"], ["run"])
        self.assertTrue(all(b.block_id.startswith("py-") for b in blocks))

    def test_python_blocks_split_lines_the_way_the_parser_counts_them(self) -> None:
        """Form feeds (Emacs-style page breaks) and a U+2028 inside a string are not line breaks. Lines used to be
        split with str.splitlines, so a single one in the file shifted the text of every later function block by
        a line, and function bodies landed in the module block."""
        from kb_pipeline.parsers.code_python import python_symbol_blocks

        src = "import os\n\x0c\ndef a():\n    s = 'x y'\n    return 1\n\n\ndef b():\n    return 2\n"
        for text in (src, src.replace("\n", "\r\n")):
            blocks = {b.title: b.text for b in python_symbol_blocks(text)}
            self.assertEqual(blocks["a"], "def a():\n    s = 'x y'\n    return 1")
            self.assertEqual(blocks["b"], "def b():\n    return 2")
            self.assertEqual(blocks["(module)"], "import os\n\x0c")


class MultiLanguageSymbolTests(unittest.TestCase):
    """tree-sitter multi-language symbol chunking: a small sample per language, checking that functions /
    classes / methods / imports / calls are extracted. Skipped where tree-sitter is not installed (the DGX
    venv has it)."""

    def setUp(self) -> None:
        try:
            import tree_sitter_language_pack  # noqa: F401
        except ImportError:
            self.skipTest("tree-sitter-language-pack not installed")

    def _syms(self, language: str, src: str):
        from kb_pipeline.parsers.code_symbols import symbol_blocks

        blocks = symbol_blocks(src, language)
        self.assertIsNotNone(blocks)
        syms = {b.metadata["symbol"]["qualname"]: b.metadata["symbol"] for b in blocks if b.metadata["symbol"].get("qualname")}
        module = next((b.metadata["symbol"] for b in blocks if b.metadata["symbol"].get("kind") == "module"), None) or blocks[0].metadata.get("module") or {}
        # Lossless: every (non-empty) line lands in exactly one block
        got = sorted(l for b in blocks for l in b.text.splitlines() if l.strip())
        self.assertEqual(got, sorted(l for l in src.splitlines() if l.strip()), language)
        return syms, module

    def test_blocks_split_lines_the_way_the_syntax_tree_counts_them(self) -> None:
        """A U+2028 inside a string, old Mac \\r line endings: when lines are split differently from the way the syntax
        tree numbers them, every function block takes its text from somewhere else."""
        from kb_pipeline.parsers.code_symbols import symbol_blocks

        src = 'const s = "a\u2028b";\nfunction a() {\n  return 1;\n}\nfunction b() {\n  return 2;\n}\n'
        for text in (src, src.replace("\n", "\r\n"), src.replace("\n", "\r")):
            blocks = {b.title: b.text for b in symbol_blocks(text, "javascript")}
            self.assertEqual(blocks["a"], "function a() {\n  return 1;\n}")
            self.assertEqual(blocks["b"], "function b() {\n  return 2;\n}")
            self.assertEqual(blocks["(module)"], 'const s = "a\u2028b";')

    def test_decorators_belong_to_the_definition_below_them(self) -> None:
        """Decorators on TypeScript class members and Rust attributes are siblings in front of the definition in the
        syntax tree: they used to end up in the previous method's block (the first method's in the class head), so
        route / test markers were separated from the methods they decorate."""
        from kb_pipeline.parsers.code_symbols import symbol_blocks

        syms, _ = self._syms("typescript", '''class Svc {
  @Get('/a')
  first() { return 1; }

  /** second route */
  @Post('/b')
  @Auth()
  second() { return 2; }
}
''')
        self.assertEqual((syms["Svc.first"]["lineno"], syms["Svc.second"]["lineno"]), (2, 6))
        self.assertEqual(syms["Svc.second"]["docstring"], "second route")
        blocks = {b.title: b.text for b in symbol_blocks("class Svc {\n  @Get('/a')\n  first() {}\n  @Post('/b')\n  second() {}\n}\n", "typescript")}
        self.assertEqual(blocks, {"Svc": "class Svc {", "Svc.first": "  @Get('/a')\n  first() {}",
                                  "Svc.second": "  @Post('/b')\n  second() {}\n}"})
        rust = "/// A thing\n#[derive(Debug)]\nstruct Foo { a: i32 }\n\nimpl Foo {\n    #[inline]\n    fn first(&self) -> i32 { 1 }\n\n    #[cfg(test)]\n    fn second(&self) -> i32 { 2 }\n}\n\n#[test]\nfn top() {}\n"
        blocks = [(b.title, b.text) for b in symbol_blocks(rust, "rust")]
        self.assertEqual(blocks, [("(module)", "/// A thing"), ("Foo", "#[derive(Debug)]\nstruct Foo { a: i32 }"), ("Foo", "impl Foo {"),
                                  ("Foo.first", "    #[inline]\n    fn first(&self) -> i32 { 1 }\n"),
                                  ("Foo.second", "    #[cfg(test)]\n    fn second(&self) -> i32 { 2 }\n}"),
                                  ("top", "#[test]\nfn top() {}")])

    def test_go_grouped_types_take_their_own_lines(self) -> None:
        """type ( A struct...; B struct... ): each type used to take the first and last line of the whole group, so the
        same text went into the blocks of both A and B."""
        from kb_pipeline.parsers.code_symbols import symbol_blocks

        src = "package x\n\ntype (\n\t// A doc\n\tA struct {\n\t\tX int\n\t}\n\tB interface {\n\t\tRun()\n\t}\n)\n\n// C doc\ntype C struct { Z int }\n"
        syms, _ = self._syms("go", src)                                          # lossless: every line lands in exactly one block
        blocks = {b.title: b.text for b in symbol_blocks(src, "go")}
        self.assertEqual(blocks["A"], "\tA struct {\n\t\tX int\n\t}")
        self.assertEqual(blocks["B"], "\tB interface {\n\t\tRun()\n\t}")
        self.assertEqual(blocks["C"], "type C struct { Z int }")
        self.assertEqual([(syms[n]["lineno"], syms[n]["end_lineno"], syms[n]["signature"], syms[n]["docstring"]) for n in "ABC"],
                         [(5, 7, "A struct", "A doc"), (8, 10, "B interface", ""), (14, 14, "type C struct", "C doc")])
        syms, _ = self._syms("go", "package x\n\n// Shapes used by the renderer\ntype (\n\tA struct {\n\t\tX int\n\t}\n\tB struct {\n\t\tY int\n\t}\n)\n")
        self.assertEqual([syms[n]["docstring"] for n in "AB"], ["Shapes used by the renderer"] * 2)   # a type without a comment of its own takes the group's

    def test_two_methods_on_one_line_do_not_break_the_file(self) -> None:
        syms, _ = self._syms("javascript", "class A {\n  a() { return 1; } b() { return 2; }\n  c() { return 3; }\n}\n")
        self.assertEqual(syms["A"]["methods"], ["a", "b", "c"])
        self.assertIn("A.c", syms)

    def test_wide_literals_do_not_take_quadratic_time(self) -> None:
        """In a bundled / generated file one statement is an array of tens of thousands of items: callee names used
        to be dequeued one at a time from a list used as a queue and checked for duplicates against an unbounded
        list, over twenty seconds for 1 MB. Collection now stops at the cap; the result is unchanged."""
        import time as time_module

        from kb_pipeline.parsers.code_symbols import MAX_CALLS, symbol_blocks

        big = "const data = [" + ",".join(f"f{i}({i})" for i in range(60000)) + "];\n"
        started = time_module.time()
        blocks = symbol_blocks(big, "javascript")
        self.assertLess(time_module.time() - started, 5.0)
        self.assertEqual(blocks[0].metadata["symbol"]["calls"], [f"f{i}" for i in range(MAX_CALLS)])

    def test_javascript_and_typescript(self) -> None:
        syms, mod = self._syms("javascript", '''import fs from "fs";
import { helper } from "./lib/util.js";
const LIMIT = 3;
/** Run it. */
export function run(x) { return helper(x) + fs.readFileSync(x); }
export const arrow = (a) => run(a);
class Engine extends Base {
  start() { return this.check(); }
  check() { return run(1); }
}
''')
        self.assertEqual(set(syms), {"run", "arrow", "Engine", "Engine.start", "Engine.check"})
        self.assertEqual((syms["run"]["kind"], syms["run"]["docstring"], syms["run"]["calls"]), ("function", "Run it.", ["helper", "fs.readFileSync"]))
        self.assertEqual(syms["Engine"]["bases"], ["Base"])
        self.assertEqual(syms["Engine.start"]["calls"], ["self.check"])
        self.assertEqual([i["module"] for i in mod["imports"]], ["fs", "./lib/util.js"])
        self.assertEqual(mod["constants"], ["LIMIT"])
        syms, mod = self._syms("typescript", '''import { A } from "./a";
export interface Shape { area(): number; }
export abstract class Base<T> implements Shape {
  protected helper(): void { this.area(); }
}
export function make(n: number): Base<number> { return new Impl(n); }
''')
        self.assertEqual(set(syms), {"Shape", "Base", "Base.helper", "make"})
        self.assertIn("Impl", syms["make"]["calls"])

    def test_go_java_rust(self) -> None:
        syms, mod = self._syms("go", '''package main
import (
    "fmt"
    "example.com/repo/lib/util"
)
const Limit = 3
type Engine struct { name string }
// Run runs it.
func (e *Engine) Run() error { util.Helper(e.name); return nil }
func main() { e := &Engine{}; e.Run(); fmt.Println(Limit) }
''')
        self.assertEqual(set(syms), {"Engine", "Engine.Run", "main"})
        self.assertEqual((syms["Engine.Run"]["kind"], syms["Engine.Run"]["docstring"], syms["Engine.Run"]["calls"]), ("method", "Run runs it.", ["util.Helper"]))
        self.assertEqual([i["module"] for i in mod["imports"]], ["fmt", "example.com/repo/lib/util"])
        syms, mod = self._syms("java", '''package com.example.app;
import java.util.List;
import com.example.lib.Util;
/** Engine doc. */
public class Engine extends Base implements Runnable {
    public Engine(String name) { this.name = name; }
    public void run() { Util.helper(name); this.check(); }
    private int check() { return 1; }
}
interface Runnable {
    void run();
}
''')
        self.assertEqual(set(syms), {"Engine", "Engine.Engine", "Engine.run", "Engine.check", "Runnable", "Runnable.run"})
        self.assertEqual((syms["Engine"]["docstring"], syms["Engine"]["bases"]), ("Engine doc.", ["Base", "Runnable"]))
        self.assertEqual(syms["Engine.run"]["calls"], ["Util.helper", "self.check"])
        self.assertEqual([(i["module"], i["name"]) for i in mod["imports"]], [("java.util", "List"), ("com.example.lib", "Util")])
        syms, mod = self._syms("rust", '''use std::collections::HashMap;
use crate::lib::util::helper;
const LIMIT: usize = 3;
/// Engine doc
pub struct Engine { name: String }
pub trait Runner {
    fn run(&self) -> u32;
}
impl Runner for Engine {
    fn run(&self) -> u32 { helper(&self.name); self.check() }
}
impl Engine {
    fn check(&self) -> u32 { LIMIT as u32 }
}
pub fn main() { let e = Engine { name: "x".into() }; e.run(); }
''')
        self.assertEqual(set(syms), {"Engine", "Runner", "Runner.run", "Engine.run", "Engine.check", "main"})
        self.assertEqual(set(syms["Engine.run"]["calls"]), {"helper", "self.check"})
        self.assertEqual([(i["module"], i["name"]) for i in mod["imports"]], [("std::collections", "HashMap"), ("crate::lib::util", "helper")])
        self.assertEqual(mod["constants"], ["LIMIT"])

    def test_c_cpp_csharp(self) -> None:
        syms, mod = self._syms("c", '''#include <stdio.h>
#include "lib/util.h"
#define LIMIT 3
struct Engine { int id; };
/* run it */
static int run(int x) { return helper(x) + LIMIT; }
int main(void) { return run(1); }
''')
        self.assertEqual(set(syms), {"Engine", "run", "main"})
        self.assertEqual((syms["run"]["docstring"], syms["run"]["calls"]), ("run it", ["helper"]))
        self.assertEqual([(i["module"], i["level"]) for i in mod["imports"]], [("<stdio.h>", 0), ("lib/util.h", 1)])
        self.assertEqual(mod["constants"], ["LIMIT"])
        syms, mod = self._syms("cpp", '''#include "lib/util.h"
namespace app {
class Engine : public Base {
public:
  int run() { return helper(id_) + check(); }
private:
  int check() { return 1; }
};
int Engine::check2() { return 1; }
}
int main() { app::Engine e; return e.run(); }
''')
        self.assertEqual(set(syms), {"app.Engine", "app.Engine.run", "app.Engine.check", "app.Engine.check2", "main"})
        self.assertEqual(syms["app.Engine"]["bases"], ["Base"])
        syms, mod = self._syms("csharp", '''using System;
using App.Lib;
namespace App {
  public class Engine : Base, IRunner {
    public Engine(string name) { Name = name; }
    public int Run() { Util.Helper(Name); return Check(); }
    private int Check() => 1;
  }
  public interface IRunner {
    int Run();
  }
}
''')
        self.assertEqual(set(syms), {"App.Engine", "App.Engine.Engine", "App.Engine.Run", "App.Engine.Check", "App.IRunner", "App.IRunner.Run"})
        self.assertEqual(syms["App.Engine"]["bases"], ["Base", "IRunner"])
        self.assertEqual([i["module"] for i in mod["imports"]], ["System", "App.Lib"])

    def test_php_ruby_swift_kotlin_scala(self) -> None:
        syms, mod = self._syms("php", '''<?php
namespace App;
use App\\Lib\\Util;
require_once "lib/util.php";
class Engine extends Base {
    public function run() { Util::helper($this->name); return $this->check(); }
    private function check() { return 1; }
}
function main() { $e = new Engine(); return $e->run(); }
''')
        self.assertEqual(set(syms), {"Engine", "Engine.run", "Engine.check", "main"})
        self.assertEqual(syms["Engine.run"]["calls"], ["Util.helper", "self.check"])
        self.assertEqual([i["module"] for i in mod["imports"]], ["App.Lib.Util", "lib/util.php"])
        syms, mod = self._syms("ruby", '''require "json"
require_relative "lib/util"
module App
  class Engine < Base
    def run
      helper(name)
      check
    end
    def check; 1; end
  end
end
def main; App::Engine.new.run; end
''')
        self.assertEqual(set(syms), {"App.Engine", "App.Engine.run", "App.Engine.check", "main"})
        self.assertIn("helper", syms["App.Engine.run"]["calls"])
        self.assertEqual(syms["App.Engine"]["bases"], ["Base"])
        self.assertEqual([i["module"] for i in mod["imports"]], ["json", "lib/util"])
        syms, mod = self._syms("swift", '''import Foundation
class Engine: Base, Runner {
    func run() -> Int { helper(name); return check() }
    private func check() -> Int { return 1 }
}
protocol Runner {
    func run() -> Int
}
func main() { let e = Engine(name: "x"); _ = e.run() }
''')
        self.assertEqual(set(syms), {"Engine", "Engine.run", "Engine.check", "Runner", "Runner.run", "main"})
        self.assertEqual(syms["Engine"]["bases"], ["Base", "Runner"])
        self.assertEqual([i["module"] for i in mod["imports"]], ["Foundation"])
        syms, mod = self._syms("kotlin", '''package com.example.app
import com.example.lib.Util
class Engine(val name: String) : Base(), Runner {
    override fun run(): Int { Util.helper(name); return check() }
    private fun check(): Int = 1
}
fun main() { Engine("x").run() }
''')
        self.assertEqual(set(syms), {"Engine", "Engine.run", "Engine.check", "main"})
        self.assertEqual(syms["Engine"]["bases"], ["Base", "Runner"])
        self.assertEqual([(i["module"], i["name"]) for i in mod["imports"]], [("com.example.lib", "Util")])
        syms, mod = self._syms("scala", '''package com.example.app
import com.example.lib.Util
class Engine(name: String) extends Base with Runner {
  def run(): Int = { Util.helper(name); check() }
  private def check(): Int = 1
}
trait Runner {
  def run(): Int
}
''')
        self.assertEqual(set(syms), {"Engine", "Engine.run", "Engine.check", "Runner", "Runner.run"})
        self.assertEqual([(i["module"], i["name"]) for i in mod["imports"]], [("com.example.lib", "Util")])

    def test_bash_lua_powershell(self) -> None:
        syms, mod = self._syms("bash", '''#!/bin/bash
source ./lib/util.sh
LIMIT=3
# run it
run() {
  helper "$1"
  check
}
function check { echo "$LIMIT"; }
run "x"
''')
        self.assertEqual(set(syms), {"run", "check"})
        self.assertEqual((syms["run"]["docstring"], syms["run"]["calls"]), ("run it", ["helper", "check"]))
        self.assertEqual(mod["constants"], ["LIMIT"])
        self.assertEqual(mod["calls"], ["run"])
        syms, mod = self._syms("lua", '''local util = require("lib.util")
local function run(x) return util.helper(x) end
Engine = {}
function Engine:start() return self:check() end
''')
        self.assertEqual(set(syms), {"run", "Engine.start"})
        self.assertEqual(syms["Engine.start"]["kind"], "method")
        syms, mod = self._syms("powershell", '''function Check { return 1 }
function Run-It { Check }
Run-It
''')
        self.assertEqual(set(syms), {"Check", "Run-It"})


class MergeBoundaryRuleTests(unittest.TestCase):
    """Three rules of the merge flow at figure / table boundaries (2026-09-05 sampling): paragraphs never go
    into TITLE, title blocks go to the figure / table regardless of its caption, and notes attach backwards
    as FOOTNOTE (one right after a figure / table attaches on the spot, without waiting for the next
    boundary)."""

    PARA = "本节说明产品的供电要求,电压范围与纹波指标见下表。" * 3   # a punctuated paragraph, not a lead-in

    def _run(self, blocks):
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks
        return merge_mineru_text_blocks(blocks, target_tokens=800)

    def test_a_paragraph_before_a_table_is_not_swallowed_into_its_title(self) -> None:
        merged = self._run([
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("p", self.PARA),
            _mb("tbl1", "| a | b |", block_type="table"),
        ])
        self.assertEqual([b.block_id for b in merged], ["tbl0", "mineru-merged-text-00001", "tbl1"])
        self.assertIsNone(merged[2].title)
        self.assertEqual(merged[1].text, self.PARA)

    def test_a_title_block_is_adopted_even_when_the_figure_has_a_caption(self) -> None:
        merged = self._run([
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("h", "直流电气特性", block_type="title"),
            _mb("tbl1", "| a | b |", block_type="table", caption="表 3-1 直流电气特性"),
        ])
        self.assertEqual([b.block_id for b in merged], ["tbl0", "tbl1"])
        self.assertEqual(merged[1].title, "直流电气特性")
        self.assertTrue(merged[1].metadata.get("adopted_titles_only"))

    def test_paragraph_then_heading_before_a_figure(self) -> None:
        merged = self._run([
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("p", self.PARA),
            _mb("h", "3.2 时序图", block_type="title"),
            _mb("img", "", block_type="image"),
        ])
        self.assertEqual([b.block_id for b in merged], ["tbl0", "mineru-merged-text-00001", "img"])
        self.assertEqual(merged[2].title, "3.2 时序图")
        self.assertEqual(merged[1].text, self.PARA)

    def test_a_note_right_after_a_table_becomes_its_footnote(self) -> None:
        from kb_pipeline.chunking.chunker import text_for_block

        merged = self._run([
            _mb("tbl", "| a | b |", block_type="table", table_markdown="| a | b |"),
            _mb("n", "注:数值为典型值"),
            _mb("h", "1.2 下一节", block_type="title"),
            _mb("p", self.PARA),
        ])
        self.assertEqual([b.block_id for b in merged], ["tbl", "mineru-merged-text-00001"])
        self.assertEqual(merged[0].text, "| a | b |\nFOOTNOTE: 注:数值为典型值")
        self.assertEqual(merged[0].table_markdown, merged[0].text)
        self.assertEqual(merged[0].metadata["footnote_block_ids"], ["n"])
        self.assertEqual(text_for_block(merged[0]).count("| a | b |"), 1)
        self.assertTrue(merged[1].text.startswith("1.2 下一节"))
        self.assertEqual(merged[1].metadata["heading_lines"], ["1.2 下一节"])

    def test_a_note_label_followed_by_a_heading_is_dropped(self) -> None:
        merged = self._run([
            _mb("img", "VISUAL SUMMARY: 架构图", block_type="image"),
            _mb("lbl", "注释:"),
            _mb("h", "第2章 概述", block_type="title"),
            _mb("p", self.PARA),
        ])
        self.assertEqual([b.block_id for b in merged], ["img", "mineru-merged-text-00001"])
        self.assertNotIn("FOOTNOTE", merged[0].text)
        self.assertNotIn("注释", merged[1].text)
        self.assertTrue(merged[1].text.startswith("第2章 概述"))

    def test_a_note_between_two_figures_goes_backward(self) -> None:
        merged = self._run([
            _mb("img1", "VISUAL SUMMARY: 图一", block_type="image"),
            _mb("lbl", "注释:"),
            _mb("body", "本图数值来自实测"),
            _mb("img2", "VISUAL SUMMARY: 图二", block_type="image"),
        ])
        self.assertEqual([b.block_id for b in merged], ["img1", "img2"])
        self.assertEqual(merged[0].text, "VISUAL SUMMARY: 图一\nFOOTNOTE: 注释: 本图数值来自实测")
        self.assertEqual(merged[0].metadata["footnote_block_ids"], ["lbl", "body"])
        self.assertIsNone(merged[1].title)


class TableHeaderCandidateTests(unittest.TestCase):
    def _rows(self, *rows):
        return [(i + 1, list(r)) for i, r in enumerate(rows)]

    def test_title_and_note_rows_are_skipped(self) -> None:
        from kb_pipeline.parsers.native_table import chunk_rows_to_blocks, detect_header_rows

        rows = self._rows(
            ["产品需求清单", "", "", ""],
            ["开源组件协议地址: Apache2.0 :http://www.apache.org/licenses/LICENSE-2.0.html", "", "", ""],
            ["模块", "功能", "优先级", "备注"],
            ["登录", "手机号登录", "P0", ""],
        )
        self.assertEqual([rn for rn, _ in detect_header_rows(rows)], [3])
        blocks = chunk_rows_to_blocks("xlsx", rows, "需求", max_tokens=200, overlap_tokens=0)
        self.assertIn("Title: 产品需求清单", blocks[0].text)
        self.assertNotIn("产品需求清单", "\n".join(b.text for b in blocks[1:]))
        self.assertIn("HEADER: 模块", blocks[1].text)

    def test_merged_title_row_is_skipped(self) -> None:
        from kb_pipeline.parsers.native_table import detect_header_rows

        rows = self._rows(["设备清单", "设备清单", "设备清单"], ["名称", "型号", "数量"], ["交换机", "S5720", "2"])
        self.assertEqual([rn for rn, _ in detect_header_rows(rows)], [2])
        rows = self._rows(["名称"], ["交换机"])            # single-column table: the first row is the header
        self.assertEqual([rn for rn, _ in detect_header_rows(rows)], [1])

    def test_dispimg_formula_becomes_a_placeholder(self) -> None:
        from kb_pipeline.parsers.native_table import normalize_cell

        self.assertEqual(normalize_cell('=DISPIMG("ID_3F2A",1)'), "(image)")
        self.assertEqual(normalize_cell("普通"), "普通")

    def test_notes_above_the_header_are_kept_whole_and_only_once(self) -> None:
        """Note rows above the header go only into the summary, and they used to be cut at 200 characters, so the
        second half was in no chunk at all; a merged note block spanning 6 rows was spread into 6 identical lines,
        the summary repeated the same passage 6 times, and chunking produced word-for-word identical chunks."""
        from openpyxl import Workbook

        from kb_pipeline.chunking.chunker import blocks_to_chunks
        from kb_pipeline.parsers.native_table import parse_native_table

        note = "填写说明:" + "".join(f"第{i}条,按实际情况如实填写,不确定的留空。" for i in range(1, 25))
        self.assertGreater(len(note), 400)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "中间件.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "中间件"
            ws["A1"] = "说明"
            ws["B1"] = note
            ws.merge_cells("A1:A6")
            ws.merge_cells("B1:H6")
            for row_no, row in enumerate((["名称", "版本", "用途", "备注"], ["nginx", "1.25", "反向代理", ""], ["redis", "7.2", "缓存", "单机"]), start=8):
                for col, value in enumerate(row, start=1):
                    ws.cell(row=row_no, column=col, value=value)
            wb.save(path)
            blocks = parse_native_table(path, max_tokens=200, overlap_tokens=0)
        summary = blocks[0]
        self.assertEqual(summary.text.count("Title: "), 1)
        self.assertIn(f"Title: 说明 {note}", summary.text)
        self.assertIn("Columns: 名称 | 版本 | 用途 | 备注", summary.text)
        self.assertNotIn("填写说明", "\n".join(b.text for b in blocks[1:]))
        chunks = blocks_to_chunks(kb_id="kb", file_key=1, content_version="v", parser_profile="p",
                                  blocks=blocks, max_tokens=200, overlap_tokens=0)
        texts = [c.text for c in chunks if c.block.block_id == summary.block_id]
        self.assertGreater(len(texts), 1)                                       # the note is longer than the budget: split, every piece within it
        self.assertTrue(all(c.token_count <= 200 for c in chunks), [c.token_count for c in chunks])
        self.assertEqual(len(set(texts)), len(texts))                           # no word-for-word identical chunks
        body = [l for t in texts for l in t.splitlines() if not l.startswith(("TITLE:", "Sheet:", "Rows:", "Columns:"))]
        self.assertEqual("".join(body), f"Title: 说明 {note}")                   # split by sentence without losing a character, and not labelled like a table row

    def test_xlsx_cells_read_the_way_the_sheet_shows_them(self) -> None:
        """The number format used to be ignored: a cell showing 6% was indexed as 0.06, one showing 92% as
        0.916666666666667, and a whole date carried 00:00:00."""
        import datetime

        from openpyxl import Workbook

        from kb_pipeline.parsers.native_table import normalize_cell, read_xlsx_rows

        cells = [
            ("覆盖率", 0.916666666666667, "0%", "92%"),
            ("税率", 0.06, "0%", "6%"),
            ("增长率", 0.546, "0.00%", "54.60%"),
            ("降幅", -0.0525, "0.0%;[Red]-0.0%", "-5.3%"),
            ("进位", 0.545, "0%", "55%"),                       # 0.54500000000000004 in the machine, the sheet shows 55%
            ("字面百分号", 0.06, '0.00"%"', "0.06"),             # a quoted % is just a character, no multiplication by 100
            ("勾选", 1, '[=1]"☑";[=0]"☐"', "1"),
            ("公式结果", 0.1 + 0.2, "General", "0.3"),
            ("整数", 3.0, "General", "3"),
            ("金额", 1234.5, '"￥"#,##0.00', "1234.5"),          # currency symbol and thousands separators are not restored
            ("日期", datetime.datetime(2024, 1, 1), "yyyy-mm-dd", "2024-01-01"),
            ("时刻", datetime.datetime(2024, 1, 1, 9, 30), "yyyy-mm-dd h:mm", "2024-01-01 09:30:00"),
            ("启用", True, "General", "TRUE"),
            ("文本", "0.06", "0%", "0.06"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "格式.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append([name for name, *_ in cells])
            ws.append([value for _, value, *_ in cells])
            for col, (_, _, number_format, _) in enumerate(cells, start=1):
                ws.cell(row=2, column=col).number_format = number_format
            ws.append(["合并的百分比", 0.25])
            ws.cell(row=3, column=2).number_format = "0%"
            ws.merge_cells("B3:B4")
            ws.cell(row=4, column=1).value = "下一行"
            wb.save(path)
            rows = read_xlsx_rows(path)[0][1]
        self.assertEqual(rows[1][1], [shown for *_, shown in cells])
        self.assertEqual([values[:2] for _, values in rows[2:]], [["合并的百分比", "25%"], ["下一行", "25%"]])   # values spread over a merged range are written the same way
        self.assertEqual(normalize_cell(0.06), "0.06")                          # csv / xls carry no format information: the number as it is
        self.assertEqual(normalize_cell(45123.0), "45123")
        # A very large number in a percentage cell: still written out, it must not fail the whole sheet; one that
        # overflows when multiplied by 100 is written as if it had no format
        self.assertEqual(normalize_cell(1e30, "0%"), "1" + "0" * 32 + "%")
        self.assertEqual(normalize_cell(1e24, "0.00%"), "1" + "0" * 26 + ".00%")
        self.assertEqual(normalize_cell(1.7e308, "0%"), normalize_cell(1.7e308))
        import xlrd

        from kb_pipeline.parsers.native_table import xls_cell_text

        self.assertEqual(xls_cell_text(xlrd.XL_CELL_NUMBER, 0.1 + 0.2, 0), "0.3")   # numbers from xls are written the same way
        self.assertEqual(parser_profile_for_path(Path("a.xlsx")), "table-native-v4")


class VisualContractTests(unittest.TestCase):
    def test_contract_v5_has_decorative_and_kind_specific_guidance(self) -> None:
        self.assertEqual(vlm.CAPTION_CONTRACT_VERSION, 5)
        self.assertIn("decorative", vlm.RESULT_SCHEMA["properties"])
        self.assertIn("decorative", vlm.RESULT_SCHEMA["required"])
        self.assertIs(vlm.EMPTY_RESULT["decorative"], False)
        guide = getattr(vlm, "_FIELD_GUIDE", "") + vlm.DEFAULT_PROMPT
        for needle in ("UI screenshot", "output mermaid", "lists every name", "decorative"):
            self.assertIn(needle, guide)
        self.assertIs(vlm.normalize_result({"kind": "photo", "decorative": "true"})["decorative"], True)
        self.assertIs(vlm.normalize_result({"kind": "photo"})["decorative"], False)

    def test_tidy_visual_text_drops_fences_and_repeats_and_caps_length(self) -> None:
        from kb_pipeline.parsers.visual_blocks import tidy_visual_text
        from kb_pipeline.utils import count_tokens

        text = "首页\n首页\n首页\n首页\n```mermaid\ngraph LR\nA-->B\n```\n设置\n" + "\n".join(f"第{i}行内容" for i in range(300))
        out, stats = tidy_visual_text(text, kind="screenshot", max_tokens=100)
        self.assertNotIn("mermaid", out)
        self.assertEqual(out.count("首页"), 2)
        self.assertEqual(stats["fences_dropped"], 1)
        self.assertEqual(stats["lines_dropped"], 2)
        self.assertGreater(stats["tokens_cut"], 0)
        self.assertLessEqual(count_tokens(out), 100)
        kept, _ = tidy_visual_text("```mermaid\ngraph LR\nA-->B\n```", kind="diagram", max_tokens=100)
        self.assertIn("mermaid", kept)

    def test_tidy_visual_blocks_and_decorative_flag_are_wired(self) -> None:
        from kb_pipeline.parsers.visual_blocks import tidy_visual_blocks

        block = _mb("img", "首页\n首页\n首页\n首页", block_type="image", metadata={"visual_kind": "screenshot"})
        self.assertEqual(tidy_visual_blocks([block]), 1)
        self.assertEqual(block.text, "首页\n首页")
        self.assertEqual(block.metadata["visual_text_tidy"]["lines_dropped"], 2)
        self.assertIn('"decorative": decorative', _repo_file("app/kb_pipeline/parsers/visual_blocks.py"))
        for name in ("pdf_enhanced", "docx_enhanced", "pptx_enhanced"):
            self.assertIn("tidy_visual_blocks(", _repo_file(f"app/kb_pipeline/parsers/{name}.py"), name)


class MarkdownAndSlideSectionTests(unittest.TestCase):
    def test_markdown_headings_are_cleaned_and_rules_dropped(self) -> None:
        from kb_pipeline.parsers.common import markdown_to_blocks

        md = "## **2.1 安装**\n\n正文一。\n\n---\n\n正文二。\n\n⸻\n\n### ==📅 2026-01-15== 周会\n\n正文三。\n"
        blocks = markdown_to_blocks(md, parser="md", parser_profile="p", doc_type="md")
        # markdown headings do not become blocks of their own: they hang on the body block's title and
        # section_path
        self.assertFalse(any(l.strip() in {"---", "⸻"} for b in blocks for l in (b.text or "").splitlines()))
        self.assertEqual([b.title for b in blocks], ["2.1 安装", "📅 2026-01-15 周会"])
        self.assertEqual(blocks[0].text, "正文一。\n\n正文二。")
        path = list(blocks[1].metadata.get("section_path") or [])
        self.assertIn("2.1 安装", path)
        self.assertEqual(path[-1], "📅 2026-01-15 周会")

    def test_pptx_slides_get_a_section_from_their_first_line(self) -> None:
        from kb_pipeline.parsers.mineru_pptx import aggregate_mineru_slide_blocks

        def blk(i, text, bt="text", slide=1):
            return _mb(f"b{i}", text, block_type=bt, doc_type="pptx", slide_idx=slide, page_idx=slide)

        blocks = [blk(1, "**产品架构**", slide=1), blk(2, "三层:接入、服务、数据。", slide=1),
                  blk(3, "VISUAL SUMMARY: 架构图", "image", slide=1), blk(4, "客户案例", slide=2), blk(5, "某某医院", slide=2)]
        out = aggregate_mineru_slide_blocks(blocks, Path("deck.pptx"))
        slides = [b for b in out if b.block_type == "slide"]
        self.assertEqual([b.metadata["section_path"] for b in slides], [["产品架构"], ["客户案例"]])
        image = [b for b in out if b.block_type == "image"][0]
        self.assertEqual(image.metadata["section_path"], ["产品架构"])
        self.assertEqual([b.block_type for b in out], ["slide", "image", "slide"])   # the image follows its own slide


class MergeBoundaryRuleFollowupTests(unittest.TestCase):
    """Two findings from the second sampling round: TOC entries do not follow a heading into TITLE; inline
    icons are not merge boundaries."""

    PARA = MergeBoundaryRuleTests.PARA

    def _run(self, blocks):
        from kb_pipeline.parsers.pdf_enhanced import merge_mineru_text_blocks
        return merge_mineru_text_blocks(blocks, target_tokens=800)

    def test_toc_lines_before_a_title_stay_in_the_text(self) -> None:
        merged = self._run([
            _mb("tbl0", "前一张表", block_type="table"),
            _mb("toc1", "技术支持....18"),
            _mb("toc2", "赛普拉斯开发者社区 18"),
            _mb("h", "引脚配置", block_type="title"),
            _mb("img", "", block_type="image"),
        ])
        self.assertEqual([b.block_id for b in merged], ["tbl0", "mineru-merged-text-00001", "img"])
        self.assertEqual(merged[2].title, "引脚配置")
        self.assertEqual(merged[1].text, "技术支持....18\n\n赛普拉斯开发者社区 18")

    def test_inline_icons_do_not_break_the_merge(self) -> None:
        merged = self._run([
            _mb("p1", self.PARA),
            _mb("icon", "", block_type="image", bbox=[46.0, 167.0, 69.0, 190.0]),
            _mb("p2", self.PARA),
        ])
        self.assertEqual([b.block_id for b in merged], ["mineru-merged-text-00001"])
        self.assertEqual(merged[0].metadata["source_block_ids"], ["p1", "p2"])
        merged = self._run([_mb("p1", self.PARA), _mb("fig", "", block_type="image", bbox=[40.0, 100.0, 400.0, 380.0])])
        self.assertEqual([b.block_id for b in merged], ["mineru-merged-text-00001", "fig"])

    def test_leader_lines_are_not_lead_ins(self) -> None:
        from kb_pipeline.headings import is_short_lead_in

        self.assertTrue(is_short_lead_in("三种访问方式"))
        self.assertFalse(is_short_lead_in("技术支持....18"))
        self.assertFalse(is_short_lead_in("栓锁电流 ..... > 140 mA"))
        self.assertFalse(is_short_lead_in("电容 …… 7"))


class VisualEntitiesAndTitleCleanupTests(unittest.TestCase):
    def test_entities_are_rendered_into_the_text(self) -> None:
        from PIL import Image
        from kb_pipeline.parsers import visual_blocks

        def fake_caption(jobs, **kwargs):
            return {job_id: {"kind": "photo", "title": "", "summary": "客户 logo 墙", "text_verbatim": "",
                             "entities": ["北岭医院", "南川医院"], "facts": [], "keywords": [], "confidence": "high",
                             "decorative": False}
                    for job_id, *_ in jobs}

        with tempfile.TemporaryDirectory() as tmp, patch.object(visual_blocks, "caption_images_parallel", fake_caption):
            img = Path(tmp) / "i.png"; Image.new("RGB", (300, 200)).save(img)
            block = ParsedBlock(parser="p", parser_profile="p", doc_type="pptx", block_type="image", text="",
                                block_id="a", visual_ref=str(img))
            visual_blocks.enrich_blocks_with_vlm([block], base_url="u", api_key="k", model_id="m",
                                                 cache_dir=Path(tmp), filter_decorative=False)
        self.assertIn("ENTITIES: 北岭医院，南川医院", block.text)
        self.assertFalse(block.metadata.get("decorative"))

    def test_title_blocks_lose_bold_markup_in_both_parsers(self) -> None:
        for name in ("mineru_pdf", "mineru_docx"):
            src = _repo_file(f"app/kb_pipeline/parsers/{name}.py")
            self.assertIn("text = clean_heading_text(text) or text", src, name)


class ParserDetailTests(unittest.TestCase):
    """Health check D2 / R8 / R9: detail improvements on the parsing side."""

    def test_generic_html_pages_are_split_by_headings(self) -> None:
        """R9: HTML without custom class names is split into sections by h1–h6 and the blocks carry
        section_path; heading text stays out of the body."""
        from kb_pipeline.parsers.html_dom import parse_html_dom

        preface = "前言文字," * 12
        html_text = (f"<html><head><title>手册</title></head><body><p>{preface}</p>"
                     "<h1>安装</h1><p>安装步骤一。</p><p>安装步骤二。</p>"
                     "<h2>依赖</h2><ul><li>依赖 A</li><li>依赖 B</li></ul>"
                     "<h1>使用</h1><p>使用说明。</p></body></html>")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "manual.html"
            path.write_text(html_text, encoding="utf-8")
            blocks = parse_html_dom(path)
        self.assertEqual([b.metadata.get("html_block_type") for b in blocks],
                         ["overview", "section", "section", "section"])
        self.assertEqual([b.metadata.get("section_path") for b in blocks[1:]], [["安装"], ["安装", "依赖"], ["使用"]])
        self.assertEqual(blocks[1].text.splitlines(), ["安装步骤一。", "安装步骤二。"])
        self.assertEqual(blocks[2].title, "依赖")
        self.assertIn("- 依赖 A", blocks[2].text)
        self.assertTrue(all(b.block_id.startswith("page-1-") for b in blocks))
        self.assertIn("前言文字", blocks[0].text)      # a UTF-8 page without a declared charset is no longer read as Latin-1

    def test_html_tables_stay_in_their_section_and_spans_are_filled(self) -> None:
        """The tables of a generic HTML page all used to come after every section, without a section path and titled
        after the first heading of the page, so nothing told which table belonged to which section; header cells
        spanning rows / columns were ignored, so the second header row was shifted as a whole."""
        from kb_pipeline.parsers.html_dom import parse_html_dom

        html_text = ("<html><head><title>手册</title></head><body><p>" + "前言文字," * 12 + "</p>"
                     "<h1>型号 A 参数</h1><p>型号 A 的说明。</p>"
                     "<table><tr><th rowspan='2'>项目</th><th colspan='2'>数值</th></tr><tr><th>最小</th><th>最大</th></tr>"
                     "<tr><td>电压</td><td>3.0</td><td>3.6</td></tr></table><p>表后的补充说明。</p>"
                     "<h2>订购信息</h2><table><tr><td>料号</td><td>封装</td></tr><tr><td>A-100</td><td>QFN</td></tr></table>"
                     "<h1>型号 B 参数</h1><p>型号 B 的说明。</p>"
                     "<table><tr><th>项目</th><th>数值</th></tr><tr><td>电压</td><td>5.0</td></tr></table></body></html>")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "manual.html"
            path.write_text(html_text, encoding="utf-8")
            blocks = parse_html_dom(path)
        self.assertEqual([(b.block_type, b.metadata.get("section_path")) for b in blocks], [
            ("text", None), ("text", ["型号 A 参数"]), ("table", ["型号 A 参数"]),
            ("table", ["型号 A 参数", "订购信息"]),                                   # this section has a table and no body
            ("text", ["型号 B 参数"]), ("table", ["型号 B 参数"])])
        tables = [b for b in blocks if b.block_type == "table"]
        self.assertEqual([b.title for b in tables], ["型号 A 参数 / Table 1", "订购信息 / Table 2", "型号 B 参数 / Table 3"])
        self.assertEqual(len({b.block_id for b in blocks}), len(blocks))
        self.assertEqual(tables[0].text.splitlines(), ["| 项目 | 数值 | 数值 |", "| --- | --- | --- |",
                                                       "| 项目 | 最小 | 最大 |", "| 电压 | 3.0 | 3.6 |"])
        self.assertIn("| 电压 | 5.0 |", tables[2].text)
        self.assertEqual(blocks[1].text.splitlines(), ["型号 A 的说明。", "表后的补充说明。"])   # the text of the tables does not go into the body
        self.assertTrue(all(b.parser_profile == "html-dom-v2" for b in blocks))
        self.assertEqual(parser_profile_for_path(Path("a.html")), "html-dom-v2")

    def test_hidden_sheets_and_rows_are_skipped(self) -> None:
        """R8: hidden sheets / hidden rows are not indexed."""
        from openpyxl import Workbook

        from kb_pipeline.parsers.native_table import read_xlsx_rows

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "t.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"
            for row in (["name", "value"], ["a", 1], ["old", 2], ["b", 3]):
                ws.append(row)
            ws.row_dimensions[3].hidden = True
            backup = wb.create_sheet("Backup")
            backup.append(["x", "y"])
            backup.sheet_state = "hidden"
            wb.save(path)
            sheets = read_xlsx_rows(path)
        self.assertEqual([name for name, *_ in sheets], ["Data"])
        self.assertEqual([values[0] for _, values in sheets[0][1]], ["name", "a", "b"])

    def test_xlsx_read_is_bounded_by_the_cells_that_hold_content(self) -> None:
        """2026-09-29 audit: a sheet with a whole row formatted has max_column 16384; building cells row by row up
        to it costs about 4.7 MB per row, tens of GB for a few thousand rows. The read stops at the last row and
        column with content; what is read is the same as without the bound."""
        from openpyxl import Workbook
        from openpyxl.styles import Font

        from kb_pipeline.parsers import native_table
        from kb_pipeline.parsers.native_table import read_xlsx_rows

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "styled.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "Data"
            ws.append(["group", "name", "count"])
            for i in range(1, 41):
                ws.append(["alpha" if i == 1 else None, f"n{i}", i])
            ws.merge_cells("A2:A41")                                   # a category label merged downwards: every row carries it
            ws.cell(row=2, column=16384).font = Font(bold=True)        # formatting out to column XFD, no value
            ws.cell(row=9000, column=2).font = Font(bold=True)         # formatting far below the data, no value
            wb.save(path)
            seen: list[dict] = []
            original = native_table.load_workbook_compat

            def spy(*args, **kwargs):
                book = original(*args, **kwargs)
                for sheet in book.worksheets:
                    real = sheet.iter_rows

                    def bounded(*a, _real=real, **k):
                        seen.append(dict(k))
                        return _real(*a, **k)

                    sheet.iter_rows = bounded
                return book

            with patch.object(native_table, "load_workbook_compat", spy):
                sheets = read_xlsx_rows(path)
        self.assertEqual(seen, [{"max_row": 41, "max_col": 3}])        # the declared range is 9000 rows x 16384 columns
        name, rows, _spans = sheets[0]
        self.assertEqual((name, len(rows)), ("Data", 41))
        self.assertEqual(rows[0], (1, ["group", "name", "count"]))
        self.assertEqual(rows[1], (2, ["alpha", "n1", "1"]))
        self.assertEqual(rows[40], (41, ["alpha", "n40", "40"]))       # the value filled down the merged range is still there
        self.assertTrue(all(len(values) == 3 for _, values in rows))

    def test_vlm_prompt_language_follows_the_kb(self) -> None:
        """D2: when the KB's output language is known, the image caption prompt names that language; if no
        language has been extracted, follow the language of the text in the image."""
        from kb_pipeline.vision.vlm import DEFAULT_PROMPT, compose_prompt, language_is_chinese, prompt_with_language

        self.assertFalse(language_is_chinese(None))
        self.assertTrue(language_is_chinese("Chinese"))
        self.assertTrue(language_is_chinese("简体中文"))
        self.assertFalse(language_is_chinese("English"))
        unknown = prompt_with_language(None, "")
        self.assertTrue(unknown.startswith(DEFAULT_PROMPT))
        self.assertIn("in the language of the text in the image", unknown)
        zh = prompt_with_language("custom", "zh-CN")
        self.assertTrue(zh.startswith("custom"))
        self.assertIn("Output language: zh-CN", zh)
        en = prompt_with_language(None, "English")
        self.assertTrue(en.startswith(DEFAULT_PROMPT))
        self.assertIn("Output language: English", en)
        self.assertIn("Output language: English", compose_prompt("figure", en))
        src = _repo_file("app/kb_pipeline/pipeline/parse_job.py")
        self.assertIn("vlm_prompt = prompt_with_language(source.vlm_prompt, source.graph_language)", src)
        self.assertNotIn("vlm_prompt=source.vlm_prompt", src)


class PdfReportParseTests(unittest.TestCase):
    """Three parse defects common in report-style PDFs: lost list items, tables broken across pages, and
    misjudged decorative images."""

    def test_mineru_list_items_are_indexed(self) -> None:
        """In MinerU 3.x the text of list items lives entirely in list_items; it used to be dropped wholesale
        (recommendation and explanation paragraphs)."""
        from kb_pipeline.parsers.common import item_text

        self.assertEqual(item_text({"type": "list", "sub_type": "text", "list_items": ["1、安装前请核对支架型号…", "2、螺栓按对角顺序拧紧…"], "page_idx": 4}),
                         "1、安装前请核对支架型号…\n2、螺栓按对角顺序拧紧…")
        self.assertEqual(item_text({"type": "list", "list_items": []}), "")
        self.assertEqual(item_text({"type": "text", "text": "正文", "list_items": ["a"]}), "正文\na")

    def test_mineru_code_blocks_are_indexed(self) -> None:
        """2026-09-29 audit: the text of MinerU's code listings / algorithm boxes is in code_body; it used to yield
        an empty string and the item was skipped as a whole, caption included. Every code item now becomes a code
        block with its caption, and the chunks carry the code."""
        from unittest import mock

        from kb_pipeline.chunking.chunker import blocks_to_chunks
        from kb_pipeline.parsers import mineru_pdf
        from kb_pipeline.parsers.common import item_text
        from kb_pipeline.parsers.pdf_enhanced import is_mergeable_text_block

        code = "def area(r):\n    return 3.14 * r * r"
        self.assertEqual(item_text({"type": "code", "sub_type": "code", "code_body": code, "page_idx": 3}), code)
        self.assertEqual(item_text({"type": "code", "code_body": "  "}), "")
        payload = {"content_list": [
            {"type": "text", "text": "Chapter 3 Areas", "text_level": 1, "page_idx": 3},
            {"type": "text", "text": "The function below returns the area of a circle.", "page_idx": 3},
            {"type": "code", "sub_type": "code", "code_body": code, "code_caption": ["Listing 3-1 Circle area"], "page_idx": 3,
             "bbox": [10, 20, 300, 80]},
            {"type": "code", "sub_type": "algorithm", "code_body": "Input: radius r\nOutput: area S\n1. S <- pi * r * r", "page_idx": 4},
            {"type": "code", "sub_type": "code", "code_body": "", "page_idx": 4},
        ]}
        with tempfile.TemporaryDirectory() as tmp, \
                mock.patch.object(mineru_pdf, "call_mineru_sync", return_value=payload):
            blocks = mineru_pdf.mineru_pdf_blocks(mineru_url="http://x", path=Path(tmp) / "book.pdf", cache_dir=Path(tmp))
        kinds = [(b.block_type, b.metadata.get("source_type")) for b in blocks]
        self.assertEqual(kinds, [("title", "text"), ("text", "text"), ("code", "code"), ("code", "code")])   # an empty code item is still skipped
        listing, algorithm = blocks[2], blocks[3]
        self.assertEqual((listing.text, listing.caption, listing.page_idx), (code, "Listing 3-1 Circle area", 4))   # pages count from 1
        self.assertIn("S <- pi * r * r", algorithm.text)
        self.assertEqual(listing.metadata["section_path"], ["Chapter 3 Areas"])                       # under its chapter
        self.assertFalse(is_mergeable_text_block(listing))                                           # not merged with the prose around it
        chunks = blocks_to_chunks(kb_id="kb_x", file_key=1, content_version="v", parser_profile="p", blocks=blocks,
                                  max_tokens=400, overlap_tokens=0)
        text = "\n".join(c.text for c in chunks)
        self.assertIn("CAPTION: Listing 3-1 Circle area", text)
        self.assertIn("def area(r):\n    return 3.14 * r * r", text)                                  # the code keeps its lines
        self.assertIn("Input: radius r", text)

    def test_split_tables_are_merged_back(self) -> None:
        """A lab report split across pages: the piece on the next page without a header (or with the header
        reprinted) is merged back into the previous table; a different column count, a separate caption or a
        merged-cells flag prevents the merge."""
        from kb_pipeline.parsers.pdf_enhanced import merge_split_tables

        def table(bid, md, page, **kw):
            return _mb(bid, md, block_type="table", table_markdown=md, page_idx=page, **kw)

        head = "| 检查项目 | 缩写 | 测量结果 | 参考区间 | 单位 |\n| --- | --- | --- | --- | --- |"
        a = table("t1", "CAPTION: 血常规\n" + head + "\n| 血小板计数 | PLT | 236 | 125—350 | ×10^9/L |\nFOOTNOTE: 注 1", 9)
        # After html_table_to_markdown the continuation's first row is also taken as a header and followed by
        # the separator line -- that is the real shape
        b = table("t2", "| 血小板分布宽度 | PDW | 13.4 | 11.0—18.0 | fL |\n| --- | --- | --- | --- | --- |\n| 单核细胞绝对值 | MONO | 0.41 | 0.10—0.60 | ×10^9/L |\nFOOTNOTE: 注 1", 10)
        c = table("t3", head + "\n| 总蛋白 | TP | 71.8 | 65.0—85.0 | g/L |", 10)                    # header reprinted
        d = table("t4", "| 项目 | 结果 |\n| --- | --- |\n| 视力 | 5.0 |", 11)                           # different column count: another table
        e = table("t5", "| 检查项目 | 缩写 | 测量结果 | 参考区间 | 单位 |\n| --- | --- | --- | --- | --- |\n| x | y | 1 | 2 | z |", 11, caption="尿常规")
        out = merge_split_tables([_mb("p", "前文"), a, b, c, d, e])
        self.assertEqual([x.block_id for x in out], ["p", "t1", "t4", "t5"])
        rows = [l for l in a.table_markdown.splitlines() if l.startswith("|")]
        self.assertEqual(len(rows), 2 + 4)                                      # header + separator + 4 data rows
        self.assertTrue(a.table_markdown.startswith("CAPTION: 血常规\n| 检查项目"))
        self.assertTrue(a.table_markdown.endswith("| 总蛋白 | TP | 71.8 | 65.0—85.0 | g/L |\nFOOTNOTE: 注 1"))
        self.assertEqual(a.text, a.table_markdown)
        self.assertEqual((a.metadata["merged_tables"], a.metadata["page_end"]), (["t2", "t3"], 10))
        # A table carrying merged-cell flags is not merged (screenshot verification maps rows against the
        # original grid); neither is one two pages away
        f = table("t6", head + "\n| a | b | 1 | 2 | c |", 12, metadata={"table_flags": [{"row": 2, "confidence": "high"}]})
        g = table("t7", "| d | e | 3 | 4 | f |\n| --- | --- | --- | --- | --- |", 12)
        self.assertEqual(len(merge_split_tables([f, g])), 2)
        h = table("t8", head + "\n| a | b | 1 | 2 | c |", 1)
        i = table("t9", "| d | e | 3 | 4 | f |\n| --- | --- | --- | --- | --- |", 5)
        self.assertEqual(len(merge_split_tables([h, i])), 2)
        # A table spanning three pages: the part on the third page continues the second page; it used to be compared
        # with the first page, and two pages apart it was not merged
        first = table("t10", head + "\n| a | b | 1 | 2 | c |", 1)
        rest = [table(f"t1{n}", f"| r{n} | x | {n} | 4 | f |\n| --- | --- | --- | --- | --- |", n + 1) for n in (1, 2, 3)]
        self.assertEqual([x.block_id for x in merge_split_tables([first, *rest])], ["t10"])
        self.assertEqual((first.metadata["merged_tables"], first.metadata["page_end"]), (["t11", "t12", "t13"], 4))
        self.assertEqual([l.split(" | ")[0] for l in first.table_markdown.splitlines()[2:]], ["| a", "| r1", "| r2", "| r3"])
        for name in ("pdf_enhanced.py", "docx_enhanced.py"):
            self.assertIn("blocks = merge_split_tables(blocks)", _repo_file(f"app/kb_pipeline/parsers/{name}"), name)

    def test_decorative_verdict_keeps_substantive_photos_and_folds_codes(self) -> None:
        """A full-page product photo tagged photo by the VLM with no facts / entities used to be folded away
        as decorative under "photo without content", while QR codes / barcodes / icons / signatures each
        became their own chunk."""
        from kb_pipeline.parsers.visual_blocks import is_decorative_result

        self.assertFalse(is_decorative_result("photo", "六张显微照片,显示了同一批样品在不同放大倍数下的晶粒结构,图像中可见多处孔隙。",
                                              facts=[], entities=[], verbatim="", flagged=False))
        self.assertTrue(is_decorative_result("photo", "一只手握着一支装有蓝色液体的试管。", facts=[], entities=[], verbatim="", flagged=False))
        self.assertFalse(is_decorative_result("photo", "一只手握着一支装有蓝色液体的试管。", facts=["x"], entities=[], verbatim="", flagged=False))
        for summary in ("一个二维码,中心嵌入了“NW”标志。", "一个标准的黑白条形码图案。", "一张手写的中文签名,内容为“张三”。",
                        "一个绿色的图标,描绘了一个人骑自行车的形象。", "一组与物流相关的抽象图标,包括卡车、包裹和时钟。"):
            self.assertTrue(is_decorative_result("screenshot", summary, facts=["f"], entities=["e"], verbatim="", flagged=False), summary)
        self.assertFalse(is_decorative_result("screenshot", "示波器屏幕显示一段方波及相关的测量参数和操作界面,右下角有一个二维码。",
                                              facts=["f"], entities=[], verbatim="", flagged=False))
        self.assertTrue(is_decorative_result("chart", "任何描述", facts=["f"], entities=[], verbatim="", flagged=True))

    def test_table_pieces_keep_the_header_behind_a_two_line_caption(self) -> None:
        """A MinerU table caption spans two lines ("blood routine" + "operator: ..."), the second without a
        prefix: the chunker used to stop when it hit it, the header row never made it into the later pieces,
        and the lab report continuation pieces were all value rows without column names."""
        from kb_pipeline.chunking.chunker import split_table_text
        from kb_pipeline.utils import count_tokens

        head = "TITLE: 初步意见\nCAPTION: 血常规\n操作者:张三 审核者:李四\n| 检查项目 | 缩写 | 测量结果 | 参考区间 | 单位 |\n| --- | --- | --- | --- | --- |"
        rows = "\n".join(f"| 项目{i} | ABC{i} | {i}.{i} | 0.0—{i} | g/L |" for i in range(40))
        pieces = split_table_text(f"{head}\n{rows}", 120)
        self.assertGreater(len(pieces), 2)
        for piece in pieces:
            self.assertTrue(piece.startswith(head + "\n"), piece[:120])
            self.assertLessEqual(count_tokens(piece), 120)
        body = [l for p in pieces for l in p.splitlines()[5:]]
        self.assertEqual(body, rows.splitlines())
        # Lead-in line too long (over a third of the budget): carry only the prefix lines and the header; the
        # long lead-in text goes first as body text and is not lost; when it alone exceeds the budget it is split
        # to the budget as a passage of text, not as a table row, so it gets no column names
        lead_in = "很长的引导说明," * 60
        long_lead = "CAPTION: 表\n" + lead_in + "\n| a | b |\n| --- | --- |"
        pieces = split_table_text(long_lead + "\n" + "\n".join(f"| {i} | {i} |" for i in range(80)), 120)
        self.assertTrue(all(p.startswith("CAPTION: 表\n| a | b |\n| --- | --- |") for p in pieces))
        self.assertTrue(all(count_tokens(p) <= 120 for p in pieces), [count_tokens(p) for p in pieces])
        rest = [l for p in pieces for l in p.splitlines()[3:]]
        self.assertEqual("".join(l for l in rest if not l.startswith("|")), lead_in)
        self.assertEqual([l for l in rest if l.startswith("|")], [f"| {i} | {i} |" for i in range(80)])
        self.assertNotIn("a: ", "\n".join(pieces))


class VisualEvidenceTests(unittest.TestCase):
    """Text in the image (caption / transcription / chart reading) is evidence, the model's FACTS are
    interpretation: when the caption says 28 and the model says 25, the text wins and a trace is kept.
    Another kind of runaway: a screenshot's FACTS swelling item by item to thousands of tokens: collapse,
    flag as runaway, retry, never cache."""

    def test_model_estimates_yield_to_on_image_text(self) -> None:
        from kb_pipeline.parsers.visual_blocks import reconcile_visual_facts

        facts = ["综合风险指数 25,中低等级", "供电可靠性指数 27", "过热风险指数 0", "EMC 测试结果 一般"]
        trusted = "综合风险指数 28 中低等级\n供电可靠性指数 27\n过热风险指数:2\nEMC 测试结果 正常"
        out, _, conflicts = reconcile_visual_facts(facts, trusted)
        self.assertEqual(out[0], "综合风险指数 28,中低等级")
        self.assertEqual(out[1], "供电可靠性指数 27")                 # consistent ones untouched
        self.assertEqual(out[2], "过热风险指数 2")
        self.assertEqual(out[3], "EMC 测试结果 一般")                        # non-numeric category conflicts are not on this path
        self.assertEqual([(c["label"], c["text_value"], c["model_value"]) for c in conflicts],
                         [("综合风险指数", "28", "25"), ("过热风险指数", "2", "0")])
        # The real-box shape: the caption names the metric while the model summary and FACTS say "current risk
        # value is 25" -- a generic label is matched against the single value in the image
        summary = "一个风险水平的进度条,显示当前风险值为25,并与同类产品平均值30比较"
        out2, sum2, conf2 = reconcile_visual_facts(["当前风险值为25", "同类产品平均值为30"], "综合风险指数 28 中低等级", summary)
        self.assertEqual((out2, sum2), (["当前风险值为28", "同类产品平均值为30"], "一个风险水平的进度条,显示当前风险值为28,并与同类产品平均值30比较"))
        self.assertEqual([(c["label"], c["text_value"], c["model_value"], c["model_label"]) for c in conf2],
                         [("综合风险指数", "28", "25", "当前风险值为"), ("综合风险指数", "28", "25", "显示当前风险值为")])
        # No collateral damage: reference ranges / normal values are not readings; metrics whose content words
        # differ (power stability index vs overall risk index) are not the same; unmatched labels stay as is
        keep = ["输出电压参考范围 3.1-5.2", "输出电压正常值 5.2", "供电稳定指数 24", "转速 3.3"]
        self.assertEqual(reconcile_visual_facts(keep, "输出电压 6.0\n综合风险指数 28\n电流 5")[0], keep)
        self.assertEqual(reconcile_visual_facts(["转速 3.3"], ""), (["转速 3.3"], "", []))
        src = _repo_file("app/kb_pipeline/parsers/visual_blocks.py")
        self.assertIn("facts, summary, value_conflicts = reconcile_visual_facts(", src)
        self.assertIn('"visual_value_conflicts": value_conflicts', src)

    def test_runaway_captions_are_collapsed_retried_and_never_cached(self) -> None:
        from kb_pipeline.vision.vlm import normalize_result

        chain = ["待处理文件夹"] + ["待处理文件夹" + "内文件夹" * i for i in range(1, 30)]
        data = normalize_result({"kind": "screenshot", "summary": "云杀毒概况页面", "facts": ["左侧导航栏显示安全管控"] + chain,
                                 "entities": ["云盘", "安全管控"], "keywords": []})
        self.assertEqual(data["facts"], ["左侧导航栏显示安全管控", "待处理文件夹"])
        self.assertTrue(data["runaway"])
        self.assertEqual(data["runaway_dropped"], 29)
        normal = normalize_result({"kind": "chart", "summary": "s", "facts": ["A 1", "B 2", "C 3"], "entities": [], "keywords": []})
        self.assertEqual(normal["facts"], ["A 1", "B 2", "C 3"])
        self.assertNotIn("runaway", normal)
        # repetition inside a single item counts too
        loop = normalize_result({"kind": "screenshot", "summary": "s", "facts": ["首页 密码 " * 60], "entities": [], "keywords": []})
        self.assertEqual((loop["facts"], loop.get("runaway")), ([], True))
        src = _repo_file("app/kb_pipeline/vision/vlm.py")
        self.assertIn('if result.get("runaway"):', src)                       # re-checked on cache read: a runaway old entry is recaptioned
        self.assertIn('if data is not None and data.get("runaway"):', src)     # runaway takes the degraded retry
        self.assertIn('data["vlm_runaway"] = True', src)                       # still runaway after retry: keep the non-repetitive facts, confidence low, no caching
        self.assertNotIn('data["facts"] = []', src)                            # 2026-09-10: facts are no longer wiped wholesale


class VlmRunawayStringFieldsTests(unittest.TestCase):
    """2026-09-10, the "automation tasks" screenshot in the product docs KB: the model wrote FACTS into
    summary, repeating "1. task list; 2. task list; ..." up to 516 items, and over 5,000 characters went
    verbatim into the chunk (3,157 tokens) and into the cache. Four places were dropping state: title /
    summary were neither checked for repetition nor capped; numbering and punctuation diluted the share of
    4-character fragments to under half, so the check missed it; visual block tidying truncated by line and
    emptied the whole block when one line exceeded the budget; and chunking added the uncapped visual_summary
    back verbatim."""

    _HEAD = "顶部导航栏包含:首页、日程、任务。;右侧有一个黑色按钮,文字为“新建任务”。;左侧边栏有多个项目列表,包括:"

    def test_repetitive_summary_is_cut_at_the_cycle_and_marked_runaway(self) -> None:
        from kb_pipeline.vision import vlm

        loop = self._HEAD + "".join(f"{i}. 任务列表;" for i in range(1, 400))
        self.assertTrue(vlm._repetitive_prose(loop))
        self.assertFalse(vlm._repetitive(loop[:2000]))                   # the old check that only strips spaces misses it, which is why it slipped through
        data = vlm.normalize_result({"kind": "screenshot", "summary": loop, "facts": ["a"], "entities": [], "keywords": []})
        self.assertTrue(data["runaway"])
        self.assertIn("新建任务", data["summary"])                        # content before the loop is kept
        self.assertNotIn("2. 任务列表", data["summary"])                  # only the first item of the loop is kept
        self.assertLessEqual(len(data["summary"]), vlm._STRING_CAPS["summary"])
        self.assertEqual(data["facts"], ["a"])
        # Whole-sentence repetition (a loop body longer than 4 characters) is caught by the compression ratio
        sentences = "".join(f"第{i}段说明了界面上第{i}个区域的用途与状态。" for i in range(1, 80))
        self.assertTrue(vlm.normalize_result({"summary": sentences}).get("runaway"))
        # Long but not repetitive: capped but not runaway; a normal sentence or two stays as is
        varied = "".join(chr(0x4E00 + (i * 7919) % 2000) for i in range(900))
        long_ok = vlm.normalize_result({"kind": "screenshot", "summary": varied})
        self.assertNotIn("runaway", long_ok)
        self.assertEqual(len(long_ok["summary"]), vlm._STRING_CAPS["summary"])
        short = "一个弹窗界面,用于将当前项目添加到指定团队,列出了多个可选团队。"
        self.assertEqual(vlm.normalize_result({"summary": short})["summary"], short)
        title = vlm.normalize_result({"title": "任务 " * 200})
        self.assertTrue(title["runaway"])
        self.assertLessEqual(len(title["title"]), vlm._STRING_CAPS["title"])

    def test_verbatim_runaway_on_one_line_is_dropped_but_pin_tables_survive(self) -> None:
        from kb_pipeline.vision import vlm

        loop = vlm.normalize_result({"summary": "s", "text_verbatim": "首页 密码 " * 300})
        self.assertEqual((loop["text_verbatim"], loop.get("runaway")), ("", True))
        pins = "\n".join(f"A{i} VSS" if i % 2 else f"A{i} VDD" for i in range(1, 80))
        rows = "\n".join(f"| 参数 {i} | {i * 3} | {i * 7} mV | 典型值 |" for i in range(1, 60))
        for legit in (pins, rows):
            ok = vlm.normalize_result({"summary": "s", "text_verbatim": legit})
            self.assertEqual(ok["text_verbatim"], legit)
            self.assertNotIn("runaway", ok)
        # A single facts item: the old rule stands, and a long item repeating whole sentences counts too
        item = "".join(f"第{i}行显示了第{i}个配置项的当前取值与说明。" for i in range(1, 40))
        self.assertEqual(vlm.normalize_result({"facts": [item]})["facts"], [])

    def test_tidy_keeps_a_truncated_first_line_instead_of_emptying_the_block(self) -> None:
        from kb_pipeline.parsers.visual_blocks import tidy_visual_text
        from kb_pipeline.utils import count_tokens

        out, stats = tidy_visual_text("VISUAL SUMMARY: " + "任务列表;" * 800, kind="screenshot", max_tokens=100)
        self.assertTrue(out.startswith("VISUAL SUMMARY: "))
        self.assertLessEqual(count_tokens(out), 100)
        self.assertGreater(stats["tokens_cut"], 0)
        # With too little budget left, no line is truncated (half a sentence is worse than none); if there is
        # enough, the line that does not fit is cut to the remaining budget and later lines are dropped
        out2, _ = tidy_visual_text("短行一\n" + "长行" * 300 + "\n短行二", max_tokens=20)
        self.assertEqual(out2, "短行一")
        facts = "FACTS: " + "；".join(f"第{i}个按钮位于工具栏右侧,用于打开设置面板" for i in range(1, 80))
        out3, _ = tidy_visual_text("VISUAL SUMMARY: 一个设置界面。\n" + facts + "\nENTITIES: WPS", kind="screenshot", max_tokens=200)
        self.assertTrue(out3.startswith("VISUAL SUMMARY: 一个设置界面。\nFACTS: 第1个按钮"))
        self.assertTrue(out3.endswith("；"))                                   # cut at the full-width semicolon, leaving no half fact
        self.assertNotIn("ENTITIES", out3)
        self.assertLessEqual(count_tokens(out3), 200)

    def test_chunker_does_not_reappend_a_truncated_summary(self) -> None:
        from kb_pipeline.chunking.chunker import text_for_block

        full = "一个很长的摘要," * 50
        b = _block("i1", "VISUAL SUMMARY: " + full[:120], block_type="image", visual_summary=full)
        self.assertEqual(text_for_block(b).count("VISUAL SUMMARY:"), 1)
        self.assertLess(len(text_for_block(b)), 200)
        b2 = _block("i2", "", block_type="image", visual_summary="短摘要")
        self.assertIn("VISUAL SUMMARY: 短摘要", text_for_block(b2))

    def test_cached_runaway_summary_is_recaptioned(self) -> None:
        import hashlib
        from unittest import mock

        from kb_pipeline.vision import vlm

        clean = {"kind": "screenshot", "title": "", "summary": "干净的摘要", "text_verbatim": "", "entities": [],
                 "facts": ["顶部有导航栏"], "keywords": [], "confidence": "high", "decorative": False}

        class _Resp:
            def __init__(self, content):
                self.choices = [SimpleNamespace(message=SimpleNamespace(content=content), finish_reason="stop")]

        class _Fake:
            def __init__(self):
                self.calls = 0
                self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

            def _create(self, **request):
                self.calls += 1
                return _Resp(json.dumps(clean, ensure_ascii=False))

        with tempfile.TemporaryDirectory() as tmp:
            cache = Path(tmp) / "c.json"
            img = Path(tmp) / "x.png"
            img.write_bytes(b"x")
            identity = {"contract_version": vlm.CAPTION_CONTRACT_VERSION, "model_id": "m",
                        "prompt_sha256": hashlib.sha256(b"P").hexdigest(), "temperature": 0.1, "top_p": 0.8,
                        "repetition_penalty": 1.05, "structured": True, "max_pixels": vlm.DEFAULT_MAX_PIXELS}
            loop = self._HEAD + "".join(f"{i}. 任务列表;" for i in range(1, 300))
            cache.write_text(json.dumps({"kind": "screenshot", "summary": loop, "confidence": "low", "_cache": identity},
                                        ensure_ascii=False), encoding="utf-8")
            fake = _Fake()
            with mock.patch.object(vlm, "_client_for", return_value=fake), \
                    mock.patch.object(vlm, "image_data_url", return_value="data:,"):
                out = vlm.caption_image(image_path=img, base_url="http://x", api_key="k", model_id="m", prompt="P", cache_json=cache)
            self.assertEqual(fake.calls, 1)                                 # the old cache entry is judged runaway and recaptioned, not reused
            self.assertEqual((out["summary"], out["vlm_cache_hit"]), ("干净的摘要", False))
            self.assertEqual(json.loads(cache.read_text(encoding="utf-8"))["summary"], "干净的摘要")


class ParserFixRegressionTests(_CodexAudit20260906TestsSupport, _CodexFinalTestsSupport, unittest.TestCase):
    """Regressions for problems found in successive re-checks, health checks and reviews; each test's docstring
    gives the source and the symptom observed at the time."""

    def test_split_long_row_pieces_get_unique_block_ids(self) -> None:  # issue 1
        from kb_pipeline.parsers.native_table import chunk_rows_to_blocks

        long_cell = "需求描述" + "很长的正文内容。" * 120
        rows = [(1, ["编号", "描述"]), (2, ["REQ-1", long_cell]), (3, ["REQ-2", "普通行"])]
        blocks = chunk_rows_to_blocks("xlsx", rows, "Sheet1", max_tokens=120, overlap_tokens=0)
        ids = [b.block_id for b in blocks]
        self.assertEqual(len(ids), len(set(ids)), ids)
        pieces = [b for b in blocks if "-rows-2-2" in b.block_id]
        self.assertGreater(len(pieces), 1)
        self.assertTrue(all(b.block_id.endswith(f"-p{i}") for i, b in enumerate(pieces, start=1)), ids)

    def test_split_long_row_keeps_column_names_and_row_identity(self) -> None:
        """2026-09-13 Codex F01: continuation pieces of a wide table must carry the column names and the row
        identity; empty columns must not vanish silently."""
        from kb_pipeline.chunking.chunker import _split_table_row, split_table_text
        from kb_pipeline.parsers.native_table import split_long_row
        from kb_pipeline.utils import count_tokens

        header = ["提交时间", "姓名", "岗位", "问题一", "问题二"]
        row = ["2026-05-01", "张三", "", "很长的回答" * 30, "另一段很长的回答" * 30]
        pieces = split_long_row(row, 60, header)
        self.assertGreater(len(pieces), 1)
        self.assertTrue(pieces[0].startswith("提交时间: 2026-05-01 | 姓名: 张三"))            # the empty "position" cell is skipped, later cells still know their column
        self.assertTrue(all(p.startswith("提交时间: 2026-05-01") for p in pieces))              # every piece starts with the row identity
        self.assertIn("问题二: 另一段很长的回答", pieces[-1])
        self.assertNotIn("岗位:", " ".join(pieces))
        self.assertTrue(all(count_tokens(p) <= 60 for p in pieces), [count_tokens(p) for p in pieces])   # a long cell alone exceeds the budget: split, every piece within it
        self.assertEqual("".join(p.split("问题一: ", 1)[1] for p in pieces if "问题一: " in p), row[3])
        self.assertEqual(split_long_row(["A值", "", "C值"], 60, None), ["A: A值 | C: C值"])   # without a header, column letters keep the position
        self.assertEqual(split_long_row(["", ""], 8, ["a", "b"]), [""])
        # markdown tables / native table blocks follow the same rule
        md = _split_table_row("| REQ-1 | | " + "描述文字" * 40 + " |", 40, ["编号", "状态", "描述"])
        self.assertGreater(len(md), 1)
        self.assertTrue(all(m.startswith("编号: REQ-1 | 描述: ") for m in md), md)               # no line holding nothing but the row identity
        self.assertEqual("".join(m[len("编号: REQ-1 | 描述: "):] for m in md), "描述文字" * 40)
        text = "SHEET: 需求\nROWS: 9-9\nHEADER: 编号 | 状态 | 描述\nREQ-9 |  | " + "很长的描述。" * 200
        out = split_table_text(text, 120)
        self.assertGreater(len(out), 5)
        self.assertTrue(all(o.splitlines()[:2] == ["SHEET: 需求", "ROWS: 9-9"] for o in out))
        self.assertTrue(all(o.splitlines()[2].startswith("编号: REQ-9 | 描述: 很长的描述。") for o in out))   # every piece carries the row identity and the column name
        self.assertNotIn("HEADER:", "\n".join(out))                                           # pieces that carry column names do not repeat the whole header
        self.assertTrue(all(count_tokens(o) <= 120 for o in out))

    def test_xls_cell_text_dates_and_booleans(self) -> None:  # issue 9
        import xlrd

        from kb_pipeline.parsers.native_table import xls_cell_text

        serial = xlrd.xldate.xldate_from_datetime_tuple((2026, 8, 20, 0, 0, 0), 0)
        self.assertEqual(xls_cell_text(xlrd.XL_CELL_DATE, serial, 0), "2026-08-20")
        with_time = xlrd.xldate.xldate_from_datetime_tuple((2026, 8, 20, 9, 30, 0), 0)
        self.assertEqual(xls_cell_text(xlrd.XL_CELL_DATE, with_time, 0), "2026-08-20 09:30:00")
        self.assertEqual(xls_cell_text(xlrd.XL_CELL_BOOLEAN, 1, 0), "TRUE")
        self.assertEqual(xls_cell_text(xlrd.XL_CELL_ERROR, 42, 0), "")
        self.assertEqual(xls_cell_text(xlrd.XL_CELL_NUMBER, 3.0, 0), "3")

    def test_row_alignment_keeps_interior_holes(self) -> None:  # alignment bonus
        from kb_pipeline.parsers.native_table import row_to_text

        self.assertEqual(row_to_text(["服务器A", "", "已下线"]), "服务器A |  | 已下线")
        self.assertEqual(row_to_text(["a", "b", "", ""]), "a | b")

    def test_cell_newlines_are_flattened_so_one_record_stays_one_line(self) -> None:
        """In-cell newlines ("Session list<LF>@owner", header "WPS Collab<LF>public") used to split one record
        over two or three lines; rows, headers and the column names / values of an over-budget split all
        collapse to one line."""
        from kb_pipeline.parsers.native_table import flat_cell, header_text, row_to_text, split_long_row

        self.assertEqual(flat_cell("会话列表\n@某某"), "会话列表 @某某")
        self.assertEqual(flat_cell("  a \r\n\n b\tc  "), "a b c")
        self.assertEqual(row_to_text(["聊天能力", "会话列表\n@某某", "移除会话", "", "1"]), "聊天能力 | 会话列表 @某某 | 移除会话 |  | 1")
        self.assertEqual(header_text(["支持情况/WPS协作\n公网", "WPS协作\n私网"]), "支持情况/WPS协作 公网 | WPS协作 私网")
        self.assertEqual(split_long_row(["a\nb", "c"], 100, ["x\ny", "z"]), ["x y: a b | z: c"])
        self.assertEqual(row_to_text(["", "", ""]), "")

    def test_gbk_files_keep_their_chinese(self) -> None:  # issue 3
        from kb_pipeline.parsers.common import read_text_smart
        from kb_pipeline.parsers.html_dom import parse_html_dom
        from kb_pipeline.parsers.router import parse_native

        with tempfile.TemporaryDirectory() as tmp:
            txt = Path(tmp) / "老文档.txt"
            txt.write_bytes("这是一段GB编码的中文说明".encode("gb18030"))
            self.assertIn("中文说明", read_text_smart(txt))
            blocks = parse_native(txt)
            self.assertIn("中文说明", blocks[0].text)

            md = Path(tmp) / "老文档.md"
            md.write_bytes("# 标题\n\n中文正文段落".encode("gb18030"))
            md_blocks = parse_native(md)
            self.assertTrue(any("中文正文" in b.text for b in md_blocks))

            page = Path(tmp) / "老页面.html"
            body = "<html><head><meta charset=gb2312><title>目录</title></head><body><p>" + "中文内容段落," * 20 + "</p></body></html>"
            page.write_bytes(body.encode("gb18030"))
            html_blocks = parse_html_dom(page)
            self.assertTrue(any("中文内容段落" in b.text for b in html_blocks))

    def test_xhtml_parses_and_fallback_strips_scripts(self) -> None:  # issue 10
        from kb_pipeline.parsers.html_dom import _raw_fallback, _safe_id, parse_html_dom

        with tempfile.TemporaryDirectory() as tmp:
            xhtml = Path(tmp) / "doc.html"
            xhtml.write_text(
                '<?xml version="1.0" encoding="utf-8"?>\n'
                "<html><head><title>规格书</title></head><body><p>"
                + "正文内容。" * 40
                + "</p></body></html>",
                encoding="utf-8",
            )
            blocks = parse_html_dom(xhtml)
            self.assertNotEqual(blocks[0].block_id, "html-raw-fallback")
            self.assertTrue(any("正文内容" in b.text for b in blocks))

        fallback = _raw_fallback(
            "<html><script>var秘密=1;alert(1)</script><body>可见文字 &amp; 实体</body></html>",
            title="t", reason="r",
        )[0]
        self.assertIn("可见文字 & 实体", fallback.text)
        self.assertNotIn("alert", fallback.text)
        self.assertNotIn("<body>", fallback.text)

        a, b = _safe_id("简介"), _safe_id("详情")
        self.assertNotEqual(a, b)
        self.assertTrue(a.startswith("id-") and b.startswith("id-"))
        self.assertEqual(_safe_id("page-2"), "page-2")

    def test_nested_table_rows_not_duplicated(self) -> None:  # issue 10
        from lxml import html as lxml_html

        from kb_pipeline.parsers.html_dom import _table_markdown

        table = lxml_html.fragment_fromstring(
            "<table><tr><th>外层</th></tr>"
            "<tr><td><table><tr><td>内层A</td></tr><tr><td>内层B</td></tr></table></td></tr></table>"
        )
        markdown = _table_markdown(table)
        self.assertEqual(markdown.count("内层A"), 1)   # inlined once via the outer cell
        self.assertNotIn("| 内层B |\n| 内层B |", markdown)
        self.assertEqual(len([l for l in markdown.splitlines() if l.startswith("|")]), 3)  # header+sep+1 row

    def test_unidentified_image_passthrough_is_capped(self) -> None:  # issue 12
        from kb_pipeline.vision import images as images_module

        with tempfile.TemporaryDirectory() as tmp:
            small = Path(tmp) / "small.png"; small.write_bytes(b"not an image")
            data, mime = images_module.load_image_for_model(small)
            self.assertEqual(data, b"not an image")
            with patch.object(images_module, "MAX_UNIDENTIFIED_PASSTHROUGH_BYTES", 4):
                with self.assertRaises(RuntimeError):
                    images_module.load_image_for_model(small)

    def test_decompression_bomb_is_skipped_before_the_vlm(self) -> None:  # issue 12
        from PIL import Image

        from kb_pipeline.models import ParsedBlock
        from kb_pipeline.parsers import visual_blocks

        with tempfile.TemporaryDirectory() as tmp:
            img = Path(tmp) / "big.png"
            Image.new("RGB", (600, 600)).save(img)
            block = ParsedBlock(parser="p", parser_profile="p", doc_type="pdf", block_type="image",
                                text="", block_id="b", visual_ref=str(img))
            old_limit = Image.MAX_IMAGE_PIXELS
            Image.MAX_IMAGE_PIXELS = 10_000  # 600x600 is now a "bomb"
            try:
                def must_not_run(jobs, **kwargs):
                    raise AssertionError("bomb reached the VLM")

                with patch.object(visual_blocks, "caption_images_parallel", must_not_run):
                    visual_blocks.enrich_blocks_with_vlm([block], base_url="u", api_key="k", model_id="m",
                                                         cache_dir=Path(tmp), filter_decorative=False)
            finally:
                Image.MAX_IMAGE_PIXELS = old_limit
            self.assertEqual(block.metadata["vlm_status"], "skipped")
            self.assertEqual(block.metadata["vlm_skip_reason"], "unreadable_image")

    def test_exif_rotated_jpeg_is_transposed(self) -> None:  # bonus with issue 12
        from io import BytesIO

        from PIL import Image

        from kb_pipeline.vision.images import load_image_for_model

        with tempfile.TemporaryDirectory() as tmp:
            src = Image.new("RGB", (40, 20), (10, 20, 30))
            exif = Image.Exif(); exif[0x0112] = 6  # rotate 90 CW to display
            path = Path(tmp) / "rotated.jpg"
            src.save(path, format="JPEG", exif=exif)
            data, mime = load_image_for_model(path)
            self.assertEqual(mime, "image/jpeg")
            with Image.open(BytesIO(data)) as out:
                self.assertEqual((out.width, out.height), (20, 40))  # transposed, not passthrough

    def test_extract_markdown_prefers_md_content_deterministically(self) -> None:  # issue 16
        from kb_pipeline.parsers.service_clients import extract_markdown

        payload = {"text": "wrong", "md_content": "right", "markdown": "also wrong"}
        for _ in range(5):
            self.assertEqual(extract_markdown(payload), "right")

    def test_compose_prompt_keeps_the_field_guide_contract(self) -> None:
        from kb_pipeline.vision import vlm

        default = vlm.compose_prompt("figure", None)
        self.assertEqual(default, vlm.DEFAULT_PROMPT + "\n\n" + vlm._FIELD_GUIDE)
        self.assertEqual(default, vlm.compose_prompt("docx_image", None))  # one default for every route
        custom = vlm.compose_prompt("figure", "半导体图片,重点识别时序图与针脚。")
        self.assertTrue(custom.startswith("半导体图片"))
        self.assertNotIn(vlm.DEFAULT_PROMPT, custom)       # custom replaces the paragraph
        self.assertIn("kind: screenshot", custom)          # field guide always appended
        img = vlm.compose_prompt("image_file", "自定义指令", filename="架构图.png")
        self.assertIn("架构图.png", img)
        self.assertIn("kind: screenshot", img)
        default_img = vlm.compose_prompt("image_file", None, filename="架构图.png")
        self.assertIn("架构图.png", default_img)
        self.assertIn(vlm.DEFAULT_PROMPT, default_img)     # standalone images also get the shared default paragraph
        with self.assertRaises(ValueError):
            vlm.compose_prompt("image_file", None)         # a standalone image must come with a file name

    def test_caption_parse_failure_defense_layers(self) -> None:
        # 1) text_verbatim must come last in the schema: truncation then only hurts the transcription's tail
        self.assertEqual(list(vlm.RESULT_SCHEMA["properties"])[-1], "text_verbatim")
        self.assertEqual(vlm.RESULT_SCHEMA["required"][-1], "text_verbatim")

        # 2) Truncated guided-decoding output is a valid JSON prefix and can be completed deterministically
        full = {"kind": "diagram", "title": "T", "summary": "S", "entities": ["e"],
                "facts": ["f1", "f2"], "keywords": ["k"], "confidence": "high",
                "text_verbatim": "line1\nline2"}
        text = json.dumps(full, ensure_ascii=False)
        cut = text[: text.index('"text_verbatim"') + len('"text_verbatim": "li')]
        fixed = vlm.complete_json_prefix(cut)
        self.assertIsNotNone(fixed)
        obj = json.loads(fixed)
        self.assertEqual(obj["summary"], "S")
        self.assertEqual(obj["facts"], ["f1", "f2"])
        parsed = vlm.parse_structured(cut)
        self.assertEqual(parsed["kind"], "diagram")
        # Truncation right after a key's colon: the dangling key is dropped, the other fields are recovered
        cut2 = text[: text.index('"text_verbatim"') + len('"text_verbatim":')]
        self.assertEqual(json.loads(vlm.complete_json_prefix(cut2))["summary"], "S")
        # Complete JSON passes through unchanged
        self.assertEqual(vlm.complete_json_prefix(text), text)

        # 3) Degenerate repeated lines are collapsed in the normalize stage; legitimate spaced repeats are kept
        loop = {"summary": "ok", "confidence": "high",
                "text_verbatim": "\n".join(["High Impedance"] * 50) + "\nEND\nVSS\nA1\nVSS"}
        norm = vlm.normalize_result(loop)
        self.assertEqual(norm["text_verbatim"], "High Impedance\nEND\nVSS\nA1\nVSS")

        # 4) When nothing parses, never stuff the raw text into summary; salvage the summary field when possible
        junk_prefix = '{"kind": "diagram", "summary": "该图展示时序关系", "text_verbatim": "aaa'
        s1 = vlm.parse_jsonish(junk_prefix)
        self.assertEqual(s1["kind"], "diagram")          # prefix completion recovers it directly
        self.assertEqual(s1["summary"], "该图展示时序关系")
        hopeless = 'not json at all "summary": "抢救出的概括" trailing junk'
        s2 = vlm.parse_jsonish(hopeless)
        self.assertEqual(s2["summary"], "抢救出的概括")
        s3 = vlm.parse_jsonish("彻底没有结构的输出")
        self.assertEqual(s3["summary"], "")              # better empty than garbage

        # 5) Requests carry a mild repetition_penalty by default, and it is part of the cache identity
        # (changing it means a rerun)
        import inspect
        sig = inspect.signature(vlm.caption_image)
        self.assertAlmostEqual(sig.parameters["repetition_penalty"].default, 1.05)

    def test_vlm_failures_are_counted_so_the_job_can_retry(self) -> None:
        """H5: when the VLM is down, each image's exception was swallowed one image at a time and the document
        was still stored as "parsed successfully"; with content_version unchanged the scan judged it
        unchanged, and the missing captions were never filled in. The parse flow must see the failure count
        to decide on a retry."""
        from kb_pipeline.parsers.visual_blocks import vlm_failure_summary

        def block(block_type: str, *, visual_ref: str | None, status: str | None) -> ParsedBlock:
            return ParsedBlock(
                parser="test", parser_profile="p", doc_type="pdf", block_type=block_type,
                text="", block_id=f"b-{block_type}-{status}", visual_ref=visual_ref,
                metadata={"vlm_status": status} if status else {},
            )

        def visual(status: str) -> ParsedBlock:
            return block("image", visual_ref="x.png", status=status)

        text_block = block("text", visual_ref=None, status=None)
        blocks = [text_block, visual("success"), visual("failed"), visual("failed")]
        summary = vlm_failure_summary(blocks)
        self.assertEqual(summary["visual_blocks"], 3)   # only blocks that went through the VLM are counted
        self.assertEqual(summary["vlm_failed"], 2)
        self.assertEqual(vlm_failure_summary([text_block]), {"visual_blocks": 0, "vlm_failed": 0})

    def test_markdown_setext_headings_and_long_fences(self) -> None:
        """B9: setext headings (=== / --- on the next line) are recognised too; ``` and # inside a
        four-backtick fence are code."""
        from kb_pipeline.parsers.common import markdown_to_blocks

        md = ("Title\n=====\n\nintro\n\nPart\n---\n\nbody\n\n````md\n```\n# not a heading\n```\n````\n\n"
              "line one\nline two\n---\ntail\n")
        blocks = markdown_to_blocks(md, parser="native", parser_profile="p", doc_type="markdown")
        self.assertEqual([(b.title, b.metadata["section_path"]) for b in blocks],
                         [("Title", ["Title"]), ("Part", ["Title", "Part"])])
        self.assertEqual(blocks[0].text, "intro")
        self.assertIn("# not a heading", blocks[1].text)
        self.assertIn("line one\nline two\ntail", blocks[1].text)
        self.assertNotIn("---", blocks[1].text)          # --- after a paragraph is still a rule, not a heading
        self.assertNotIn("=====", blocks[0].text)

    def test_f02_glued_table_rows_are_flagged_and_only_verified_splits_are_applied(self) -> None:
        """F02: MinerU glued the stacked 75/75/57 mA into 757557 / mAmAmA. Detection flags per row (unit
        repeated n=3 + a long digit string = high confidence); the repair is applied only when the VLM's split
        reassembles into the original string; if it does not, or the call fails, the original text and the
        flag are kept."""
        from kb_pipeline.parsers import table_repair
        from kb_pipeline.parsers.common import grid_to_markdown, html_table_to_grid, html_table_to_markdown
        from kb_pipeline.parsers.table_check import (detect_table_ambiguity, expand_glued_rows, table_ambiguity_flags,
                                                    table_ambiguity_summary, unit_repeat, unrepaired_glued_values)

        self.assertEqual((unit_repeat("mAmAmA"), unit_repeat("°C°C"), unit_repeat("mA"), unit_repeat("Vpp")), (3, 2, 1, 1))
        self.assertEqual((unit_repeat("mm"), unit_repeat("**"), unit_repeat("输出数据" * 6), unit_repeat("pFpF")), (1, 1, 1, 2))
        # Capacitance table: three test conditions in one row is normal; with no glued values / units nothing
        # is flagged (a first-round false positive on the real box)
        cap = ("<table><tr><td>参数</td><td>说明</td><td>测试条件</td><td>最大值</td><td>单位</td></tr>"
               "<tr><td>$C_{IN}$</td><td>输入电容</td><td>$T_A = 25°C$ 、f=1 MHz、 $V_{CC} = V_{CC(typ)}$</td><td>10</td><td>pF</td></tr></table>")
        self.assertEqual(table_ambiguity_flags(cap), [])
        rev = ("<table><tr><td>版本</td><td>ECN</td><td>日期</td></tr><tr><td>**</td><td>5148011</td><td>02/23/2016</td></tr>"
               "<tr><td>*A</td><td>5290011</td><td>03/01/2017</td></tr></table>")
        self.assertEqual(table_ambiguity_flags(rev), [])
        flags = table_ambiguity_flags(self.NVSRAM_ROW_HTML)
        kinds = {(f["row"], f["col"], f["kind"], f["n"], f["confidence"]) for f in flags}
        self.assertEqual(kinds, {(2, 2, "multi_condition", 4, "high"), (2, 5, "glued_values", 3, "high"), (2, 6, "glued_units", 3, "high")})
        self.assertTrue(all(f["key"].startswith("$I_{CC1}$") for f in flags))
        self.assertEqual(unrepaired_glued_values(flags, None), ["757557"])
        self.assertEqual(unrepaired_glued_values(flags, {"status": "verified", "rows": [2]}), [])
        # A normal table is not flagged; nor is a column of 5-digit order codes; only a digit string clearly
        # longer than the rest of its column gets low confidence
        self.assertEqual(table_ambiguity_flags("<table><tr><td>a</td><td>b</td></tr><tr><td>1</td><td>2.5</td></tr></table>"), [])
        codes = "<table><tr><td>码</td><td>值</td></tr><tr><td>A</td><td>12345</td></tr><tr><td>B</td><td>23456</td></tr><tr><td>C</td><td>34567</td></tr></table>"
        self.assertEqual(table_ambiguity_flags(codes), [])
        odd = "<table><tr><td>参数</td><td>最大值</td></tr><tr><td>a</td><td>75</td></tr><tr><td>b</td><td>57</td></tr><tr><td>c</td><td>757557</td></tr></table>"
        self.assertEqual([(f["kind"], f["confidence"]) for f in table_ambiguity_flags(odd)], [("glued_values", "low")])
        grid = html_table_to_grid(self.NVSRAM_ROW_HTML)
        self.assertEqual(html_table_to_markdown(self.NVSRAM_ROW_HTML), grid_to_markdown(grid))
        expanded = expand_glued_rows(grid, {2: {5: ["75", "75", "57"], 6: ["mA", "mA", "mA"], 2: ["t_RC = 20 ns", "t_RC = 25 ns", "t_RC = 45 ns"]}})
        self.assertEqual(len(expanded), len(grid) + 2)
        self.assertEqual([r[5] for r in expanded[2:5]], ["75", "75", "57"])
        self.assertEqual([r[2] for r in expanded[2:5]], ["t_RC = 20 ns", "t_RC = 25 ns", "t_RC = 45 ns"])
        self.assertEqual(expanded[2][0], expanded[4][0])                            # unsplit columns are copied as is
        self.assertEqual(expand_glued_rows(grid, {2: {5: ["75", "75"], 6: ["mA"] * 3}}), grid)   # inconsistent piece counts: no expansion
        self.assertTrue(table_repair.verify_split("757557", ["75", "75", "57"]))
        self.assertTrue(table_repair.verify_split("mAmAmA", ["mA", "mA", "mA"]))
        self.assertFalse(table_repair.verify_split("757557", ["75", "75", "75"]))
        self.assertFalse(table_repair.verify_split("757557", ["757557"]))

        def block(tmp):
            crop = Path(tmp) / "t.jpg"; crop.write_bytes(b"jpg")
            return ParsedBlock(parser="mineru", parser_profile="p", doc_type="pdf", block_type="table", block_id="t1",
                               text=html_table_to_markdown(self.NVSRAM_ROW_HTML) + "\nFOOTNOTE: 注 8",
                               table_markdown=html_table_to_markdown(self.NVSRAM_ROW_HTML), visual_ref=str(crop),
                               metadata={"table_flags": flags, "source_item": {"table_body": self.NVSRAM_ROW_HTML}})

        good = [{"key": "$I_{CC1}$", "cells": [["I_CC1"], ["平均电流 V_CC"], ["t_RC = 20 ns", "t_RC = 25 ns", "t_RC = 45 ns"],
                                                ["-"], ["-"], ["75", "75", "57"], ["mA", "mA", "mA"]]}]
        with tempfile.TemporaryDirectory() as tmp:
            b = block(tmp)
            calls: list[dict] = []

            def fake(image, **kw):
                calls.append(kw); return good
            stats = table_repair.repair_ambiguous_tables([b], base_url="u", api_key="k", model_id="m", cache_dir=Path(tmp), transcribe=fake)
            self.assertEqual((stats["flagged"], stats["verified"]), (1, 1))
            self.assertEqual(calls[0]["row_keys"], ["$I_{CC1}$"])
            self.assertEqual(b.metadata["table_repair"]["status"], "verified")
            self.assertEqual((b.metadata["table_repair"]["rows"], b.metadata["table_repair"]["n"]), ([2], 3))
            self.assertNotIn("757557", b.text)
            self.assertIn("| t_RC = 45 ns | - | - | 57 | mA |", b.text)
            self.assertTrue(b.text.endswith("FOOTNOTE: 注 8"))                     # the note after the table is kept
            self.assertEqual(b.table_markdown, b.text)
            self.assertTrue(all(f.get("repaired") for f in b.metadata["table_flags"]))
            self.assertEqual(table_ambiguity_summary([b]), {"flagged": 1, "verified": 1, "unverified": 0})
            # Does not reassemble into the original string: no rewrite, flag kept
            b2 = block(tmp)
            bad = [{"key": "$I_{CC1}$", "cells": [["I_CC1"], ["x"], ["c"], ["-"], ["-"], ["75", "75", "75"], ["mA", "mA", "mA"]]}]
            stats = table_repair.repair_ambiguous_tables([b2], base_url="u", api_key="k", model_id="m", cache_dir=Path(tmp), transcribe=lambda *a, **k: bad)
            self.assertEqual((stats["verified"], stats["unverified"]), (0, 1))
            self.assertEqual(b2.metadata["table_repair"]["status"], "unverified")
            self.assertIn("757557", b2.text)
            # VLM call failed: parsing is unaffected, recorded as failed
            b3 = block(tmp)

            def boom(*a, **k):
                raise RuntimeError("vlm down")
            stats = table_repair.repair_ambiguous_tables([b3], base_url="u", api_key="k", model_id="m", cache_dir=Path(tmp), transcribe=boom)
            self.assertEqual((stats["failed"], b3.metadata["table_repair"]["status"]), (1, "failed"))
            self.assertIn("757557", b3.text)
        self.assertEqual(table_repair._parse_rows('```json\n{"rows": [{"key": "k", "cells": [["a", "b"], "c", []]}]}\n```'),
                         [{"key": "k", "cells": [["a", "b"], ["c"], []]}])
        # The real-box VLM shape: first cell missing, the dash merged into the value column, conditions in <sub>;
        # matching by content still verifies n=3
        real = [{"key": "I_{CC1}", "cells": [["平均电流 V<sub>CC</sub>", "t<sub>RC</sub> = 20 ns", "t<sub>RC</sub> = 25 ns", "t<sub>RC</sub> = 45 ns",
                                               "无输出负载下取得的值（I<sub>OUT</sub> = 0 mA）"], ["—", "—", "75", "75", "57"], ["mA", "mA", "mA", "mA", "mA"]]}]
        repairs, reasons = table_repair.plan_repairs(grid, flags, real)
        self.assertEqual(reasons, [])
        self.assertEqual(repairs[2][5], ["75", "75", "57"])
        self.assertEqual(repairs[2][6], ["mA", "mA", "mA"])
        self.assertEqual(repairs[2][2], ["t_RC = 20 ns", "t_RC = 25 ns", "t_RC = 45 ns"])       # <sub> converted to the manual's own underscore notation
        self.assertEqual(table_repair.clean_part("V<sub>CC</sub> + 0.5 V<sup>2</sup> <b>x</b>"), "V_CC + 0.5 V^2 x")
        self.assertNotIn(0, repairs[2])                                            # the symbol cell is not split
        with tempfile.TemporaryDirectory() as tmp:
            b4 = block(tmp)
            seen: list[dict] = []
            table_repair.repair_ambiguous_tables([b4], base_url="u", api_key="k", model_id="m", cache_dir=Path(tmp),
                                                 transcribe=lambda image, **kw: (seen.append(kw), real)[1])
            self.assertEqual(b4.metadata["table_repair"]["status"], "verified")
            self.assertEqual((seen[0]["width"], seen[0]["header"][:3]), (7, ["参数", "说明", "测试条件"]))
            self.assertIn("| t_RC = 45 ns | - | - | 57 | mA |", b4.text)
            self.assertNotIn("<sub>", b4.text)

    def test_visual_correction_only_touches_the_matched_span(self) -> None:
        from kb_pipeline.parsers.visual_blocks import reconcile_visual_facts

        # Two identical numbers in one item: only the voltage one is wrong. The old code replaced the first 10
        # in the whole sentence, corrupting the correct current
        out, _, conflicts = reconcile_visual_facts(["电流 10,电压 10"], "电流 10\n电压 20")
        self.assertEqual(out, ["电流 10,电压 20"])
        self.assertEqual([(c["label"], c["text_value"], c["model_value"]) for c in conflicts], [("电压", "20", "10")])
        # "current" is a substring of "leakage current": exact equality wins and the label listed first no
        # longer grabs it
        self.assertEqual(reconcile_visual_facts(["漏电流 0.2"], "电流 10\n漏电流 0.2")[::2], (["漏电流 0.2"], []))
        out3, _, conf3 = reconcile_visual_facts(["漏电流 0.5"], "电流 10\n漏电流 0.2")
        self.assertEqual((out3, [(c["label"], c["text_value"]) for c in conf3]), (["漏电流 0.2"], [("漏电流", "0.2")]))
        # Containment matches several candidates (input current / output current): ambiguous, left alone
        self.assertEqual(reconcile_visual_facts(["电流 7"], "输入电流 10\n输出电流 5")[::2], (["电流 7"], []))
        # Numbers with different units cannot be compared: 0.1 V and 100 mV are two spellings of one value,
        # so the number is left and no conflict is recorded; with the same unit it is corrected as usual
        self.assertEqual(reconcile_visual_facts(["电压 100 mV"], "电压 0.1 V")[::2], (["电压 100 mV"], []))
        out6, _, conf6 = reconcile_visual_facts(["电压 3.0 V"], "电压 3.3 V")
        self.assertEqual((out6, [c["text_value"] for c in conf6]), (["电压 3.3 V"], ["3.3"]))
        # The same number appears twice in the summary; only the occurrence belonging to that metric changes
        _, summary, _ = reconcile_visual_facts([], "综合风险指数 34", "同龄平均 31;显示当前风险值为31")
        self.assertEqual(summary, "同龄平均 31;显示当前风险值为34")

    def test_a_runaway_transcription_does_not_stall_the_parse(self) -> None:
        """A few thousand characters repeated by the model carry neither punctuation nor digits, and backtracking of
        the "label value" pattern on such a segment grows with the cube of its length: 1,600 characters take 27
        seconds, so one picture could hold a whole document in the image description stage (2026-09-30, a manual
        with 166 pictures). A segment longer than SEGMENT_MAX_CHARS is not a reading and is skipped; short segments
        are compared as before."""
        import time

        from kb_pipeline.parsers.visual_blocks import SEGMENT_MAX_CHARS, _label_num_spans, reconcile_visual_facts

        runaway = "用户管理 " * 4000
        started = time.monotonic()
        self.assertEqual(_label_num_spans(runaway, max_numbers=3), [])
        self.assertEqual(_label_num_spans(runaway + "设备 12", single_only=True), [])
        self.assertEqual(_label_num_spans(runaway), [])
        out, _, conflicts = reconcile_visual_facts(["总分 95"], runaway + "\n总分 98", runaway)
        self.assertLess(time.monotonic() - started, 2.0)
        self.assertEqual((out, conflicts[0]["text_value"]), (["总分 98"], "98"))      # the runaway segment is skipped, the real reading next to it still applies
        long_label = "甲" * (SEGMENT_MAX_CHARS - 4) + " 12"
        self.assertEqual(len(_label_num_spans(long_label)), 1)
        self.assertEqual(_label_num_spans("甲" * SEGMENT_MAX_CHARS + " 12"), [])

    def test_same_label_with_several_values_is_left_alone(self) -> None:
        from kb_pipeline.parsers.visual_blocks import reconcile_visual_facts

        self.assertEqual(reconcile_visual_facts(["电压 5 V"], "电压 3.3 V\n电压 5 V")[::2], (["电压 5 V"], []))
        # 2026-09-13 Codex F06: a flat table's full header row + the first id is not a "label - value" pair; a
        # number the model read correctly must not be changed to 1
        table = ("商品编号 商品名称 商品种类 销售单价 进货单价 登记日期 0001 T恤衫 衣服 1000 500 2009-09-20 "
                 "0002 打孔器 办公用品 500 320 2009-09-11")
        keep = ["商品编号0002对应的商品名称是打孔器,销售单价500,进货单价320。"]
        self.assertEqual(reconcile_visual_facts(keep, table), (keep, "", []))
        self.assertEqual(reconcile_visual_facts(["销售单价 500"], "商品编号 销售单价 进货单价 1")[::2], (["销售单价 500"], []))   # a label of three Chinese words is a header
        self.assertEqual(reconcile_visual_facts(["编号 5"], "编号 0007")[::2], (["编号 5"], []))                             # a leading zero marks an id, not a reading
        self.assertEqual(reconcile_visual_facts(["心率 72"], "心率 66 血压 118 76 体重 62")[::2], (["心率 72"], []))          # four numbers in one segment: a table row, not readings
        out7, _, conf7 = reconcile_visual_facts(["心率 72"], "心率 66 血压 118")                                            # every number has its own label: trusted
        self.assertEqual((out7, conf7[0]["text_value"]), (["心率 66"], "66"))
        # 2026-09-14 Codex R01: citation numbers / numbers inside names / incompletely split table rows are
        # not readings
        keep2 = ["五年生存率 90.1%", "I 期五年生存率为 90.1%"]
        self.assertEqual(reconcile_visual_facts(keep2, "肠癌五年生存率[1-2]\n90.1% 72.6% 53.8% 10.4%")[::2], (keep2, []))
        self.assertEqual(reconcile_visual_facts(["五年生存率 90.1%"], "肠癌五年生存率[1]")[::2], (["五年生存率 90.1%"], []))
        self.assertEqual(reconcile_visual_facts(["Price 20"], "Product Price Stock 1 20 50")[::2], (["Price 20"], []))
        self.assertEqual(reconcile_visual_facts(["MACD 12"], "缩量 MACD(8,17,9)")[::2], (["MACD 12"], []))
        self.assertEqual(reconcile_visual_facts(["安装量 12"], "安装量 Top 8")[::2], (["安装量 12"], []))
        self.assertEqual(reconcile_visual_facts(["总分 95"], "总分(满分100)")[::2], (["总分 95"], []))
        out8, _, conf8 = reconcile_visual_facts(["总分 95"], "总分 98\n注[1]")                                              # a citation line in its own segment does not affect the real reading
        self.assertEqual((out8, conf8[0]["text_value"]), (["总分 98"], "98"))
        self.assertEqual(reconcile_visual_facts(["电压 4 V"], "电压 3.3 V\n电压 5 V")[::2], (["电压 4 V"], []))
        out, _, conf = reconcile_visual_facts(["电压 4 V"], "电压 5 V\n电压 5 V")            # a repeated identical value is safe to use
        self.assertEqual((out, [c["text_value"] for c in conf]), (["电压 5 V"], ["5"]))

    def test_csv_row_numbers_survive_blank_lines(self) -> None:
        from kb_pipeline.parsers.native_table import parse_native_table, read_csv_rows_numbered

        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "t.csv"
            p.write_text("名称,数量\n\n设备,42\n", encoding="utf-8")
            self.assertEqual(read_csv_rows_numbered(p), [(1, ["名称", "数量"]), (3, ["设备", "42"])])
            blocks = parse_native_table(p, max_tokens=400, overlap_tokens=0)
            data = [b for b in blocks if not (b.metadata or {}).get("summary")]
            self.assertEqual((data[0].row_start, data[0].row_end), (3, 3))                   # the data sits on line 3 of the source file, not on filtered line 2

    def test_code_suffixes_fall_back_to_text_when_syntax_parsing_is_unavailable(self) -> None:
        from unittest import mock

        from kb_pipeline.parsers import code_symbols, router

        self.assertIn(".mjs", code_symbols.LANGUAGE_BY_SUFFIX)
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "m.mjs"
            p.write_text("export const a = 1;\n", encoding="utf-8")
            with mock.patch.object(code_symbols, "parse_code", return_value=None):
                blocks = router.parse_native(p)
            self.assertTrue(blocks)


class PdfTextLayerTests(unittest.TestCase):
    """Whole pages lose their CJK text when MinerU cannot render non-embedded Chinese fonts (spot check: an ECG report
    kept only digits on 6 pages): compare against the text layer page by page, confirm through the fonts, render and
    resend only the lost pages, splice them back, mark pages still lost as degraded."""

    @staticmethod
    def _block(page, text, block_type="text", bbox=None, **meta):
        from kb_pipeline.models import ParsedBlock
        return ParsedBlock(parser="mineru", parser_profile="p", doc_type="pdf", block_type=block_type, text=text,
                           block_id=f"b-{page}-{abs(hash(text)) % 1000}", page_idx=page, bbox=bbox, metadata=dict(meta))

    def test_text_layer_and_rendering_run_on_a_real_pdf(self) -> None:
        """Text-layer counts, font checks and page rendering run on the real PyMuPDF: the other tests stub it, so a
        broken import would be swallowed by "skip when not installed"."""
        try:
            import pymupdf
        except ImportError:
            self.skipTest("PyMuPDF not installed")
        from kb_pipeline.parsers import pdf_textlayer as tl

        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "a.pdf"
            with pymupdf.open() as doc:
                doc.new_page(width=300, height=200).insert_text(          # china-s is a non-embedded Chinese font
                    (20, 40), "窦性心律 未见明显异常 建议定期复查 心电图报告结论正常范围", fontname="china-s", fontsize=9)
                doc.new_page(width=300, height=200).insert_text((20, 40), "plain page 12 34", fontname="helv", fontsize=9)
                doc.save(str(source))
            self.assertEqual(tl.text_layer_cjk_by_page(source), {1: 27, 2: 0})
            self.assertEqual(tl.page_details(source, [1, 2, 9], {}), {1: (27, 27), 2: (0, 0)})      # out-of-range page numbers are ignored
            self.assertEqual(tl.page_details(source, [1], {1: [[0, 0, 1000, 1000]]}), {1: (0, 0)})   # text inside figure regions does not count
            rendered = tl.render_pages(source, [1], Path(tmp) / "rendered" / "a.pdf")
            with pymupdf.open(str(rendered)) as doc:
                self.assertEqual(doc.page_count, 1)
                self.assertEqual((doc[0].rect.width, doc[0].rect.height), (300.0, 200.0))
                self.assertEqual(len(doc[0].get_images()), 1)
            self.assertIsNone(tl.text_layer_cjk_by_page(Path(tmp) / "missing.pdf"))

    def test_candidate_pages_compare_layer_with_blocks_and_skip_figures(self) -> None:
        from kb_pipeline.parsers.pdf_textlayer import blocks_cjk_by_page, candidate_pages, figure_boxes_by_page

        blocks = [self._block(0, "心电图报告 结论正常"), self._block(1, "12 34 56"), self._block(2, "图里的汉字不算", "chart", bbox=[100, 100, 900, 600]),
                  self._block(3, "英文页 abc")]
        got = blocks_cjk_by_page(blocks)
        self.assertEqual(got, {0: 9, 1: 0, 3: 3})                                  # figure blocks do not count
        self.assertEqual(candidate_pages({0: 9, 1: 40, 2: 30, 3: 3, 4: 0}, got), [1, 2])   # page 3 has too little Chinese to count
        self.assertEqual(figure_boxes_by_page(blocks), {2: [[100.0, 100.0, 900.0, 600.0]]})

    def test_repair_renders_only_confirmed_pages_and_splices_them_back(self) -> None:
        from kb_pipeline.parsers import pdf_textlayer as tl

        original = [self._block(0, "第一章 总则", section_path=["第一章"]), self._block(1, "1 2 3", section_path=["第一章"]),
                    self._block(2, "4 5 6"), self._block(3, "第二章 正文照旧", section_path=["第二章"])]
        layer = {0: 5, 1: 40, 2: 40, 3: 8}
        rendered_pages = []
        rendered_blocks = [self._block(1, "第一页救回来了 窦性心律 未见明显异常 建议定期复查"), self._block(2, "7 8 9")]   # small-PDF pages are 1-based
        with patch.object(tl, "render_pages", side_effect=lambda src, pages, out, **kw: rendered_pages.append(list(pages)) or out):
            merged, info = tl.repair_lost_text_layer(
                original, path=Path("/x.pdf"), rendered_path=Path("/tmp/none/x.pdf"), layer=layer,
                detail=lambda pages, boxes: {1: (40, 40), 2: (40, 40)},     # both pages: CJK drawn with non-embedded fonts
                parse_rendered=lambda p: rendered_blocks)
        self.assertEqual(rendered_pages, [[1, 2]])                               # only the two confirmed pages are rendered
        self.assertEqual((info["candidates"], info["lost_before"], info["recovered"], info["lost"]), ([1, 2], [1, 2], [1], [2]))
        self.assertEqual([b.page_idx for b in merged], [0, 1, 2, 3])
        page1 = merged[1]
        self.assertTrue(page1.block_id.startswith("rendered-") and page1.metadata["rendered_page"])
        self.assertEqual(page1.metadata["section_path"], ["第一章"])              # inherits the preceding block's section
        self.assertIs(merged[2], original[2])                                    # a page that did not recover keeps its original blocks
        self.assertIs(merged[3], original[3])

    def test_repair_skips_pages_whose_chinese_is_embedded_or_inside_figures(self) -> None:
        from kb_pipeline.parsers import pdf_textlayer as tl

        original = [self._block(0, "封面"), self._block(1, "1 2 3")]
        with patch.object(tl, "render_pages") as render:
            blocks, info = tl.repair_lost_text_layer(original, path=Path("/x.pdf"), rendered_path=Path("/tmp/none"),
                                                     layer={0: 2, 1: 40}, detail=lambda pages, boxes: {1: (40, 0)},   # all embedded fonts
                                                     parse_rendered=lambda p: [])
            self.assertIs(blocks, original)
            self.assertEqual((info["candidates"], info["lost"]), ([1], []))
            render.assert_not_called()
            # the second parse blows up: the main flow is unaffected, the original result is used, the page stays lost
            def boom(p):
                raise RuntimeError("mineru down")
            blocks, info = tl.repair_lost_text_layer(original, path=Path("/x.pdf"), rendered_path=Path("/tmp/none"),
                                                     layer={0: 2, 1: 40}, detail=lambda pages, boxes: {1: (40, 40)}, parse_rendered=boom)
            self.assertIs(blocks, original)
            self.assertEqual(info["lost"], [1])
            self.assertIn("mineru down", info["error"])
        # no text-layer information (PyMuPDF missing / unreadable): the whole check is skipped
        blocks, info = tl.repair_lost_text_layer(original, path=Path("/x.pdf"), rendered_path=Path("/tmp/none"), layer={},
                                                 parse_rendered=lambda p: [])
        self.assertIs(blocks, original)
        self.assertFalse(info["checked"])

    def test_degraded_mark_reaches_the_chunk_payload_and_search_row(self) -> None:
        from kb_pipeline.models import UnifiedChunk
        from kb_pipeline.parsers.pdf_textlayer import mark_degraded_pages
        from kb_pipeline.pipeline.parse_job import _payload_for_chunk
        from kb_search.evidence import source_row

        merged = self._block(2, "跨页正文", page_start=2, page_end=3)
        single = self._block(5, "别页")
        table = self._block(3, "1 | 2", "table")
        self.assertEqual(mark_degraded_pages([merged, single, table], [3]), 2)      # merged blocks match by page range; page 5 untouched
        self.assertEqual(merged.metadata["degraded"], "text_layer_cjk_lost")
        self.assertNotIn("degraded", single.metadata)
        settings = SimpleNamespace(cache_dir=Path("/cache"), embedding_model_id="m")
        row = {"kb_id": "k", "file_key": 1, "source_path": "r/a.pdf", "rel_path": "a.pdf", "filename": "a.pdf",
               "dir": "", "content_version": "v", "size": 1, "mtime": 1, "mime_type": "application/pdf"}
        payload = _payload_for_chunk(settings, row, UnifiedChunk("u", 0, "t", table, 1), 1)
        self.assertEqual(payload["degraded"], "text_layer_cjk_lost")
        self.assertIsNone(_payload_for_chunk(settings, row, UnifiedChunk("u2", 1, "t", single, 1), 2).get("degraded"))
        self.assertEqual(source_row(1, {"point_id": "p"}, {"text": "t", **payload})["degraded"], "text_layer_cjk_lost")

    def test_second_pass_failures_are_told_apart(self) -> None:
        """Resending the rendered pages to MinerU fails: every such failure used to be swallowed, the lost pages were
        marked degraded and indexed as usual, the scan then judged the file unchanged and it was never retried.
        Transient failures (unreachable, 5xx) must be able to propagate; an explicit rejection by MinerU (4xx) and
        a rendering failure would fail the same way on retry, so they degrade in place."""
        from kb_pipeline.parsers import pdf_textlayer as tl
        from kb_pipeline.parsers.mineru_pdf import MinerUServiceError

        original = [self._block(0, "封面"), self._block(1, "1 2 3")]

        def repair(parse_rendered, **kw):
            return tl.repair_lost_text_layer(original, path=Path("/x.pdf"), rendered_path=Path("/tmp/none"), layer={0: 2, 1: 40},
                                             detail=lambda pages, boxes: {1: (40, 40)}, parse_rendered=parse_rendered, **kw)

        def failing(exc):
            def parse(_path):
                raise exc
            return parse

        with patch.object(tl, "render_pages") as render:
            for exc in (MinerUServiceError("500 Server Error", status_code=500), MinerUServiceError("connection reset"),
                        RuntimeError("boom")):
                with self.assertRaises(type(exc)):
                    repair(failing(exc), raise_transient=True)
                blocks, info = repair(failing(exc))                                  # last attempt: degrade in place
                self.assertEqual((blocks, info["lost"]), (original, [1]))
            blocks, info = repair(failing(MinerUServiceError("422 Unprocessable", status_code=422)), raise_transient=True)
            self.assertEqual((blocks, info["lost"], info["rendered"]), (original, [1], False))
            self.assertIn("422", info["error"])
            render.side_effect = OSError("disk full")                                # rendering failed: a retry would fail the same way
            blocks, info = repair(lambda p: [], raise_transient=True)
            self.assertEqual((info["lost"], info["error"]), ([1], "OSError: disk full"))
            render.reset_mock(side_effect=True)
            blocks, info = repair(None, raise_transient=True)                        # check only, no repair
            self.assertEqual((blocks, info["lost_before"], info["lost"]), (original, [1], [1]))
            render.assert_not_called()

    def _parse_pdf(self, tmp, fake_mineru, **kw):
        """Run parse_pdf_enhanced once: the main parse yields two pages, and page 1 is confirmed to have lost its
        text; fake_mineru stands in for MinerU, rendering is stubbed."""
        from kb_pipeline.parsers import pdf_enhanced, pdf_textlayer as tl

        with patch.object(pdf_enhanced, "mineru_pdf_blocks", fake_mineru), \
                patch.object(tl, "text_layer_cjk_by_page", return_value={1: 40, 2: 7}), \
                patch.object(tl, "page_details", return_value={1: (40, 40)}), \
                patch.object(tl, "render_pages", side_effect=lambda src, pages, out, **_: out) as render:
            blocks = pdf_enhanced.parse_pdf_enhanced(
                mineru_url="http://x", vlm_base_url="u", vlm_api_key="", vlm_model_id="m", vlm_concurrency=1,
                path=Path(tmp) / "a.pdf", cache_dir=Path(tmp) / "cache", **kw)
        return blocks, render

    def test_transient_second_pass_failure_retries_the_job_and_degrades_on_the_last_attempt(self) -> None:
        from kb_pipeline.parsers.mineru_pdf import MinerUServiceError

        def fake_mineru(*, path, **kw):
            if path.name.endswith(".lost-pages.pdf"):
                raise MinerUServiceError("500 Server Error", status_code=500)
            return [self._block(1, "1 2 3"), self._block(2, "第二页正文照旧")]

        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(MinerUServiceError):
                self._parse_pdf(tmp, fake_mineru)                                    # retries left: the job fails and runs again
            blocks, _ = self._parse_pdf(tmp, fake_mineru, text_layer_repair="final")
        self.assertEqual([b.metadata.get("degraded") for b in blocks], ["text_layer_cjk_lost"])
        self.assertIn("第二页正文照旧", blocks[0].text)

    def test_preview_mode_never_sends_the_rendered_pages_to_mineru(self) -> None:
        """The chunk preview runs in the web process: without a resend result in the parse cache the pages are only
        marked; with one, only the cache is read."""
        seen = []

        def fake_mineru(*, path, cached_only=False, **kw):
            seen.append((path.name, cached_only))
            if path.name.endswith(".lost-pages.pdf"):
                return [self._block(1, "第一页救回来了 窦性心律 未见明显异常 建议定期复查 心电图报告结论正常")]
            return [self._block(1, "1 2 3"), self._block(2, "第二页正文照旧")]

        with tempfile.TemporaryDirectory() as tmp:
            blocks, render = self._parse_pdf(tmp, fake_mineru, text_layer_repair="cached")
            self.assertEqual(seen, [("a.pdf", False)])
            render.assert_not_called()
            self.assertEqual([b.metadata.get("degraded") for b in blocks], ["text_layer_cjk_lost"])
            cached = Path(tmp) / "cache" / "rendered" / "mineru" / "result.json"
            cached.parent.mkdir(parents=True)
            cached.write_text("{}", encoding="utf-8")
            blocks, _ = self._parse_pdf(tmp, fake_mineru, text_layer_repair="cached")
            self.assertEqual(seen[1:], [("a.pdf", False), ("a.lost-pages.pdf", True)])
            self.assertIn("窦性心律", blocks[0].text)
            self.assertNotIn("degraded", blocks[0].metadata)

    @staticmethod
    def _fake_mineru_service(payload, posts):
        """Stand-in for the MinerU HTTP session: records the form of every request and returns a fixed result."""
        class Session:
            def __enter__(self):
                return self

            def __exit__(self, *exc):
                return False

            def post(self, url, files=None, data=None, timeout=None):
                posts.append(dict(data))
                return SimpleNamespace(raise_for_status=lambda: None, json=lambda: json.loads(json.dumps(payload)))
        return Session

    def test_mineru_cache_is_reused_only_for_the_same_file_and_backend(self) -> None:
        """The MinerU result cache used to check the file size only: after MINERU_BACKEND changed, a re-parse was
        served the old backend's result as is, and the new backend was never called."""
        import os

        from kb_pipeline.parsers import service_clients

        posts: list[dict] = []
        session = self._fake_mineru_service({"backend": "vlm-engine", "version": "3.4.4", "content_list": []}, posts)
        with tempfile.TemporaryDirectory() as tmp, patch.object(service_clients, "_local_service_session", session), \
                patch.dict(os.environ, {"MINERU_BACKEND": "vlm-engine", "MINERU_REUSE_CACHE": "1"}):
            src = Path(tmp) / "a.pdf"
            src.write_bytes(b"%PDF-1.4 demo")
            out = Path(tmp) / "mineru" / "result.json"
            call = lambda **kw: service_clients.call_mineru_sync("http://x", src, out, out.with_suffix(".md"), **kw)
            with self.assertRaisesRegex(RuntimeError, "no cached MinerU result"):
                call(cached_only=True)                                        # a cache-only caller: no cache means an error, no request
            self.assertEqual(posts, [])
            call()
            call()                                                            # same file, same backend: reused
            self.assertEqual(call(cached_only=True)["_kb_source"], {"name": "a.pdf", "size": 13, "backend": "vlm-engine"})
            self.assertEqual([p["backend"] for p in posts], ["vlm-engine"])
            with patch.dict(os.environ, {"MINERU_BACKEND": "pipeline"}):
                call()                                                        # backend changed: requested again
                self.assertEqual([p["backend"] for p in posts], ["vlm-engine", "pipeline"])
                call()
                self.assertEqual(len(posts), 2)
            src.write_bytes(b"%PDF-1.4 another")
            call()                                                            # file changed: requested again
            self.assertEqual(len(posts), 3)
            out.write_text(json.dumps({"content_list": [], "_kb_source": {"name": "a.pdf", "size": src.stat().st_size}}), encoding="utf-8")
            call()                                                            # the stamp records no backend: nothing to compare, accepted as before
            self.assertEqual(len(posts), 3)

    def test_parse_job_picks_the_mode_from_the_attempts_left(self) -> None:
        """With retries left a transient failure makes the job retry, and the last attempt degrades in place;
        callers outside a job (the chunk preview) read only the cache by default."""
        from kb_pipeline.pipeline import parse_job

        seen = []

        class Stop(Exception):
            pass

        def fake_parse_blocks(*args, **kw):
            seen.append(kw.get("text_layer_repair"))
            raise Stop()

        def fake_pdf(**kw):
            seen.append(kw["text_layer_repair"])
            return []

        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "a.pdf"
            pdf.write_bytes(b"%PDF-1.4 demo")
            settings = SimpleNamespace(max_file_bytes=0, job_max_retries=5, cache_dir=Path(tmp) / "cache", mineru_url="http://x",
                                       vlm_base_url="u", vlm_api_key="", vlm_model_id="m", vlm_concurrency=1, vlm_temperature=0.1,
                                       vlm_top_p=0.8, vlm_max_tokens=64, vlm_structured_output=True, image_max_pixels=1,
                                       mineru_timeout_seconds=5)
            source = SimpleNamespace(vlm_prompt=None, graph_language=None, block_merge_tokens=800, max_tokens=400, overlap_tokens=0)
            row = {"physical_path": str(pdf), "size": 13, "kb_id": "kb_x", "file_key": 1, "content_version": "v", "source_path": "r/a.pdf"}
            with patch.object(parse_job.db, "get_file_by_id", return_value=row), \
                    patch.object(parse_job, "_source_for_file", return_value=source), \
                    patch.object(parse_job, "verify_source_file"), \
                    patch.object(parse_job, "_parse_blocks", fake_parse_blocks):
                for retry_count in (0, 4, 5):
                    job = {"job_id": "j", "file_id": "f", "parser_profile": "p", "retry_count": retry_count, "locked_by": None}
                    with self.assertRaises(Stop):
                        parse_job.process_parse_job(None, settings, job)
            self.assertEqual(seen, ["retry", "retry", "final"])
            with patch.object(parse_job, "parse_pdf_enhanced", fake_pdf):
                parse_job._parse_blocks(settings, source, pdf, Path(tmp) / "cache", "p", row)
                parse_job._parse_blocks(settings, source, pdf, Path(tmp) / "cache", "p", row, text_layer_repair="final")
            self.assertEqual(seen[3:], ["cached", "final"])
