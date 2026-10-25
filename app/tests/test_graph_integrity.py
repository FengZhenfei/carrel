"""Regressions for the pinned whole-repository audit; synthetic inputs, no external services."""
from pathlib import Path
from types import SimpleNamespace as NS
import json
import time
from unittest.mock import patch

import pytest

from kb_pipeline import db, discovery
from kb_pipeline.graph import corrections, currency
from kb_pipeline.graph.units import Unit


def unit(uid="u", path="rows.xlsx", text=None):
    return Unit(unit_id=uid, doc_id="d"+uid, rel_path=path, section_path=["Sheet"], block_ids=["b"],
                chunk_uids=["c"+uid], point_ids=["p"+uid], n_tokens=50,
                text=text or "SHEET: Sheet\nROWS: 1-3\nHEADER: ID | Name | Owner | Price\nE01 | Widget | Alice | 100\nE02 | Widget | Bob | 200\nE03 | Other | Cara | 300")


def table_fixture(units):
    from kb_pipeline.graph.tables import TableExtractor
    from kb_pipeline.graph.extract import ExtractionSchema
    schema = ExtractionSchema(entity_types=("product","person"), parent_types={"product":"entity", "person":"agent"})
    tx = TableExtractor(units, schema=schema)
    tx.load_profiles({key: {"row_type":"product", "key_column":"ID", "name_column":"Name", "date_column":"",
                           "references":[{"column":"Owner", "target_type":"person", "predicate":"owned_by"}]}
                      for tables in tx.unit_tables.values() for key, _ in tables})
    return tx, schema


def test_table_identity_reaches_relationships_and_facts():
    from kb_pipeline.graph.merge import merge_extractions
    from kb_pipeline.graph.facts import link_facts
    u=unit();tx,schema=table_fixture([u]);ex=tx.extract(u)
    graph=merge_extractions([u], {u.unit_id:{"entities":ex.entities,"relations":ex.relations}},
                           entity_types=schema.entity_types, upper_parents=schema.parent_types)
    facts=tx.facts(u);link_facts(facts,graph["entities"],units_by_id={u.unit_id:u})
    ids={e["key"]:e["external_ids"][0]["value"] for e in graph["entities"] if e.get("external_ids")}
    assert {(ids[f["subject_key"]],f["value"]) for f in facts if f["property"]=="Price"} == {("E01","100"),("E02","200"),("E03","300")}
    assert {(ids[r["source_key"]],r["target"]) for r in graph["relations"]} == {("E01","Alice"),("E02","Bob"),("E03","Cara")}
    vague=[{"subject":"Widget","property":"Price","value":"999","unit_id":u.unit_id}]
    link_facts(vague,graph["entities"],units_by_id={u.unit_id:u})
    assert not vague[0].get("subject_key")


def test_different_cross_document_conditions_do_not_supersede():
    from kb_pipeline.graph.reconcile import reconcile_facts
    from kb_pipeline.graph.facts import normalize_fact
    rows=[]
    for n,temp,day in [(1,"25 C","2025-01-01"),(2,"85 C","2026-01-01")]:
        f=normalize_fact({"subject":"Part X","property":"Leakage","value":str(n),"unit":"mA","valid_from":day,"conditions":{"temperature":temp}})
        f.update(id=f"f{n}",doc_id=f"d{n}");rows.append(f)
    reconcile_facts(rows)
    assert all(not r.get("superseded_by") for r in rows)


def test_pdf_full_text_dedup_and_metadata(tmp_path):
    from kb_pipeline.parsers import mineru_pdf
    from kb_pipeline.parsers.common import walk_dicts,item_text
    prefix="shared terms "*30
    payload={"content_list":[{"type":"text","text":prefix+tail,"page_idx":0} for tail in ("100","200","100")]}
    with patch.object(mineru_pdf,"call_mineru_sync",return_value=payload):
        blocks=mineru_pdf.mineru_pdf_blocks(mineru_url="unused",path=tmp_path/"x.pdf",cache_dir=tmp_path)
    assert len(blocks)==2 and any(b.text.endswith("200") for b in blocks)
    assert not any(item_text(x) for x in walk_dicts({"_kb_source":{"name":"x.pdf"}}))


def test_bad_mineru_cache_is_retried_and_bad_response_not_saved(tmp_path):
    from kb_pipeline.parsers import service_clients as sc
    p=tmp_path/"x.pdf";p.write_bytes(b"pdf")
    cache=tmp_path/"result.json";cache.write_text(json.dumps({"_kb_source":{"name":"x.pdf","size":3}}))
    class Session:
        def __enter__(self):return self
        def __exit__(self,*a):pass
        def post(self,*a,**kw):return NS(raise_for_status=lambda:None,json=lambda:{})
    with patch.object(sc,"_local_service_session",return_value=Session()) as post:
        for _ in range(2):
            with pytest.raises(RuntimeError,match="no valid document"):
                sc.call_mineru_sync("unused",p,cache,tmp_path/"result.md")
        assert post.call_count==2
    assert not sc.valid_mineru_result(json.loads(cache.read_text()))


@pytest.mark.parametrize("value",[0,False])
def test_correction_falsy_value_is_not_wildcard(value):
    corr=corrections.parse_corrections({"facts":{"reject":[{"subject":"X","property":"price","value":value}]}})
    kept,n=corrections.apply_fact_corrections([{"subject":"X","property":"price","value":str(value)},
                                             {"subject":"X","property":"price","value":"100"}],corr)
    assert n==1 and kept[0]["value"]=="100"


@pytest.mark.parametrize("doc",[{"entities":{"merge":42}},{"entities":{"separate":True}},{"documents":42},{"facts":{"reject":True}},{"relations":True}])
def test_malformed_corrections_are_validation_errors(doc):
    assert corrections.parse_corrections(doc)["errors"]


def test_future_effective_current_matches_as_of():
    old={"superseded_by":"new","superseded_on":"2027-01-01"}
    new={"valid_from":"2027-01-01"}
    assert currency.row_status(old,day="2026-10-25")=="current"
    assert currency.row_status(new,day="2026-10-25")=="not_started"
    assert currency.row_status(old,day="2027-01-01")=="superseded"
    assert currency.row_status(new,day="2027-01-01")=="current"
    assert currency.next_change_day([old,new,{"valid_until":"2026-11","lapses":True}],today="2026-10-25")=="2026-12-01"
    assert currency.next_change_day([old,new],today="2027-01-01")==""


def test_xlsx_fallback_preserves_hidden_format_date_and_merge(tmp_path):
    import datetime
    from openpyxl import Workbook
    from kb_pipeline.parsers import native_table as nt
    w=Workbook();s=w.active;s.title="Rates"
    s.append(["Group","Rate","Date"]);s.append(["A",.06,datetime.datetime(2026,1,2)])
    s["B2"].number_format="0%";s.append([None,.12,datetime.datetime(2026,2,3)]);s["B3"].number_format="0%"
    s.merge_cells("A2:A3");s.append(["hidden",999]);s.row_dimensions[4].hidden=True
    hidden=w.create_sheet("Hidden");hidden.append(["secret",888]);hidden.sheet_state="hidden"
    p=tmp_path/"x.xlsx";w.save(p)
    expected=nt.read_xlsx_rows(p)
    with patch.object(nt,"load_workbook_compat",side_effect=ValueError("bad backend")):
        assert nt.read_xlsx_rows(p)==expected


@pytest.mark.parametrize("code",[401,403,404,400,422,429])
def test_visual_service_configuration_errors_fail_job(tmp_path,code):
    from PIL import Image
    from kb_pipeline.embedding.visual import VisualEmbeddingClient
    image=tmp_path/"x.png";Image.new("RGB",(4,4)).save(image)
    client=VisualEmbeddingClient(base_url="http://unused",api_key="synthetic",model_id="test",dim=2,retry=1)
    with patch.object(client.session,"post",return_value=NS(status_code=code,text="invalid API key or model")):
        with pytest.raises(RuntimeError):client.embed_images([(image,None)])


def test_table_profile_change_does_not_invalidate_unrelated_unit():
    from kb_pipeline.graph.unit_cache import fingerprints
    prose=unit("prose","a.pdf","a normal paragraph");a=unit("a","a.xlsx");b=unit("b","b.xlsx")
    tx,_=table_fixture([prose,a,b]);before=fingerprints("base",[prose,a,b],set(),tx)
    key=tx.unit_tables[b.unit_id][0][0];tx.profiles[key]["references"]=[]
    after=fingerprints("base",[prose,a,b],set(),tx)
    assert before[prose.unit_id]==after[prose.unit_id]=="base"
    assert before[a.unit_id]==after[a.unit_id]
    assert before[b.unit_id]!=after[b.unit_id]


def test_version_ring_keeps_build_references():
    from kb_pipeline.limits import push_schema_version
    result=push_schema_version([{"id":str(i)} for i in range(4)],{"id":"new"},active_id="0",protected_ids={"1","2","3"})
    assert {r["id"] for r in result}=={"new","0","1","2","3"}


def test_schema_spaces_round_trip_and_error_languages():
    from kb_pipeline.ontology import normalize_model,to_linkml,from_linkml,SchemaError
    model,_=normalize_model({"types":[{"name":"sales order","is_a":"entity"},{"name":"staff member","is_a":"agent"}],
                            "predicates":[{"name":"owned_by","source_parents":["sales order"],"target_parents":["staff member"]}]})
    result,_=from_linkml(to_linkml(model,name="test"))
    pred=next(p for p in result["predicates"] if p["name"]=="owned_by")
    assert pred["source_parents"]==["sales order"] and pred["target_parents"]==["staff member"]
    with pytest.raises(SchemaError) as err:normalize_model({})
    assert err.value.details[0]["code"]=="no_concrete_types"
    assert "non-abstract" in err.value.details[0]["message"]["en"]


def test_gc_reenroll_between_selection_and_claim_survives(tmp_path):
    from kb_pipeline import maintenance
    mirror=tmp_path/"mirror";(mirror/"test").mkdir(parents=True)
    state=tmp_path/"s.db";db.init_db(state)
    with db.connect(state) as con:
        src,_=discovery.enroll(con,mirror,"test")
        con.execute("UPDATE kb_sources SET status='inactive',inactive_at=1,inactive_reason='unenrolled'")
    settings=NS(state_db=state,mirror_root=mirror,runtime_dir=tmp_path,corrections_dir=None)
    def directories(_):
        with db.connect(state) as con:
            with patch.object(discovery,"discover_directories",return_value=["test"]):discovery.enroll(con,mirror,"test")
        return ["test"]
    with patch.object(maintenance,"service_busy",return_value=(False,[])),patch.object(discovery,"discover_directories",side_effect=directories),patch.object(maintenance,"_hard_delete_kb") as drop:
        result=maintenance.kb_sources_gc(settings,retention_days=7)
    assert not drop.called and not result["dropped"]


def test_api_rejects_null_protected_fields():
    from kb_server.api import _reject_bad_config
    from fastapi import HTTPException
    for field in ("graph_schema_versions","graph_entity_types","graph_paused","graph_schema_active"):
        with pytest.raises(HTTPException):_reject_bad_config({field:None})


@pytest.mark.parametrize("doc",[
    {"entities":{"aliases":{"X":{"bad":"nested"}}}},
    {"facts":{"reject":[{"subject":"X","value":{}}]}},
    {"documents":[{"path":"old.pdf","valid_until":"2026-02-31"}]},
    {"documents":[{"path":"old.pdf","valid_from":"2027-01-01","valid_until":"2026-12-31"}]},
])
def test_nested_correction_values_and_dates_fail_validation(doc):
    assert corrections.parse_corrections(doc)["errors"]


def test_final_status_filter_exceeds_storage_cap_and_reports_bad_rules():
    from kb_search import service, docstatus
    source=NS(kb_id="kb1");settings=NS(sources={"kb1":source})
    corr=corrections.parse_corrections({"documents":[{"path":"old/*","status":"archived"}]})
    paths=[f"old/{i:05}.pdf" for i in range(docstatus.EXCLUDED_MAX+1)]
    scope={"mode":"current","day":"2026-10-25"}
    def candidate(path):return {"kb_id":"kb1","point_id":path,"payload":{"text":"evidence","rel_path":path}}
    with patch.object(docstatus,"corrections_of",return_value=corr),patch.object(docstatus,"_doc_paths",return_value=paths),patch.object(service,"rerank_active",return_value=False):
        excluded=docstatus.excluded_paths(settings,source,scope)
        assert paths[-1] not in excluded
        assert not docstatus.allowed(settings,source,None,scope)
        result=service._rank(settings,NS(rerank_n=64),None,"query",[[candidate(paths[-1]),candidate("current.pdf")]],until=time.monotonic()+10,scope=scope)
        assert [c["point_id"] for c in result["ordered"]]==["current.pdf"]
        corr["errors"]=["documents must be a list"]
        result=service._rank(settings,NS(rerank_n=64),None,"query",[[candidate("current.pdf")]],until=time.monotonic()+10,scope=scope)
        assert not result["ordered"] and "invalid_corrections" in result["degraded"][0]


def test_document_date_boundary_and_removed_corrections_require_refresh(tmp_path):
    from kb_search import docstatus
    settings=NS(state_db=tmp_path/"s.db",corrections_dir=tmp_path)
    source=NS(kb_id="kb1",collection="kb1")
    p=tmp_path/"kb1.yaml";p.write_text("documents:\n  - {path: old.pdf, valid_until: 2026-10-24}\n")
    corr=corrections.load_corrections(p);db.init_db(settings.state_db)
    with db.connect(settings.state_db) as con:
        bid=db.begin_graph_build(con,source_key="kb1",kb_id="kb1",source_collection="kb1",graph_version="v1")
        db.finish_graph_build(con,bid,status="done",manifest={"corrections_sha":corrections.effective_sha(corr),
                              "document_status_sha":corrections.status_fingerprint(corr,["old.pdf"],today="2026-10-24")})
    with patch.object(db,"active_chunk_refs",return_value=[{"rel_path":"old.pdf"}]),patch.object(docstatus.time,"strftime",return_value="2026-10-25"):
        assert docstatus.graph_refresh_pending(settings,source)
        p.write_text("documents: []\n")
        assert docstatus.graph_refresh_pending(settings,source)


def test_summaries_and_pages_drop_archived_and_future_sources():
    from kb_pipeline.graph.current_view import filter_descriptions,graph_for_pages
    from kb_search import docstatus
    corr=corrections.parse_corrections({"documents":[{"path":"old.pdf","status":"archived"}]})
    rows=[{"description":"mixed old summary","descriptions":["old price 100","current price 200"],
           "description_sources":[{"source":"old.pdf"},{"source":"current.pdf"}]}]
    filter_descriptions(rows,corr,today="2026-10-25")
    assert rows[0]["descriptions"]==["current price 200"] and not rows[0]["description"]
    graph={"documents":{"old":{"status":"archived"},"new":{}},
           "entities":[{"key":"old","doc_ids":["old"]},{"key":"new","doc_ids":["new"]}],
           "relations":[{"doc_ids":["new"],"valid_from":"2027-01-01"}],"specs":[]}
    projected,_=graph_for_pages(graph,[],today="2026-10-25")
    assert [e["key"] for e in projected["entities"]]==["new"] and not projected["relations"]
    with patch.object(docstatus,"corrections_of",return_value=corr):
        assert not docstatus.page_allowed(None,None,{"source_paths":["current.pdf","old.pdf"]},{"mode":"current"})
        assert not docstatus.page_allowed(None,None,{"source_paths":["current.pdf"],"time_mode":"history"},{"mode":"current"})
        assert not docstatus.page_allowed(None,None,{"compiled_day":"2026-10-24"},{"mode":"as_of","day":"2025-01-01"})


def test_cleanup_nested_error_is_nonzero(capsys):
    from kb_pipeline import cli
    args=NS(cleanup_command="parse-assets-gc",env_file=None,retention_days=7,dry_run=True)
    with patch.object(cli,"load_settings",return_value=NS()),patch.object(cli,"wait_out_short_locks"),patch.object(cli,"parse_assets_gc",return_value={"errors":[]}),patch.object(cli,"kb_sources_gc",return_value={"errors":["delete failed"]}):
        assert cli.cmd_cleanup(args)==1
    assert "delete failed" in capsys.readouterr().out


def test_prior_resolution_cannot_cross_current_type_scope_or_upper():
    from kb_pipeline.graph.resolution import merge_entities
    base={"key":"a","title":"A","type":"product","upper":"entity","scope":"global"}
    for changed in ({"type":"person","upper":"agent"},{"scope":"document:d2"},{"upper":"agent"}):
        other={**base,"key":"b","title":"B",**changed}
        rows,_,stats=merge_entities([base,other],[],[(0,1)],require_compatible_identity=True)
        assert len(rows)==2 and stats["merged_away"]==0


def test_unique_record_id_links_document_scoped_reference():
    from kb_pipeline.graph.merge import link_external_ids
    rows=[{"key":"feature","title":"Feature A","type":"candidate feature","upper":"concept","scope":"doc-a",
           "external_ids":[{"scheme":"feature_id","value":"F03"}]},
          {"key":"reference","title":"F03","type":"candidate feature","upper":"concept","scope":"doc-b"}]
    linked,_,stats=link_external_ids(rows,[])
    assert len(linked)==1 and linked[0]["key"]=="feature"
    assert stats["external_id_mentions_linked"]==1


def test_legacy_cache_migration_and_pruning_preserve_references(tmp_path):
    from kb_pipeline.graph import unit_cache
    state=tmp_path/"s.db";db.init_db(state)
    with db.connect(state) as con:
        for fp in ("old","unused1","unused2","unused3"):
            db.save_graph_extraction(con,kb_id="kb1",unit_id="u",fingerprint=fp,entities=[{"name":"X"}],relations=[],model="test",calls=1)
        db.save_graph_extraction(con,kb_id="kb1",unit_id="table",fingerprint="old",entities=[{"name":"Y"}],relations=[],model="table",calls=0)
        bid=db.begin_graph_build(con,source_key="kb1",kb_id="kb1",source_collection="kb1",graph_version="v1",cache_fingerprint="config")
        db.finish_graph_build(con,bid,status="done",manifest={"extract":{"fingerprint":"old"}})
        assert unit_cache.migrate_legacy(con,NS(kb_id="kb1"),"changed",{"u":"new"})==0
        assert unit_cache.migrate_legacy(con,NS(kb_id="kb1"),"config",{"u":"new","table":"new"})==1
        assert unit_cache.load(con,"kb1",{"u":"new"})["u"]["entities"][0]["name"]=="X"
        db.update_graph_build_manifest(con,bid,{"cache_fingerprints":{"extract":{"u":"old"}}})
        assert db.prune_graph_extractions(con,"kb1",["u"])==3
        remaining={r[0] for r in con.execute("SELECT fingerprint FROM graph_extractions WHERE kb_id='kb1'")}
        assert remaining=={"old","unused3","new"}


def test_schema_type_response_is_validated_before_caching():
    from kb_pipeline.graph import schema
    class Checked(Exception):pass
    class Client:
        def chat(self,messages,**kw):
            if isinstance(messages,list):
                validate=kw["validate"]
                assert not validate('{"entity_types": []}')
                assert validate('{"entity_types": ["product"]}')
                raise Checked()
            return "sample domain"
    with pytest.raises(Checked):schema.suggest(Client(),["sample input"])


def test_lock_script_failure_preserves_existing_output(tmp_path):
    import subprocess,sys,ast
    root=Path(__file__).resolve().parents[2]
    script=root/"scripts/kb-lock-deps.py";ast.parse(script.read_text())
    project=tmp_path/"pyproject.toml";project.write_text('[project]\ndependencies = ["carrel-nonexistent-test-package"]\n')
    target=tmp_path/"requirements.lock";target.write_text("keep-original\n")
    result=subprocess.run([sys.executable,str(script),str(project),"--output",str(target)],capture_output=True,text=True)
    assert result.returncode!=0 and "not installed:" in result.stderr
    assert target.read_text()=="keep-original\n"


def test_dependency_lock_includes_transitive_extras(monkeypatch):
    import importlib.util

    path = Path(__file__).resolve().parents[2] / "scripts/kb-lock-deps.py"
    spec = importlib.util.spec_from_file_location("lock_deps_test", path)
    script = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(script)
    requires = {"client": ["transport[http2]>=1"], "transport": ["h2>=1; extra == 'http2'"],
                "h2": ["hpack>=1"], "hpack": []}
    monkeypatch.setattr(script.metadata, "distribution", lambda name:
                        NS(metadata={"Name": name}, version="1.0", requires=requires[name]))
    # transport is visited without extras first, then requested with http2 by client.
    assert script.closure(["client", "transport"]) == {name: f"{name}==1.0" for name in requires}


def test_incremental_build_end_to_end_cache_delete_rename_resume_and_publish(tmp_path,monkeypatch):
    """Real orchestration/SQLite/merge/pages/vector publication with only service boundaries faked."""
    from qdrant_client import QdrantClient
    from kb_pipeline.graph import build as gb, units as gu, neo4j_import as ni
    from kb_pipeline.graph.extract import ExtractionResult
    from kb_pipeline.graph.llm import LLMSpec,ChatClient
    from _support import _schema_entry

    settings=NS(state_db=tmp_path/"state.db",runtime_dir=tmp_path/"runtime",graph_work_dir=tmp_path/"graph",
                mirror_root=tmp_path/"mirror",corrections_dir=tmp_path/"corrections",qdrant_url="unused",qdrant_api_key="",
                graph_neo4j_import_after_build=True,graph_neo4j_import_batch_size=32,graph_llm_timeout_seconds=2,
                graph_llm_concurrency=1,graph_circuit_fails=0,embedding_dim=4,embedding_model_id="synthetic")
    (settings.mirror_root/"test").mkdir(parents=True);settings.corrections_dir.mkdir()
    db.init_db(settings.state_db)
    with db.connect(settings.state_db) as con:
        source,_=discovery.enroll(con,settings.mirror_root,"test")
        discovery.set_config(con,source.kb_id,{"graph_enabled":True,"graph_schema_versions":[_schema_entry("s1",["product"])],
                                             "graph_schema_active":"s1","graph_entity_types":["product"],"graph_parent_types":{"product":"entity"}})
        source=discovery.source_from_row(settings.mirror_root,con.execute("SELECT * FROM kb_sources").fetchone())
    corr=settings.corrections_dir/f"{source.kb_id}.yaml"
    corpus=[]; refs=[]; extracted=[]
    def set_corpus(items):
        corpus.clear();refs.clear()
        for index,(path,name) in enumerate(items):
            uid=f"{path}-{name}"
            c=gu.ChunkRef(point_id=uid,chunk_uid=uid,doc_id=f"{source.kb_id}:{path}",content_version=uid,
                          chunk_index=0,block_id="b1",block_type="text",section_path=["Body"],text=f"{name} is documented here.",
                          n_tokens=10,rel_path=path,filename=path)
            corpus.append(c)
            refs.append({"point_id":uid,"chunk_uid":uid,"doc_id":c.doc_id,"rel_path":path,"content_version":uid,"text_sha":db.chunk_text_sha(c.text)})
    class Extractor:
        fingerprint="synthetic-entity-v1";max_tokens=100
        def __init__(self,*a,**kw):pass
        def extract(self,u):
            name=u.text.split()[0];extracted.append((u.rel_path,name))
            return ExtractionResult(unit_id=u.unit_id,entities=[{"name":name,"type":"product","types":{"product":1},"mentions":1,"descriptions":[u.text]}],
                                    relations=[],calls=1,stats={"unit_kind":"body"})
    q=QdrantClient(":memory:")
    spec=LLMSpec(name="synthetic",base_url="http://unused",api_key="",model_id="synthetic",protocol="openai")
    live={"version":None,"fail":False}
    def activate(settings,*,source,graph_version):
        if live["fail"]:raise RuntimeError("synthetic activation failure")
        live["version"]=graph_version;return {"active":graph_version}
    monkeypatch.setenv("KB_GRAPH_VIEWS_NO_NARRATE","1")
    monkeypatch.setattr(gb,"qdrant_client",lambda *a,**kw:q)
    monkeypatch.setattr(db,"active_chunk_refs",lambda *a,**kw:list(refs))
    monkeypatch.setattr(gu,"fetch_chunk_payloads",lambda *a,**kw:list(corpus))
    monkeypatch.setattr(gb,"GraphExtractor",Extractor)
    monkeypatch.setattr(gb,"resolve_llm_specs",lambda *a,**kw:{"extract":spec,"summarize":spec})
    monkeypatch.setattr(gb,"graph_embedder",lambda *a,**kw:NS(embed=lambda texts:[[1.,0.,0.,0.] for _ in texts]))
    monkeypatch.setattr(gb,"_phase_retry_config",lambda:{"retries":1,"backoff":0})
    monkeypatch.setattr(ChatClient,"chat",lambda *a,**kw:"1. no\n2. no\n3. no")
    monkeypatch.setattr(ni,"neo4j_driver",lambda *a:NS(verify_connectivity=lambda:None,close=lambda:None))
    monkeypatch.setattr(ni,"import_graph_to_neo4j",lambda *a,**kw:{"imported":True})
    monkeypatch.setattr(ni,"activate_neo4j_graph_version",activate)
    def run(version,**kw):
        return gb.build_graph(settings,source_key=source.kb_id,source=source,graph_version=version,run_gc=False,**kw)
    def graph(version):return json.loads(gb.graph_paths(settings,source,version).graph_file.read_text())
    def aliases():return {a.alias_name:a.collection_name for a in q.get_aliases().aliases}
    set_corpus([("one.pdf","A100"),("two.pdf","B200")]);run("v1")
    assert len(extracted)==2 and live["version"]=="v1"
    set_corpus([("one.pdf","A100"),("two.pdf","B200"),("three.pdf","C300")]);second=run("v2",incremental=True)
    assert len(extracted)==3 and second["extract"]["cached"]==2
    set_corpus([("one.pdf","A100"),("renamed.pdf","C300")]);run("v3",incremental=True)
    assert {d["rel_path"] for d in graph("v3")["documents"].values()}=={"one.pdf","renamed.pdf"}
    assert {e["title"] for e in graph("v3")["entities"]}=={"A100","C300"}
    old_alias=aliases();before=len(extracted)
    corr.write_text("entities:\n  reject: [A100]\n")
    live["fail"]=True
    with pytest.raises(RuntimeError,match="activation failure"):run("v4",incremental=True)
    assert aliases()==old_alias and live["version"]=="v3"
    # Change corrections while stopped. Resume must redo derived phases, but not model extraction.
    corr.write_text("entities:\n  reject: [C300]\n")
    live["fail"]=False
    resumed=run("v4",incremental=True,allow_existing_graph_version=True)
    assert "merge" not in resumed["resumed_phases"] and "extract" in resumed["resumed_phases"]
    assert {e["title"] for e in graph("v4")["entities"]}=={"A100"}
    assert len(extracted)==before and live["version"]=="v4"
    run("v5")
    assert {e["key"] for e in graph("v4")["entities"]}=={e["key"] for e in graph("v5")["entities"]}
    assert all(q.collection_exists(name) for name in old_alias.values())
    q.close()
